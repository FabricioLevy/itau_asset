# ─────────────────────────────────────────────
# ntnb/__main__.py
# Cota de NTN-B com e sem reinvestimento do cupom, com dados ANBIMA do Broadcast+, e carga na tabela
# DBSX337.dbo.ntnb_cotas — substitui a cota_acumulada do dbsx337.dbo.FUNDOS_DADOS_ATIVOS (ex.:
# BRSTNCNTB690 = NTN-B 15/08/2060) quando o banco não está acessível.
#
# Títulos (input/titulos.csv, 1 linha por NTN-B): codigo_IAM_ativo;vencimento;data_inicio;cota;data_cota
#   codigo_IAM_ativo = ISIN, o mesmo código do FUNDOS_DADOS_ATIVOS. A série começa na data_inicio (ou no
#   1º PU, se o título foi emitido depois) e é escalada para a série reinvestindo valer "cota" em
#   data_cota (vazia = 1º dia) — com a cota_acumulada do banco num dia qualquer, a série continua a dele.
#   Sem cota, começa em 1. input/titulos_do_banco.sql gera essas linhas a partir do cadastro do DBSX337.
#   NTN-B vencida não entra: o Broadcast apaga o código no vencimento (#BCPERM).
#
# Saída (./output/):
#   ntnb_<aaaammdd>.xlsx: 1 aba por NTN-B — Data | Taxa indicativa | PU | VNA | Cupom pago |
#     Cota reinvestindo o cupom | Cota sem reinvestir | Cota só preço; abas "Cupons" e "Fonte".
#   ntnb_cotas_<aaaammdd>.csv: as mesmas séries no layout da dbo.ntnb_cotas, para gravar no SQL
#     numa máquina com acesso ao banco (--carregar).
#
# SQL (--sql ou --carregar): DBSX337.dbo.ntnb_cotas (DDL em input/create_ntnb_cotas.sql), pela conexão
#   do vettafi/src/sql_connection.py (SQNPRC028; credenciais SQL_LOGIN e SQL_PASSWORD). Cada carga
#   substitui a série inteira de cada título do arquivo (DELETE + INSERT na mesma transação).
#   cota_acumulada = reinvestindo o cupom (mesmo nome do FUNDOS_DADOS_ATIVOS).
#
# Fontes (Broadcast+, função BCH do add-in Excel, campo ULT — mesma conexão do captacao.py):
#   NTNB<aammdd>    taxa indicativa ANBIMA (% a.a.) (ex.: NTNB600815 = NTN-B 15/08/2060)
#   VNANTB          VNA da NTN-B (ANBIMA)
#   NB<aammdd>P-AE  PU indicativo ANBIMA — conferência, e reserva nos dias sem taxa
#   Códigos achados no cache do terminal (bctab.cache) e no painel IMA-B; o histórico vai até a emissão
#   (ou antes de 2021).
#   PU usado = PU ANBIMA pela taxa e o VNA (pu_anbima: reproduz o PU do Broadcast nos ~15 mil dias bons).
#   O feed de PU do Broadcast às vezes troca o PU entre títulos (24-27/08/2026, 15/03/2023, 19/02/2021)
#   ou traz o do dia seguinte/anterior; a taxa não. Dias divergentes saem na aba "Avisos" do Excel.
#
# Cálculo (base = cota no 1º dia, igual para as 3 séries; vem da "cota" da lista):
#   cupom (R$) no dia de pagamento = VNA x ((1,06)^(1/2) - 1), sem truncar (como o banco).
#     Pagamento: dia 15 do mês de cupom (fev/ago p/ vencimento em agosto, mai/nov p/ maio) ou
#     o dia útil seguinte (calendário ANBIMA); o PU ANBIMA desse dia já é ex-cupom.
#   reinvestindo:   cota(t) = cota(t-1) x (PU(t) + cupom(t)) / PU(t-1)  — cupom recompra o título
#   sem reinvestir: cota(t) = base x (PU(t) + cupons recebidos até t) / PU(0) — cupom fica em caixa, sem render
#   só preço:       cota(t) = base x PU(t) / PU(0)  — o cupom sai e não volta
#   Título vencido: a série vai até o último PU antes do vencimento (o resgate não entra).
#   Conferido contra o FUNDOS_DADOS_ATIVOS (BRSTNCNTB690, 02/01 a 27/02/2023, cupom de 15/02):
#   a cota_acumulada do banco é a série "reinvestindo".
#
# Requisitos: Broadcast — terminal Broadcast+ aberto e logado + Excel 64-bit com o add-in; pywin32,
#   pandas, openpyxl (anaconda). SQL — ODBC Driver 18 for SQL Server, pyodbc, sqlalchemy.
#
# Uso (de dentro de ntnb/, ou "python ntnb" a partir de itau_asset/):
#   python __main__.py                   (todas as NTN-B da lista, até hoje: Excel + CSV)
#   python __main__.py BRSTNCNTB690      (só as escolhidas)
#   python __main__.py --sql             (e grava no SQL: máquina com Broadcast e acesso ao banco)
#   python __main__.py --carregar        (só grava no SQL o CSV mais recente de output/: máquina com
#                                         acesso ao banco, sem Broadcast)
# ─────────────────────────────────────────────

import argparse
import functools
import glob
import math
import os
import sys
import time
from datetime import date, datetime, timedelta
from decimal import Decimal

import numpy as np
import pandas as pd

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(AQUI, "..", "captacao", "src"))
import captacao as cp  # noqa: E402 — conexão Excel/Broadcast (_com, add-in, erros do BCH) e calendário ANBIMA

LISTA_PADRAO = os.path.join(AQUI, "input", "titulos.csv")
SAIDA_DIR = os.path.join(AQUI, "output")
SQL_CONNECTION_DIR = os.path.join(AQUI, "..", "vettafi", "src")  # sql_connection.py (SQNPRC028/DBSX337)

CODIGO_VNA = "VNANTB"
# (1,06)^(1/2) - 1 sem truncar, como o FUNDOS_DADOS_ATIVOS (o Tesouro trunca em 2,956301% e o cupom
# na 6ª casa: ~R$ 0,00002 a menos por cupom, 4e-9 na cota)
FATOR_CUPOM = 1.06 ** 0.5 - 1
CUPOM_ANBIMA = 2.956301            # cupom semestral em % do VNA na cotação ANBIMA
INICIO_20NOV = date(2023, 12, 26)  # 1º dia em que a ANBIMA conta o 20/11 (Lei 14.759/2023) até os fluxos
TOLERANCIA_PU = 1e-5               # PU do Broadcast x PU pela taxa (o Broadcast às vezes tem menos casas)

TABELA_SQL = "dbo.ntnb_cotas"
# coluna -> casas decimais, igual ao DECIMAL(p,s) de input/create_ntnb_cotas.sql (None = texto/data)
COLUNAS_SQL = {"codigo_IAM_ativo": None, "date": None, "vencimento": None, "taxa_indicativa": 4,
               "pu": 6, "vna": 6, "cupom": 8, "cota_acumulada": 16, "cota_sem_reinvestir": 16,
               "cota_preco": 16}
OPCIONAIS_SQL = {"taxa_indicativa", "vna"}  # as colunas NULL da tabela


def codigos(venc):
    aammdd = f"{venc:%y%m%d}"
    return {"pu": f"NB{aammdd}P-AE", "taxa": f"NTNB{aammdd}", "vna": CODIGO_VNA}


def nome(venc):
    return f"NTNB {venc:%Y}"


# ── Lista de títulos ─────────────────────────────────────────────────────────
def _data_flex(s):
    s = s.strip()
    return datetime.strptime(s, "%Y-%m-%d" if "-" in s else "%d/%m/%Y").date()


def _numero(s):
    s = (s or "").strip()
    if not s or s.upper() == "NULL":
        return None
    return float(s.replace(".", "").replace(",", ".") if "," in s else s)


def carregar_lista(caminho):
    """titulos.csv -> {codigo_IAM_ativo: {"venc", "inicio", "cota", "data_cota"}}. Separador ; ou tab
    (resultado do SSMS colado); linha com # é comentário e cabeçalho repetido é ignorado."""
    try:
        texto = open(caminho, encoding="utf-8-sig").read()
    except UnicodeDecodeError:
        texto = open(caminho, encoding="cp1252").read()
    titulos, vencs = {}, {}
    for n, linha in enumerate(texto.splitlines(), start=1):
        if not linha.strip() or linha.lstrip().startswith("#"):
            continue
        campos = [c.strip() for c in linha.split("\t" if "\t" in linha else ";")]
        if campos[0].lower() == "codigo_iam_ativo":
            continue
        try:
            codigo = campos[0].upper()
            venc, inicio = _data_flex(campos[1]), _data_flex(campos[2])
            cota = _numero(campos[3] if len(campos) > 3 else "")
            data_cota = campos[4] if len(campos) > 4 and campos[4].upper() not in ("", "NULL") else None
            data_cota = _data_flex(data_cota) if data_cota else None
        except (IndexError, ValueError) as e:
            raise ValueError(f"{caminho}, linha {n}: {linha.strip()!r} ({e})") from None
        erro = ("NTN-B vence em 15/05 ou 15/08" if venc.day != 15 or venc.month not in (5, 8)
                else "data_inicio depois do vencimento" if inicio >= venc
                else "cota precisa ser > 0" if cota is not None and cota <= 0
                else "data_cota sem cota" if data_cota and cota is None
                else "data_cota fora do período do título" if data_cota and not inicio <= data_cota < venc
                else f"título repetido (linha {titulos[codigo]['linha']})" if codigo in titulos
                else f"vencimento repetido ({vencs[venc]})" if venc in vencs else None)
        if erro:
            raise ValueError(f"{caminho}, linha {n}: {erro}")
        titulos[codigo] = {"venc": venc, "inicio": inicio, "cota": 1.0 if cota is None else cota,
                           "data_cota": data_cota, "linha": n}
        vencs[venc] = codigo
    if not titulos:
        raise ValueError(f"nenhuma NTN-B em {caminho}")
    return titulos


# ── Broadcast ────────────────────────────────────────────────────────────────
def _formula(codigo, inicio, fim):
    return f'=BCH("{codigo}","DRF;ULT","{inicio:%d/%m/%Y}","{fim:%d/%m/%Y}","ASC=DRF","TITULO=S")'


def _serie_vazia():
    return pd.Series(dtype=float, index=pd.DatetimeIndex([]))


def _ler_serie(ws):
    """Bloco DRF | ULT devolvido pelo BCH -> Series(data -> valor)."""
    valores = cp._com(lambda: ws.Range("A1").CurrentRegion.Value)
    if not isinstance(valores, tuple) or len(valores) < 2:
        return _serie_vazia()
    df = pd.DataFrame([list(v) for v in valores[1:]], columns=[str(t).strip().upper() for t in valores[0]])
    df = df.apply(pd.to_numeric, errors="coerce")
    # DRF é data OLE Automation; erro do Excel (#N/A...) chega como inteiro negativo e o BCH devolve 0 sem cotação
    df = df[df["DRF"].between(1, 2958465) & (df["ULT"] > 0)]
    datas = pd.to_datetime(df["DRF"], unit="D", origin="1899-12-30").dt.normalize()
    return pd.Series(df["ULT"].to_numpy(), index=datas).groupby(level=0).last().sort_index()


def consultar_bch(cods, inicio, fim, timeout=120.0):
    """Um BCH por código (inicio..fim) num Excel oculto e separado.
    Devolve {codigo: Series | 'mensagem de erro'}."""
    import pythoncom
    import win32com.client

    if not os.path.exists(cp.ADDIN_XLL):
        raise RuntimeError(f"add-in do Broadcast não encontrado em {cp.ADDIN_XLL!r}")
    if not cp.terminal_no_ar():
        raise RuntimeError("terminal Broadcast+ fechado: abra e logue o Broadcast antes de rodar")

    pythoncom.CoInitialize()
    xl = None
    try:
        xl = win32com.client.DispatchEx("Excel.Application")  # instância própria
        cp._com(lambda: setattr(xl, "Visible", False))
        cp._com(lambda: setattr(xl, "DisplayAlerts", False))
        if not cp._com(lambda: xl.RegisterXLL(cp.ADDIN_XLL)):
            raise RuntimeError("Excel não carregou o add-in do Broadcast")
        wb = cp._com(lambda: xl.Workbooks.Add())
        atributo = cp._atributo_formula(wb)
        formulas = {c: _formula(c, inicio, fim) for c in cods}

        # uma aba por código: o BCH "derrama" linhas e colunas a partir de A1
        abas = {}
        for i, c in enumerate(cods):
            abas[c] = cp._com(lambda: wb.Worksheets(1) if i == 0
                              else wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count)))
            cp._com(lambda: setattr(abas[c].Range("A1"), atributo, formulas[c]))

        # o add-in conecta no terminal de forma assíncrona: reescreve a fórmula enquanto a célula
        # mostrar #BCONN/#N/A (sem mexer em quem já está buscando)
        topo = {}
        t0 = ultimo_reenvio = time.time()
        while time.time() - t0 < timeout:
            pythoncom.PumpWaitingMessages()
            time.sleep(1.0)
            topo = {c: str(cp._com(lambda: abas[c].Range("A1").Text)).strip().upper() for c in cods}
            pendentes = [c for c in cods if topo[c] in cp.PENDENTES]
            if not pendentes:
                break
            if time.time() - ultimo_reenvio > 6:
                for c in pendentes:
                    if topo[c] in cp.REENVIAR:
                        cp._com(lambda: setattr(abas[c].Range("A1"), atributo, formulas[c]))
                ultimo_reenvio = time.time()

        res = {}
        for c in cods:
            t = topo.get(c, "")
            if t == "DRF":
                res[c] = _ler_serie(abas[c])
            elif t in cp.PENDENTES:
                res[c] = f"sem resposta em {int(time.time() - t0)}s ({t or 'vazio'})"
            else:
                res[c] = f"{t}: {cp.ERROS_BC.get(t, 'erro desconhecido do add-in')}"
        cp._com(lambda: wb.Close(False), tentativas=20)
        return res
    finally:
        if xl is not None:
            try:
                xl.Quit()
            except Exception:  # noqa
                pass
        pythoncom.CoUninitialize()


# ── Cálculo ──────────────────────────────────────────────────────────────────
def proximo_dia_util(d):
    while not cp.dia_util(d):
        d += timedelta(days=1)
    return d


def meses_cupom(venc):
    return sorted({venc.month, (venc.month + 5) % 12 + 1})


def datas_pagamento(venc, inicio, fim):
    """Pagamentos de cupom em (inicio, fim] antes do vencimento: dia 15 dos meses de cupom ou o dia
    útil seguinte. Quem entra no dia de um pagamento compra o título ex-cupom, então o 1º dia não recebe."""
    datas = []
    for ano in range(inicio.year, fim.year + 1):
        for mes in meses_cupom(venc):
            d = proximo_dia_util(date(ano, mes, 15))
            if inicio < d <= fim and d < venc:
                datas.append(d)
    return datas


@functools.lru_cache(maxsize=2)
def _feriados(com_20nov):
    """Feriados ANBIMA para contar os dias úteis até os fluxos (sem o 20/11 antes de INICIO_20NOV)."""
    return np.array(sorted(d for a in range(2000, 2080) for d in cp.feriados_nacionais(a)
                           if com_20nov or (d.month, d.day) != (11, 20)), dtype="datetime64[D]")


def pu_anbima(venc, d, taxa, vna):
    """PU ANBIMA da NTN-B no dia d: VNA x cotação, cotação = fluxos (cupom de 2,956301% e 100 no
    vencimento) descontados pela taxa indicativa em du/252, truncada na 4ª casa; PU truncado na 6ª."""
    fluxos = np.array([date(a, m, 15) for a in range(d.year, venc.year + 1) for m in meses_cupom(venc)
                       if d < date(a, m, 15) <= venc], dtype="datetime64[D]")
    if not len(fluxos):
        return float("nan")
    du = np.busday_count(np.datetime64(d), fluxos, holidays=_feriados(d >= INICIO_20NOV))
    desconto = (1 + taxa / 100) ** (du / 252)
    cotacao = math.floor(((CUPOM_ANBIMA / desconto).sum() + 100 / desconto[-1]) * 1e4) / 1e4
    return math.floor(vna * cotacao / 100 * 1e6) / 1e6


def montar_tabela(df, venc):
    """df com pu e vna por dia -> + cupom e as 3 cotas, que começam em 1 no 1º dia."""
    inicio, fim = df.index[0].date(), df.index[-1].date()

    df["cupom"] = 0.0
    for d in map(pd.Timestamp, datas_pagamento(venc, inicio, fim)):
        if d not in df.index:
            raise RuntimeError(f"sem preço no dia do cupom {d:%d/%m/%Y}: a série do Broadcast tem buraco")
        if pd.isna(df.at[d, "vna"]):
            raise RuntimeError(f"sem VNA ({CODIGO_VNA}) no dia do cupom {d:%d/%m/%Y}")
        df.at[d, "cupom"] = df.at[d, "vna"] * FATOR_CUPOM

    fator = (df["pu"] + df["cupom"]) / df["pu"].shift()
    fator.iloc[0] = 1.0
    df["cota_reinvestindo"] = fator.cumprod()
    df["cota_sem_reinvestir"] = (df["pu"] + df["cupom"].cumsum()) / df["pu"].iloc[0]
    df["cota_preco"] = df["pu"] / df["pu"].iloc[0]
    return df


def dias_faltando(datas):
    """Dias úteis (calendário ANBIMA) entre o 1º e o último dia da série que não vieram do Broadcast."""
    d, fim, faltam = datas[0].date(), datas[-1].date(), []
    presentes = {x.date() for x in datas}
    while d <= fim:
        if cp.dia_util(d) and d not in presentes:
            faltam.append(d)
        d += timedelta(days=1)
    return faltam


# (chave, nível, texto) das notas de calcular(), na ordem em que saem no console e na aba "Avisos"
NOTAS = [("emissao", "INFO", "série começa no 1º preço (título emitido depois da data_inicio)"),
         ("diverge", "AVISO", "PU do Broadcast trocado/defasado: usado o PU ANBIMA pela taxa"),
         ("so_taxa", "INFO", "sem PU no Broadcast: usado o PU ANBIMA pela taxa"),
         ("so_pu", "AVISO", "sem taxa no Broadcast: PU do Broadcast sem conferência"),
         ("nao_util", "AVISO", "preço em dia não útil descartado"),
         ("falta", "AVISO", "dia útil sem preço no Broadcast (o retorno do dia seguinte cobre o intervalo)"),
         ("vencido", "INFO", "título vencido: a série vai até o último preço (o resgate não entra)")]


def calcular(codigo, t, pu_bc, vna, taxa, fim):
    """Série de um título desde a data_inicio, escalada para a reinvestindo valer t["cota"] em
    t["data_cota"] (ou no 1º dia). pu_bc/taxa: Series ou a mensagem de erro do BCH.
    Devolve (df, {chave de NOTAS: [datas]})."""
    if isinstance(taxa, str):
        print(f"[AVISO] {nome(t['venc'])} ({codigo}): taxa indisponível ({taxa}): PU só do Broadcast, sem conferência")
        taxa = _serie_vazia()
    if isinstance(pu_bc, str):
        if taxa.empty:
            raise RuntimeError(f"PU {codigos(t['venc'])['pu']}: {pu_bc}")
        print(f"[AVISO] {nome(t['venc'])} ({codigo}): PU do Broadcast indisponível ({pu_bc}): só o PU pela taxa")
        pu_bc = _serie_vazia()
    inicio = pd.Timestamp(t["inicio"])
    pu_bc, taxa = pu_bc.loc[inicio:], taxa.loc[inicio:]
    dias = pu_bc.index.union(taxa.index)
    # o Broadcast às vezes traz preço em dia sem ANBIMA (sábado 12/07/2025 com o PU da segunda;
    # NTNB 2055 no domingo 28/07/2024 com o PU = VNA)
    util = np.array([cp.dia_util(d.date()) for d in dias], dtype=bool)
    notas = {"nao_util": list(dias[~util].date)}
    df = pd.DataFrame({"pu_broadcast": pu_bc, "taxa": taxa}, index=dias[util]).join(vna.rename("vna"), how="left")

    calc = pd.Series([pu_anbima(t["venc"], d.date(), tx, v) if pd.notna(tx) and pd.notna(v) else np.nan
                      for d, tx, v in zip(df.index, df["taxa"], df["vna"])], index=df.index, dtype=float)
    tem_bc, tem_calc = df["pu_broadcast"].notna(), calc.notna()
    notas["diverge"] = list(df.index[tem_bc & tem_calc & ((df["pu_broadcast"] / calc - 1).abs() > TOLERANCIA_PU)].date)
    notas["so_taxa"] = list(df.index[~tem_bc & tem_calc].date)
    notas["so_pu"] = list(df.index[tem_bc & ~tem_calc].date)
    df["pu"] = calc.fillna(df["pu_broadcast"])
    df = df.dropna(subset=["pu"])
    if df.empty:
        raise RuntimeError(f"sem preço desde {t['inicio']:%d/%m/%Y}")
    if df.index[0] != inicio:
        notas["emissao"] = [df.index[0].date()]
    df = montar_tabela(df, t["venc"])

    ancora = pd.Timestamp(t["data_cota"] or df.index[0])
    if ancora not in df.index:
        raise RuntimeError(f"sem preço na data_cota {ancora:%d/%m/%Y}: escolha um dia com preço")
    df[["cota_reinvestindo", "cota_sem_reinvestir", "cota_preco"]] *= t["cota"] / df.at[ancora, "cota_reinvestindo"]

    notas["falta"] = dias_faltando(df.index)
    if t["venc"] <= fim:
        notas["vencido"] = [t["venc"]]
    return df, notas


def imprimir_notas(notas):
    """notas: {chave: {data: [títulos]}} -> uma linha por tipo."""
    for chave, nivel, texto in NOTAS:
        por_dia = sorted(notas.get(chave, {}).items())
        if por_dia:
            itens = [f"{d:%d/%m/%Y} ({', '.join(n) if len(n) <= 3 else f'{len(n)} títulos'})" for d, n in por_dia]
            print(f"[{nivel}] {texto} — {len(itens)} dia(s): {', '.join(itens[:12])}"
                  f"{f' ... (lista completa na aba Avisos)' if len(itens) > 12 else ''}")


# ── Saída ────────────────────────────────────────────────────────────────────
COLUNAS = [("taxa", "Taxa indicativa (% a.a.)", "0.0000"),
           ("pu", "PU (R$)", "#,##0.000000"),
           ("vna", "VNA (R$)", "#,##0.000000"),
           ("cupom", "Cupom pago (R$)", "#,##0.000000"),
           ("cota_reinvestindo", "Cota reinvestindo o cupom", "0.00000000000000"),
           ("cota_sem_reinvestir", "Cota sem reinvestir (cupom em caixa)", "0.00000000000000"),
           ("cota_preco", "Cota só preço (PU)", "0.00000000000000"),
           ("pu_broadcast", "PU do Broadcast (R$)", "#,##0.000000")]
COLUNAS_DATA = {"Data", "Pagamento", "Vencimento", "Início", "Data da cota"}


def salvar_xlsx(series, titulos, notas, caminho):
    """{codigo: (venc, df)} -> 1 aba por NTN-B + "Cupons" (todas) + "Fonte" + "Avisos" (as notas de calcular)."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    cupons, fonte = [], []
    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        for codigo, (venc, df) in series.items():
            cotas = df[[c for c, _, _ in COLUNAS]].rename(columns={c: n for c, n, _ in COLUNAS})
            cotas.index.name = "Data"
            cotas.reset_index().to_excel(xw, sheet_name=nome(venc), index=False)
            c = df.loc[df["cupom"] > 0, ["vna", "cupom", "pu"]].rename(
                columns={"vna": "VNA (R$)", "cupom": "Cupom (R$)", "pu": "PU ex-cupom (R$)"})
            c.insert(0, "Título", nome(venc))
            cupons.append(c)
            cods = codigos(venc)
            fonte.append({"Título": nome(venc), "codigo_IAM_ativo": codigo, "Vencimento": venc,
                          "PU": cods["pu"], "Taxa": cods["taxa"], "VNA": cods["vna"],
                          "Início": df.index[0].date(), "Cota no início": df["cota_reinvestindo"].iloc[0],
                          "Cota informada": titulos[codigo]["cota"],
                          "Data da cota": titulos[codigo]["data_cota"] or df.index[0].date()})
        cupons = pd.concat(cupons).sort_index()
        cupons.index.name = "Pagamento"
        cupons.reset_index().to_excel(xw, sheet_name="Cupons", index=False)
        pd.DataFrame(fonte).to_excel(xw, sheet_name="Fonte", index=False)
        pd.DataFrame([{"Data": d, "Aviso": texto, "Títulos": ", ".join(n)}
                      for chave, _, texto in NOTAS for d, n in sorted(notas.get(chave, {}).items())],
                     columns=["Data", "Aviso", "Títulos"]).to_excel(xw, sheet_name="Avisos", index=False)

        formatos ={n: f for _, n, f in COLUNAS} | {"VNA (R$)": "#,##0.000000", "Cupom (R$)": "#,##0.000000",
                                                     "PU ex-cupom (R$)": "#,##0.000000",
                                                     "Cota no início": "0.00000000000000",
                                                     "Cota informada": "0.00000000000000"}
        for ws in xw.book.worksheets:
            ws.freeze_panes = "B2"
            for j, cel in enumerate(ws[1], start=1):
                cel.font = Font(bold=True)
                ws.column_dimensions[get_column_letter(j)].width = max(14, len(str(cel.value)) + 2)
                fmt = "dd/mm/yyyy" if cel.value in COLUNAS_DATA else formatos.get(cel.value)
                if fmt:
                    for (c,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        c.number_format = fmt
        ws = xw.book["Fonte"]
        ws.append([])
        ws.append(["Cupom", "VNA x ((1,06)^(1/2) - 1) no dia 15 de pagamento ou dia útil seguinte"])
        ws.append(["Fonte", "PU, VNA e taxa ANBIMA via Broadcast+ (BCH)"])
        ws.append(["Gerado em", f"{datetime.now():%d/%m/%Y %H:%M}"])


def tabela_sql(series):
    """{codigo: (venc, df)} -> linhas no layout da dbo.ntnb_cotas."""
    partes = []
    for codigo, (venc, df) in series.items():
        p = df.rename(columns={"taxa": "taxa_indicativa", "cota_reinvestindo": "cota_acumulada"})
        p["codigo_IAM_ativo"] = codigo
        p["date"] = df.index.date
        p["vencimento"] = venc
        partes.append(p.reset_index(drop=True))
    return pd.concat(partes, ignore_index=True)[list(COLUNAS_SQL)]


def ler_csv(caminho):
    tab = pd.read_csv(caminho, sep=";", dtype={"codigo_IAM_ativo": str})
    faltam = [c for c in COLUNAS_SQL if c not in tab.columns]
    if faltam:
        raise ValueError(f"{caminho}: faltam as colunas {faltam} (não é um ntnb_cotas_*.csv?)")
    for c in ("date", "vencimento"):
        tab[c] = pd.to_datetime(tab[c], format="%Y-%m-%d").dt.date
    return tab[list(COLUNAS_SQL)]


def _valor_sql(v, casas):
    if pd.isna(v):
        return None
    if casas is None:
        return v
    return Decimal(repr(float(v))).quantize(Decimal(1).scaleb(-casas))


def gravar_sql(tab):
    """Substitui no SQL a série inteira de cada título de tab (DELETE + INSERT na mesma transação)."""
    repetidas = tab.duplicated(["codigo_IAM_ativo", "date"], keep=False)
    if repetidas.any():
        raise ValueError(f"linhas repetidas na chave (codigo_IAM_ativo, date):\n{tab[repetidas]}")
    obrigatorias = [c for c in COLUNAS_SQL if c not in OPCIONAIS_SQL]
    vazias = tab[obrigatorias].isna().any()
    if vazias.any():
        raise ValueError(f"valores vazios em colunas NOT NULL: {list(vazias[vazias].index)}")
    linhas = [tuple(_valor_sql(v, casas) for v, casas in zip(linha, COLUNAS_SQL.values()))
              for linha in tab[list(COLUNAS_SQL)].itertuples(index=False, name=None)]
    titulos = sorted(tab["codigo_IAM_ativo"].unique())

    sys.path.insert(0, SQL_CONNECTION_DIR)
    from sql_connection import disconnect, get_connection

    sql = get_connection()
    try:
        conn = sql.connection
        if conn is None:
            raise RuntimeError(f"sem conexão com o banco {sql.db_name}")
        colunas = ", ".join(f"[{c}]" for c in COLUNAS_SQL)
        marcas = ", ".join("?" for _ in COLUNAS_SQL)
        try:
            cursor = conn.cursor()
            cursor.fast_executemany = True
            cursor.executemany(f"DELETE FROM {TABELA_SQL} WHERE codigo_IAM_ativo = ?", [(c,) for c in titulos])
            cursor.executemany(f"INSERT INTO {TABELA_SQL} ({colunas}) VALUES ({marcas})", linhas)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
    finally:
        disconnect()

    print(f"[OK] {len(linhas)} linhas gravadas em {sql.db_name}.{TABELA_SQL}:")
    for codigo, g in tab.groupby("codigo_IAM_ativo"):
        print(f"  {codigo} (NTN-B {g['vencimento'].iloc[0]:%d/%m/%Y}): "
              f"{g['date'].min():%d/%m/%Y} a {g['date'].max():%d/%m/%Y} ({len(g)} dias)")


def _gravar_sql_ou_sair(tab):
    try:
        gravar_sql(tab)
    except Exception as e:  # noqa — erro do banco (tabela inexistente, permissão, credenciais)
        sys.exit(f"[ERRO] gravação no SQL falhou: {e}\n"
                 f"       (a tabela existe? input/create_ntnb_cotas.sql; credenciais SQL_LOGIN/SQL_PASSWORD)")


def _data(s):
    return datetime.strptime(s, "%d/%m/%Y").date()


def main(argv=None):
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa
            pass

    hoje = datetime.now().date()
    ap = argparse.ArgumentParser(description="Cota de NTN-B com e sem reinvestimento do cupom (ANBIMA via "
                                             "Broadcast+) e carga na DBSX337.dbo.ntnb_cotas.")
    ap.add_argument("titulos", nargs="*", help="ISINs da lista (padrão: todos)")
    ap.add_argument("--lista", default=LISTA_PADRAO,
                    help="CSV com as NTN-B: codigo_IAM_ativo;vencimento;data_inicio;cota_inicio "
                         "(padrão: input/titulos.csv)")
    ap.add_argument("--fim", type=_data, default=hoje,
                    help="último dia dd/mm/aaaa (padrão: hoje; vem até o último dia com PU)")
    ap.add_argument("--timeout", type=float, default=120.0, help="espera máxima (s) pela resposta do Broadcast (padrão 120)")
    ap.add_argument("--sql", action="store_true",
                    help=f"grava também no SQL ({TABELA_SQL}): precisa de acesso ao banco e SQL_LOGIN/SQL_PASSWORD")
    ap.add_argument("--carregar", nargs="?", const="", metavar="CSV",
                    help="só grava no SQL um ntnb_cotas_*.csv gerado antes (sem arquivo: o mais recente de "
                         "output/); não usa o Broadcast")
    args = ap.parse_args(argv)

    if args.carregar is not None:
        caminho = args.carregar or max(glob.glob(os.path.join(SAIDA_DIR, "ntnb_cotas_*.csv")), default=None)
        if not caminho:
            ap.error(f"nenhum ntnb_cotas_*.csv em {SAIDA_DIR}")
        print(f"[INFO] carregando {caminho}")
        try:
            tab = ler_csv(caminho)
        except (OSError, ValueError) as e:
            ap.error(str(e))
        _gravar_sql_ou_sair(tab)
        return

    if args.fim > hoje:
        ap.error("--fim no futuro")
    try:
        titulos = carregar_lista(args.lista)
    except (OSError, ValueError) as e:
        ap.error(str(e))
    if args.titulos:
        pedidos = [c.upper() for c in args.titulos]
        fora = [c for c in pedidos if c not in titulos]
        if fora:
            ap.error(f"fora da lista {args.lista}: {', '.join(fora)}")
        titulos = {c: titulos[c] for c in pedidos}

    inicio = min(t["inicio"] for t in titulos.values())
    cods = {c: codigos(t["venc"]) for c, t in titulos.items()}
    consulta = [CODIGO_VNA] + [cods[c][k] for c in titulos for k in ("pu", "taxa")]
    print(f"[INFO] {len(titulos)} NTN-B ({', '.join(nome(t['venc']) for t in titulos.values())}) | "
          f"{inicio:%d/%m/%Y} a {args.fim:%d/%m/%Y} | {len(consulta)} séries do Broadcast")
    try:
        res = consultar_bch(consulta, inicio, args.fim, args.timeout)
    except RuntimeError as e:
        sys.exit(f"[ERRO] {e}")

    vna = res[CODIGO_VNA]
    if isinstance(vna, str) or vna.empty:
        sys.exit(f"[ERRO] {CODIGO_VNA} (VNA): {vna if isinstance(vna, str) else 'sem dados no período'}")
    series, notas = {}, {}  # notas: {chave: {data: [títulos]}} — o mesmo dia costuma valer para vários títulos
    for c, t in titulos.items():
        try:
            df, n = calcular(c, t, res[cods[c]["pu"]], vna, res[cods[c]["taxa"]], args.fim)
        except RuntimeError as e:
            print(f"[ERRO] {nome(t['venc'])} ({c}): {e} — fica de fora")
            continue
        series[c] = (t["venc"], df)
        for chave, dias in n.items():
            for d in dias:
                notas.setdefault(chave, {}).setdefault(d, []).append(nome(t["venc"]))
    if not series:
        sys.exit("[ERRO] nenhuma NTN-B calculada")
    imprimir_notas(notas)

    for c, (venc, df) in series.items():
        u = df.iloc[-1]
        print(f"[OK] {nome(venc)} ({c}): {len(df)} dias, {int((df['cupom'] > 0).sum())} cupons | "
              f"{df.index[-1]:%d/%m/%Y}: reinvestindo {u['cota_reinvestindo']:.10f} | "
              f"sem reinvestir {u['cota_sem_reinvestir']:.10f} | só preço {u['cota_preco']:.10f}")

    fim = max(df.index[-1] for _, df in series.values())
    os.makedirs(SAIDA_DIR, exist_ok=True)
    xlsx = os.path.join(SAIDA_DIR, f"ntnb_{fim:%Y%m%d}.xlsx")
    salvar_xlsx(series, titulos, notas, xlsx)
    print(f"[OK] {xlsx}")
    tab = tabela_sql(series)
    csv = os.path.join(SAIDA_DIR, f"ntnb_cotas_{fim:%Y%m%d}.csv")
    tab.to_csv(csv, sep=";", index=False)
    print(f"[OK] {csv} ({len(tab)} linhas p/ o SQL)")

    if args.sql:
        _gravar_sql_ou_sair(tab)


if __name__ == "__main__":
    main()
