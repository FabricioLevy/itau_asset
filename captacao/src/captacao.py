# ─────────────────────────────────────────────
# captacao.py
# Foto diária de fundos de investimento (cota, captação, resgate, cotistas, PL) no
# layout do comdinheiro — dados do Broadcast+ (AE), completados pela CVM.
#
# Saída (../output/captacao/captacao_<yyyymmdd>.xlsx):
#   aba "Captacao": mesmo layout do export do comdinheiro — título na linha 1 (as fontes
#     usadas), cabeçalho na linha 2 (com a data da análise nos nomes das colunas), 1 linha
#     por fundo:
#     Data da Análise | Código Anbima da classe da cota do fundo | CNPJ do Fundo |
#     Razão Social do Fundo | Valor da Cota | Captação no Dia R$ | Resgate no Dia R$ |
#     Número Cotistas | Patrimônio R$
#   aba "Fontes": trilha de auditoria — valor do Broadcast e da CVM lado a lado, de onde
#     saiu cada campo, dados provisórios, checagem do PL e divergências entre as fontes.
#
# Fontes:
#   1) Broadcast+ — série histórica via add-in Excel (função BCH), que conversa com o
#      terminal (bcsys32.exe) em 127.0.0.1:8100 (aesp.cfg). O script abre um Excel OCULTO
#      e separado (DispatchEx — não mexe em Excel aberto), carrega o add-in e escreve um
#      =BCH(...) por fundo, de D-10 a D. Campos: DRF data | ULT cota | CAPT captação |
#      RESG resgate | PATR patrimônio | QCOT cotistas.
#      Código do fundo no Broadcast = "FI" + código ANBIMA da classe, só a parte numérica
#      (C0001390643 -> FI1390643; é assim que o AE.Funds.Plugin monta o ticker). Para
#      achar: busca de Fundos do terminal pelo CNPJ -> coluna "Código Anbima".
#   2) CVM (dados abertos) — Informe Diário (inf_diario_fi_AAAAMM.zip; até 2020 em
#      HIST/inf_diario_fi_AAAA.zip) e cadastro (registro_fundo_classe.zip: razão social —
#      o comdinheiro usa a mesma denominação — e tipo de classe). Cache em ../input/cvm/.
#      O informe diário só cobre classes FIF: FIDC, FIP, FII, FIAGRO, FIIM (ETF) etc. não
#      entregam, então para elas só há o Broadcast.
#
# Consistência (aplicada igualmente à série do Broadcast e à da CVM, usando o dia útil
#   anterior disponível, D-1):
#   - Identidade do PL: num dia normal PL(D) - PL(D-1) x cota(D)/cota(D-1) = captação -
#     resgate, com diferença de poucos reais (coluna "Checagem PL" da aba Fontes).
#   - PRÉVIA: cotistas = 0 com PL > 0 — enquanto a ANBIMA não fecha o dia, o Broadcast
#     traz a cota certa, mas captação/resgate/cotistas zerados e o PL de D-1.
#   - LINHA REPETIDA: D idêntica a D-1 (cota, PL, captação, resgate, cotistas), num dia útil,
#     com fluxo no dia ou num fundo cuja cota muda todo dia (>= 80% dos dias da janela) —
#     o administrador reenviou D-1. Marca TODOS os campos, inclusive a cota.
#   - PL PROVISÓRIO: PL = PL(D-1) corrigido pela cota (|dif| <= TOL_PL_ROLADO) apesar de
#     fluxo líquido relevante (> FLUXO_RELEVANTE), ou PL idêntico ao de D-1 com a cota
#     mexendo — o PL ainda não incorporou o dia.
#     Não conta AMORTIZAÇÃO lançada como resgate: a cota cai no valor amortizado (PL(D-1) x
#     (1 + rendimento típico da cota - razão), 10% de tolerância) e o PL só rola — num
#     resgate comum o PL cai de novo pelo valor resgatado. Vai só uma nota em Observações.
#   - POSSÍVEL PRÉVIA: captação = resgate = 0, cotistas iguais aos de D-1 e PL rolado/
#     repetido, num fundo que tem fluxo quase todo dia (>= 80% dos dias anteriores da
#     janela) e fluxo típico > FLUXO_TIPICO_RELEVANTE — placeholder do administrador.
#     Fundo pequeno com dia sem movimento de verdade não é marcado.
#   - Se D-1 também estava marcado, a checagem do PL de D não é conclusiva (nota).
#   - Classe que só informa por subclasse à CVM: a soma não tem cota, então cada subclasse
#     é checada; subclasse marcada, em prévia, ou que informava na janela e faltou no dia
#     (e não está encerrada, com PL 0) deixa a soma provisória.
#   Limiares: TOL_PL_ROLADO = max(R$ 1; 1e-7 x PL); FLUXO_RELEVANTE = max(R$ 1.000;
#   1e-5 x PL); FLUXO_TIPICO_RELEVANTE = max(R$ 1.000; 0,1% x PL).
#
# Regras de combinação (por campo):
#   - Broadcast prevalece; campo que o Broadcast não tem -> CVM; sem nenhuma das duas ->
#     em branco (como no comdinheiro).
#   - Campo provisório no Broadcast (captação/resgate/cotistas/PL numa prévia ou possível
#     prévia; todos numa linha repetida; PL num PL provisório) -> usa a CVM se ela tiver o
#     dia e não estiver provisória nesse campo.
#   - Os zeros de uma prévia (captação/resgate/cotistas), de qualquer fonte, nunca vão para
#     a planilha: sem outra fonte, ficam em branco.
#   - Valor provisório que ficar na planilha sai com "(provisório)" na coluna Fonte da aba
#     Fontes, e o motivo na coluna Observações.
#   - Código de SUBCLASSE ANBIMA (S...): o ticker "FI"+número vale p/ a classe, então o
#     Broadcast não é consultado; a CVM só é usada se a classe não tiver várias subclasses.
#
# Nome do arquivo: captacao_<yyyymmdd>.xlsx p/ a lista; captacao_<yyyymmdd>_selecao.xlsx
#   quando os fundos vêm da linha de comando. Se já existir a planilha do dia com MAIS fundos
#   vindos do Broadcast que a rodada atual (terminal fechado, --sem-broadcast, falha), a
#   anterior é preservada e a nova vai p/ _so_cvm / _parcial_<hhmmss>.
#
# Lista de fundos (--lista, padrão ../input/fundos_captacao.csv): CSV (; , ou tab, UTF-8,
#   cp1252 ou "Texto Unicode" do Excel) ou XLSX com colunas de CNPJ, código ANBIMA e
#   (opcional) nome — reconhecidas pelo nome ("CNPJ do Fundo"/"cnpj", "Código Anbima..."/
#   "codigo_anbima", "Razão Social..."/"nome"); linhas de título antes do cabeçalho são
#   puladas, então o próprio export do comdinheiro serve de lista. Linhas do mesmo fundo
#   (mesmo CNPJ ou código) são juntadas. CNPJ que o Excel transformou em número (sem os
#   zeros à esquerda) é recuperado se o dígito verificador conferir. Em CSV, linha que
#   começa com "# " é comentário.
#   Na linha de comando também valem CNPJs e códigos ANBIMA (C0000142697, 142697 ou
#   FI142697); o que estiver na lista completa o que faltar (CNPJ <-> código).
#
# Requisitos: p/ o Broadcast, terminal aberto e logado + Excel 64-bit com o add-in; sem
#   terminal, sai só com dados da CVM (com aviso). pywin32, pandas, openpyxl (anaconda).
#
# Uso:
#   python captacao.py                          (lista padrão, dia útil anterior a hoje)
#   python captacao.py --data 18/09/2026
#   python captacao.py C0000142697 00.807.777/0001-62 --data 18/09/2026
#   python captacao.py --lista "C:\...\export_comdinheiro.xlsx"
#   python captacao.py --dias 20                (últimos 20 dias úteis até o dia útil anterior)
#   python captacao.py --inicio 01/09/2026 --data 18/09/2026
#   Período: 1 arquivo (captacao_<inicio>_a_<fim>.xlsx) com 1 linha por fundo e dia útil; as
#   colunas não levam a data (ela está em "Data da Análise"); dia sem dado fica em branco.
# ─────────────────────────────────────────────

import argparse
import codecs
import csv
import math
import os
import re
import shutil
import socket
import statistics
import sys
import time
import unicodedata
import urllib.error
import urllib.request
import zipfile
from datetime import date, datetime, timedelta

import pandas as pd


AE_BROADCAST_DIR = r"C:\Program Files (x86)\AE Broadcast"
ADDIN_XLL = os.path.join(AE_BROADCAST_DIR, "Broadcast.AddIn64.xll")

AQUI = os.path.dirname(os.path.abspath(__file__))
LISTA_PADRAO = os.path.join(AQUI, "..", "input", "fundos_captacao.csv")
CACHE_CVM = os.path.join(AQUI, "..", "input", "cvm")
SAIDA_DIR = os.path.join(AQUI, "..", "output", "captacao")

CVM_BASE = "https://dados.cvm.gov.br/dados/FI"
CVM_INF_MENSAL = CVM_BASE + "/DOC/INF_DIARIO/DADOS/inf_diario_fi_{:%Y%m}.zip"
CVM_INF_ANUAL = CVM_BASE + "/DOC/INF_DIARIO/DADOS/HIST/inf_diario_fi_{:%Y}.zip"
CVM_CADASTRO = CVM_BASE + "/CAD/DADOS/registro_fundo_classe.zip"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
# tipos do informe diário pós-RCVM 175; na migração a CVM chegou a publicar a mesma classe
# também como 'FI' — nesse caso vale a linha nova
CVM_TIPOS_NOVOS = ("CLASSES - FIF", "CLASSE FIF/FAPI")

CAMPOS = ["cota", "captacao", "resgate", "cotistas", "patrimonio"]
# campo BCH -> campo interno (a ordem é a da fórmula)
BCH_CAMPOS = [("DRF", "data"), ("ULT", "cota"), ("CAPT", "captacao"), ("RESG", "resgate"),
              ("PATR", "patrimonio"), ("QCOT", "cotistas")]
# coluna do Informe Diário da CVM -> campo interno
CVM_CAMPOS = {"VL_QUOTA": "cota", "CAPTC_DIA": "captacao", "RESG_DIA": "resgate",
              "NR_COTST": "cotistas", "VL_PATRIM_LIQ": "patrimonio"}
FLUXOS = ["captacao", "resgate", "cotistas", "patrimonio"]  # o que uma prévia deixa errado

LAYOUT = [("data", "Data da Análise"),
          ("codigo_anbima", "Código Anbima da classe da cota do fundo"),
          ("cnpj", "CNPJ do Fundo"),
          ("razao_social", "Razão Social do Fundo"),
          ("cota", "Valor da Cota {d}"),
          ("captacao", "Captação no Dia R$ {d}"),
          ("resgate", "Resgate no Dia R$ {d}"),
          ("cotistas", "Número Cotistas {d}"),
          ("patrimonio", "Patrimônio R$ {d}")]
NOMES_CAMPOS = {"cota": "Cota", "captacao": "Captação", "resgate": "Resgate",
                "cotistas": "Cotistas", "patrimonio": "PL"}

# códigos de erro do add-in (Broadcast.AddIn.chm, "Problemas Comuns")
ERROS_BC = {
    "#BCONN": "Excel não conectou ao terminal (terminal fechado ou deslogado)",
    "#BCPERM": "código inexistente ou sem permissão no Broadcast",
    "#BCSYMBOL": "código não cadastrado no Broadcast",
    "#BCPROXY": "erro na consulta histórica (parâmetros, período ou permissão)",
    "#BCDATE": "período anterior ao início da base histórica",
    "#BCSYNTAX": "erro de sintaxe na fórmula BCH (ex.: data futura)",
    "#BCLIMIT": "limite de consultas históricas excedido",
}
# textos que a célula mostra enquanto o add-in ainda está conectando/buscando
PENDENTES = {"", "#N/A", "#N/D", "#BCONN", "#GETTING_DATA"}
REENVIAR = {"", "#N/A", "#N/D", "#BCONN"}  # #GETTING_DATA = consulta em andamento: não reescreve
JANELA_BCH_DIAS = 10     # consulta D-10..D: acha D, D-1 (checagem do PL) e o último dia disponível
DESISTE_BCONN_S = 45     # tudo ainda #BCONN depois disso (ou de 40% do --timeout) = terminal deslogado
# códigos de erro do Excel (#N/A, #VALOR!...) chegam pelo COM como inteiros nessa faixa
ERRO_EXCEL_MIN, ERRO_EXCEL_MAX = -2146826300, -2146826200


def _tem(v):
    return v is not None and not (isinstance(v, float) and math.isnan(v))


def _sem_acento(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower().strip()


def _brl(v):
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ── CNPJ e código ANBIMA ─────────────────────────────────────────────────────
_CNPJ_MASCARA = re.compile(r"^[0-9A-Z]{2}\.[0-9A-Z]{3}\.[0-9A-Z]{3}/[0-9A-Z]{4}-\d{2}$")


def normalizar_cnpj(s):
    """'66.198.405/0001-94' -> '66198405000194' (maiúsculas, só alfanuméricos)."""
    return re.sub(r"[^0-9A-Z]", "", str(s or "").upper())


def cnpj_valido(cnpj):
    """Confere os 2 dígitos verificadores. Vale p/ o CNPJ alfanumérico (Receita, jul/2026):
    cada caractere entra com valor ord(c) - 48, mesmos pesos do módulo 11."""
    c = normalizar_cnpj(cnpj)
    if len(c) != 14 or not c[12:].isdigit():
        return False
    vals = [ord(ch) - 48 for ch in c[:12]]
    for pesos in ([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2], [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]):
        r = sum(v * p for v, p in zip(vals, pesos)) % 11
        vals.append(0 if r < 2 else 11 - r)
    return "".join(map(str, vals[12:])) == c[12:]


def _limpar_celula(s):
    """Desfaz o que o Excel costuma fazer com identificadores numéricos: '142697.0' ->
    '142697'; '1.7024000153E+10' -> '17024000153'. Notação científica que já perdeu
    dígitos (ex.: '3.06278E+11', formato Geral do Excel) fica como está — não dá p/ recuperar."""
    s = str(s if s is not None else "").strip()
    if re.fullmatch(r"\d+\.0+", s):
        return s.split(".")[0]
    m = re.fullmatch(r"(\d+)(?:\.(\d+))?[eE]\+?(\d+)", s)
    if m:
        digitos = len((m.group(1) + (m.group(2) or "")).lstrip("0"))
        inteiro = str(int(float(s))) if len(m.group(3)) < 3 else ""
        # só confia se a mantissa trouxer os dígitos todos (CNPJ sem zeros à esquerda tem >= 11)
        if inteiro and digitos >= min(len(inteiro), 11):
            return inteiro
    return s


def ler_cnpj(s):
    """CNPJ normalizado (14 caracteres) ou None. Aceita: com máscara (numérico ou
    alfanumérico), 14 dígitos, alfanumérico sem máscara (se o DV conferir) e CNPJ que
    virou número no Excel, sem os zeros à esquerda (se o DV conferir)."""
    s = _limpar_celula(s).upper()
    if _CNPJ_MASCARA.match(s) or re.fullmatch(r"\d{14}", s):
        return normalizar_cnpj(s)
    # sem máscara: só com DV conferindo e filial != 0000 (não existe; aparece quando o
    # Excel truncou o número e completou com zeros)
    if re.fullmatch(r"[0-9A-Z]{12}\d{2}", s) and cnpj_valido(s) and s[8:12] != "0000":
        return s
    if re.fullmatch(r"\d{11,13}", s) and cnpj_valido(s.zfill(14)) and s.zfill(14)[8:12] != "0000":
        return s.zfill(14)
    return None


def formatar_cnpj(cnpj):
    c = normalizar_cnpj(cnpj)
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if len(c) == 14 else (cnpj or "")


def ler_codigo_anbima(s):
    """'C0000142697' / 'S0000000574' / '142697' / 'FI142697' -> ('C', 142697); senão None."""
    m = re.fullmatch(r"(?:FI)?([CS])?(\d{1,10})", _limpar_celula(s).upper())
    if not m or int(m.group(2)) == 0:
        return None
    return (m.group(1) or "C", int(m.group(2)))


def codigo_anbima(prefixo, num):
    return f"{prefixo}{num:010d}" if num else ""


def codigo_broadcast(num):
    return f"FI{num}" if num else ""


# ── lista de fundos ───────────────────────────────────────────────────────────
def _ler_texto(caminho):
    """Texto do arquivo: UTF-16 ('Texto Unicode' do Excel), UTF-8 ou cp1252."""
    with open(caminho, "rb") as fh:
        bruto = fh.read()
    if bruto.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return bruto.decode("utf-16")
    amostra = bruto[:400]
    if amostra and amostra.count(b"\x00") > len(amostra) // 4:  # UTF-16 sem BOM: ~metade é NUL
        return bruto.decode("utf-16-le", errors="replace")
    try:
        return bruto.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    # mistura (ex.: lista UTF-8 com uma linha acrescentada em ANSI): decide linha a linha
    linhas = []
    for ln in bruto.split(b"\n"):
        try:
            linhas.append(ln.decode("utf-8"))
        except UnicodeDecodeError:
            linhas.append(ln.decode("cp1252", errors="replace"))
    return "\n".join(linhas).lstrip("﻿")


# critérios de reconhecimento das colunas (sobre o texto do cabeçalho sem acento e minúsculo)
_NAO_E_DO_FUNDO = ("admin", "gestor", "custod", "auditor", "control", "distrib", "master")
_NAO_E_CODIGO = ("classif", "tipo", "categoria", "nome", "razao", "cvm", "isin", "cnpj")
_COL_CNPJ = [lambda c: "cnpj" in c and ("fundo" in c or "classe" in c) and not any(o in c for o in _NAO_E_DO_FUNDO),
             lambda c: "cnpj" in c and not any(o in c for o in _NAO_E_DO_FUNDO)]
_COL_COD = [lambda c: re.search(r"\bcod(igo)?\b", c) and "anbima" in c,
            lambda c: "broadcast" in c,
            lambda c: re.match(r"cod(igo)?\b", c) and not any(o in c for o in _NAO_E_CODIGO),
            lambda c: "anbima" in c and not any(o in c for o in _NAO_E_CODIGO)]
_COL_NOME = [lambda c: "razao" in c, lambda c: "nome" in c]


def _achar_coluna(celulas, niveis, pular=()):
    """Índice da 1ª coluna que atende ao 1º nível de critério que tiver candidato."""
    for criterio in niveis:
        cands = [j for j, c in enumerate(celulas) if j not in pular and criterio(c)]
        if cands:
            return cands[0]
    return None


def _pontuar_cabecalho(celulas):
    """Quantas colunas (CNPJ, código, nome) a linha reconhece — 0 se nem CNPJ nem código.
    Célula longa ou começando com '#' (título, comentário) não conta como cabeçalho."""
    celulas = [c if len(c) <= 60 and not c.startswith("#") else "" for c in map(_sem_acento, celulas)]
    col_cnpj = _achar_coluna(celulas, _COL_CNPJ)
    col_cod = _achar_coluna(celulas, _COL_COD, pular=(col_cnpj,))
    if col_cnpj is None and col_cod is None:
        return 0
    col_nome = _achar_coluna(celulas, _COL_NOME, pular=(col_cnpj, col_cod))
    return sum(c is not None for c in (col_cnpj, col_cod, col_nome))


def _eh_comentario(celulas):
    """XLSX: linha de comentário = 1ª célula '# texto' (ou só '#' com o resto vazio). '#N/D',
    '#REF!' (erro de fórmula) e uma coluna de numeração chamada '#' são dados."""
    if not celulas:
        return False
    primeira = str(celulas[0])
    return bool(re.match(r"\s*#\s+\S", primeira)
                or (re.match(r"\s*#\s*$", primeira) and not any(str(c).strip() for c in celulas[1:])))


# texto: comentário = linha física que começa com "# " (cerquilha + espaço) ou só "#" — também
# entre aspas, como o Excel salva; "#\t..." (coluna '#' numa lista com tab) e "#N/D" são dados
_COMENTARIO_TXT = re.compile(r'\s*"?#( |"?\s*$)')
_PARECE_ID = re.compile(r"[0-9A-Z]{2}\.[0-9A-Z]{3}\.[0-9A-Z]{3}/[0-9A-Z]{4}-\d{2}|\bC\d{10}\b|\d{6,}")


def _linhas_do_arquivo(caminho):
    """[(nº da linha no arquivo, [células])] — XLSX ou texto (separador ; tab ou ,), já sem
    linhas vazias e comentários."""
    if os.path.splitext(caminho)[1].lower() in (".xlsx", ".xlsm", ".xls"):
        bruto = pd.read_excel(caminho, header=None, dtype=str).fillna("")
        linhas = [(i, [str(v) for v in lin]) for i, lin in enumerate(bruto.itertuples(index=False), start=1)]
        return [(i, lin) for i, lin in linhas if any(str(c).strip() for c in lin) and not _eh_comentario(lin)]

    texto = _ler_texto(caminho)
    fisicas = texto.splitlines()
    comentarios = {i for i, ln in enumerate(fisicas, start=1) if _COMENTARIO_TXT.match(ln)}
    # separador = o que melhor reconhece um cabeçalho nas 1as linhas não comentadas (empate: ; tab ,)
    amostra = [ln for i, ln in enumerate(fisicas, start=1) if ln.strip() and i not in comentarios][:12]
    sep = max(";\t,", key=lambda s: max([_pontuar_cabecalho(next(csv.reader([ln], delimiter=s)))
                                         for ln in amostra] or [0]))

    def engoliu_linhas(reg):
        # célula com várias linhas em que alguma linha interna tem o separador e cara de
        # identificador = uma aspa sem par engoliu outros fundos (texto com quebra de linha
        # legítima, mesmo citando um código, não tem o separador)
        return any(sep in interna and _PARECE_ID.search(interna)
                   for c in reg if "\n" in c for interna in c.split("\n")[1:])

    # arquivo inteiro de uma vez (célula com quebra de linha entre aspas); se uma aspa sem par
    # engolir linhas, cai p/ uma linha física por registro
    linhas = []
    leitor = csv.reader(texto.splitlines(keepends=True), delimiter=sep, strict=True)
    try:
        inicio = 1
        for reg in leitor:
            if leitor.line_num - inicio > 3 or engoliu_linhas(reg):
                raise csv.Error("registro atravessando linhas de outros fundos")
            linhas.append((inicio, reg))
            inicio = leitor.line_num + 1
    except csv.Error:
        linhas = [(i, next(csv.reader([ln], delimiter=sep))) for i, ln in enumerate(fisicas, start=1)]
    return [(i, lin) for i, lin in linhas if any(str(c).strip() for c in lin) and i not in comentarios]


def carregar_lista(caminho):
    """Lê a lista de fundos (CSV/XLSX) e junta linhas do mesmo fundo.
    Devolve [{cnpj, prefixo, num, nome, entrada, avisos}], na ordem do arquivo."""
    linhas = _linhas_do_arquivo(caminho)
    if not linhas:
        return []
    arquivo = os.path.basename(caminho)

    # cabeçalho = a linha (entre as 10 primeiras) que mais reconhece colunas — o export do
    # comdinheiro tem título antes; comentário/título citando "CNPJ" não ganha do cabeçalho
    pontos = [_pontuar_cabecalho(lin) for _, lin in linhas[:10]]
    if not pontos or max(pontos) == 0:
        raise ValueError(f"{caminho}: não achei o cabeçalho (coluna de CNPJ e/ou código ANBIMA)")
    cab = pontos.index(max(pontos))
    cab_bruto = linhas[cab][1]
    celulas = [_sem_acento(c) for c in cab_bruto]
    col_cnpj = _achar_coluna(celulas, _COL_CNPJ)
    col_cod = _achar_coluna(celulas, _COL_COD, pular=(col_cnpj,))
    col_nome = _achar_coluna(celulas, _COL_NOME, pular=(col_cnpj, col_cod))
    for nome_col, col, chave in (("CNPJ", col_cnpj, "cnpj"), ("código ANBIMA", col_cod, "anbima")):
        if col is not None and sum(chave in c for c in celulas) > 1:
            print(f"[AVISO] {arquivo}: mais de uma coluna de {nome_col}; usando {cab_bruto[col]!r}")
    if col_cod is None:
        print(f"[AVISO] {arquivo}: não achei coluna de código ANBIMA — o Broadcast só será consultado "
              f"para códigos passados na linha de comando")
    if col_cnpj is None:
        print(f"[AVISO] {arquivo}: não achei coluna de CNPJ do fundo — sem CVM e sem razão social")

    def celula(lin, col):
        return lin[col].strip() if col is not None and col < len(lin) else ""

    fundos = []
    for nlin, lin in linhas[cab + 1:]:
        cnpj_bruto, cod_bruto = celula(lin, col_cnpj), celula(lin, col_cod)
        if not cnpj_bruto and not cod_bruto:
            print(f"[AVISO] {arquivo}, linha {nlin}: colunas de CNPJ e código vazias "
                  f"({'; '.join(c for c in lin if str(c).strip())[:80]!r}) — ignorada.")
            continue
        f = dict(cnpj="", prefixo="C", num=None, nome=celula(lin, col_nome),
                 entrada=cnpj_bruto or cod_bruto, avisos=[])
        if cnpj_bruto:
            f["cnpj"] = ler_cnpj(cnpj_bruto) or ""
            if not f["cnpj"]:
                f["avisos"].append(f"CNPJ não reconhecido na lista (linha {nlin}): {cnpj_bruto!r}")
        if cod_bruto:
            cod = ler_codigo_anbima(cod_bruto)
            if cod:
                f["prefixo"], f["num"] = cod
            else:
                f["avisos"].append(f"código ANBIMA não reconhecido na lista (linha {nlin}): {cod_bruto!r}")
        if not f["cnpj"] and not f["num"]:
            print(f"[AVISO] {arquivo}, linha {nlin}: sem CNPJ nem código ANBIMA válidos "
                  f"({cnpj_bruto!r} / {cod_bruto!r}) — ignorada.")
            continue
        fundos.append(f)
    return _juntar(fundos)


def _juntar(fundos):
    """Junta registros do mesmo fundo, sem depender da ordem das linhas: pares CNPJ+código
    iguais viram um só; linha só com CNPJ (ou só com código) completa o par que tiver o
    mesmo CNPJ (código). CNPJ com dois códigos (ou o contrário) fica em dois, com aviso."""
    def cod(f):
        return (f["prefixo"], f["num"]) if f["num"] else None

    def absorver(dest, f):
        if not dest["nome"]:
            dest["nome"] = f["nome"]
        dest["avisos"] += [a for a in f["avisos"] if a not in dest["avisos"]]

    completos, ordem, pos = {}, [], {}  # pos: id(registro) -> 1ª linha da lista em que o fundo aparece
    for i, f in enumerate(fundos):
        if f["cnpj"] and f["num"]:
            chave = (f["cnpj"], cod(f))
            if chave in completos:
                absorver(completos[chave], f)
            else:
                completos[chave] = dict(f, avisos=list(f["avisos"]))
                ordem.append(completos[chave])
                pos[id(completos[chave])] = i
    por_cnpj, por_cod = {}, {}
    for g in ordem:
        por_cnpj.setdefault(g["cnpj"], []).append(g)
        por_cod.setdefault(cod(g), []).append(g)
    for g in ordem:
        outros = [x for x in por_cnpj[g["cnpj"]] + por_cod[cod(g)] if x is not g]
        for x in outros:
            aviso = (f"lista associa {formatar_cnpj(g['cnpj'])} e {codigo_anbima(*cod(g))} "
                     f"também a {formatar_cnpj(x['cnpj'])} / {codigo_anbima(*cod(x))} — confira")
            if aviso not in g["avisos"]:
                g["avisos"].append(aviso)

    soltos = {}
    for i, f in enumerate(fundos):
        if f["cnpj"] and f["num"]:
            continue
        alvo = (por_cnpj.get(f["cnpj"]) if f["cnpj"] else por_cod.get(cod(f))) or []
        chave = ("cnpj", f["cnpj"]) if f["cnpj"] else ("cod", cod(f))
        dest = alvo[0] if alvo else soltos.get(chave)
        if dest is None:
            dest = soltos[chave] = dict(f, avisos=list(f["avisos"]))
            ordem.append(dest)
        else:
            absorver(dest, f)
        pos[id(dest)] = min(pos.get(id(dest), i), i)
    return sorted(ordem, key=lambda g: pos[id(g)])  # ordem de 1ª aparição no arquivo


def resolver_fundos(entradas, lista):
    """Traduz as entradas da linha de comando (CNPJ ou código ANBIMA) em fundos, completando
    CNPJ <-> código pela lista. Sem entradas, usa a lista inteira.
    Devolve (fundos, entradas_invalidas)."""
    invalidas = []
    if not entradas:
        fundos = [dict(f, avisos=list(f["avisos"])) for f in lista]
    else:
        brutos = []
        for ent in entradas:
            cnpj = ler_cnpj(ent)
            cod = None if cnpj else ler_codigo_anbima(ent)
            if not cnpj and not cod:
                print(f"[AVISO] {ent!r} não é CNPJ nem código ANBIMA — ignorado.")
                invalidas.append(ent)
                continue
            achados = [f for f in lista if (cnpj and f["cnpj"] == cnpj)
                       or (cod and f["num"] and (f["prefixo"], f["num"]) == cod)]
            if not achados:
                achados = [dict(cnpj=cnpj or "", prefixo=cod[0] if cod else "C", num=cod[1] if cod else None,
                                nome="", avisos=[])]
            brutos += [dict(f, entrada=ent, avisos=list(f["avisos"])) for f in achados]
        fundos = _juntar(brutos)

    for f in fundos:
        if f["cnpj"] and not cnpj_valido(f["cnpj"]):
            f["avisos"].append("CNPJ com dígito verificador inválido")
        if not f["num"]:
            f["avisos"].append("sem código ANBIMA: Broadcast não consultado (só CVM)")
        if not f["cnpj"]:
            f["avisos"].append("sem CNPJ: CVM não consultada (sem razão social)")
    return fundos, invalidas


# ── Broadcast (Excel + add-in) ────────────────────────────────────────────────
# RPC_E_CALL_REJECTED, RPC_E_SERVERCALL_RETRYLATER, VBA_E_IGNORE (Excel ocupado em modo de edição/cálculo)
_EXCEL_OCUPADO = (-2147418111, -2147417846, -2146777998)
# atributos que o script usa no Excel; no late binding do pywin32 uma chamada recusada por
# "Excel ocupado" pode virar AttributeError('<nome>.<atributo>') ou
# AttributeError("Property '<nome>.<atributo>' can not be set.")
_ATRIBUTOS_EXCEL = {"Text", "Formula2", "Formula", "CurrentRegion", "Value", "Worksheets", "Add", "Close",
                    "Count", "RegisterXLL", "Visible", "DisplayAlerts", "Workbooks", "Range", "Quit"}


def _recusa_do_excel(e):
    m = re.fullmatch(r"(?:Property ')?.*\.(\w+)(?:' can not be set\.)?", str(e))
    return bool(m and m.group(1) in _ATRIBUTOS_EXCEL)


def _com(fn, tentativas=120, espera=0.5):
    """Chama fn() repetindo enquanto o Excel recusar a chamada por estar ocupado
    (acontece enquanto o add-in calcula/recebe dados)."""
    import pywintypes
    for i in range(tentativas):
        try:
            return fn()
        except pywintypes.com_error as e:
            if e.hresult not in _EXCEL_OCUPADO or i == tentativas - 1:
                raise
        except AttributeError as e:
            if not _recusa_do_excel(e) or i == tentativas - 1:
                raise
        time.sleep(espera)


def _porta_terminal():
    try:
        txt = open(os.path.join(AE_BROADCAST_DIR, "aesp.cfg"), encoding="latin-1").read()
        m = re.search(r"^\s*PORTA\s*=\s*(\d+)", txt, re.M)
        return int(m.group(1)) if m else 8100
    except OSError:
        return 8100


def terminal_no_ar():
    """True se o terminal Broadcast está aceitando conexões do add-in (porta do aesp.cfg)."""
    try:
        with socket.create_connection(("127.0.0.1", _porta_terminal()), timeout=2):
            return True
    except OSError:
        return False


def _formula_bch(num, inicio, fim):
    campos = ";".join(c for c, _ in BCH_CAMPOS)
    return (f'=BCH("{codigo_broadcast(num)}","{campos}","{inicio:%d/%m/%Y}","{fim:%d/%m/%Y}",'
            f'"ASC=DRF","TITULO=S")')


def _ler_resultado(ws):
    """Converte o bloco devolvido pelo BCH (título na 1ª linha) em DataFrame."""
    valores = _com(lambda: ws.Range("A1").CurrentRegion.Value)
    colunas = [n for _, n in BCH_CAMPOS]
    if not isinstance(valores, tuple) or len(valores) < 2:
        return pd.DataFrame(columns=colunas)
    nomes = dict(BCH_CAMPOS)
    df = pd.DataFrame([list(v) for v in valores[1:]],
                      columns=[nomes.get(str(t).strip().upper(), str(t).lower()) for t in valores[0]])
    for col in df.columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")
        # célula com erro do Excel (#N/A, #VALOR!...) chega como inteiro negativo gigante
        df.loc[df[col].between(ERRO_EXCEL_MIN, ERRO_EXCEL_MAX), col] = float("nan")
    # DRF vem como data OLE Automation (nº de dias desde 30/12/1899)
    df.loc[(df["data"] < 1) | (df["data"] > 2958465), "data"] = float("nan")
    df["data"] = pd.to_datetime(df["data"], unit="D", origin="1899-12-30").dt.normalize()
    df = df.dropna(subset=["data"]).sort_values("data")
    return df.drop_duplicates("data", keep="last").reset_index(drop=True)


def consultar_broadcast(nums, inicio, fim, timeout=120.0, lote=40):
    """Roda um BCH por código (inicio..fim) num Excel oculto, em lotes.
    Devolve {num: DataFrame | 'mensagem de erro'} — erro num fundo ou lote não derruba os demais."""
    import pythoncom
    import win32com.client

    if not os.path.exists(ADDIN_XLL):
        raise RuntimeError(f"add-in do Broadcast não encontrado em {ADDIN_XLL!r}")

    resultados = {}
    pythoncom.CoInitialize()
    xl = None
    try:
        xl = win32com.client.DispatchEx("Excel.Application")  # instância própria
        _com(lambda: setattr(xl, "Visible", False))
        _com(lambda: setattr(xl, "DisplayAlerts", False))
        if not _com(lambda: xl.RegisterXLL(ADDIN_XLL)):
            raise RuntimeError("Excel não carregou o add-in do Broadcast")

        atributo = None  # "Formula2" (Excel 365/2021, matriz dinâmica) ou "Formula" (versões antigas)
        for ini in range(0, len(nums), lote):
            parte = nums[ini:ini + lote]
            wb = None
            try:
                wb = _com(lambda: xl.Workbooks.Add())
                atributo = atributo or _atributo_formula(wb)
                resultados.update(_rodar_lote(wb, parte, inicio, fim, timeout, pythoncom, atributo))
            except Exception as e:  # noqa — Excel caiu/travou: guarda o que já veio e para
                for n in nums[ini:]:
                    resultados.setdefault(n, f"não consultado: erro no Excel/Broadcast ({e!r})")
                break
            finally:
                if wb is not None:
                    try:
                        _com(lambda: wb.Close(False), tentativas=20)
                    except Exception:  # noqa
                        pass
            # terminal sem conexão: não adianta insistir nos próximos lotes
            if all("#BCONN" in str(resultados.get(n)) for n in parte):
                for n in nums[ini + lote:]:
                    resultados[n] = "não consultado: terminal sem conexão (#BCONN)"
                break
        return resultados
    finally:
        if xl is not None:
            try:
                xl.Quit()
            except Exception:  # noqa
                pass
        pythoncom.CoUninitialize()


def _atributo_formula(wb):
    """'Formula2' se o Excel tiver matriz dinâmica (365/2021); senão 'Formula' (o add-in
    Excel-DNA redimensiona o resultado sozinho). Sem isso, num Excel antigo o 'membro
    inexistente' pareceria 'Excel ocupado' e cada fundo esperaria 1 min à toa."""
    import pywintypes
    rng = _com(lambda: wb.Worksheets(1).Range("A1"))
    try:
        _com(lambda: rng._oleobj_.GetIDsOfNames("Formula2"))
        return "Formula2"
    except pywintypes.com_error as e:
        if e.hresult == -2147352570:  # DISP_E_UNKNOWNNAME
            print("[AVISO] Excel sem Range.Formula2 (versão anterior à 2021): usando Range.Formula")
            return "Formula"
        raise


def _rodar_lote(wb, nums, inicio, fim, timeout, pythoncom, atributo="Formula2"):
    abas = {}
    erros = {}    # fundo sem fórmula escrita: não há o que ler
    falhas = {}   # erro passageiro ao consultar/reescrever: a leitura final tenta de novo
    textos = {}

    def escrever(n):
        f = _formula_bch(n, inicio, fim)
        _com(lambda: setattr(abas[n].Range("A1"), atributo, f))

    def topo(n, tentativas=120):
        return str(_com(lambda: abas[n].Range("A1").Text, tentativas=tentativas)).strip().upper()

    # uma aba por fundo: o BCH "derrama" linhas e colunas a partir de A1
    falhas_seguidas = 0
    for i, n in enumerate(nums):
        if falhas_seguidas >= 3:  # Excel travado: não insiste nos demais
            erros[n] = "não consultado: Excel parou de responder ao criar as abas"
            continue
        try:
            abas[n] = _com(lambda: wb.Worksheets(1) if i == 0
                           else wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count)))
            escrever(n)
            falhas_seguidas = 0
        except Exception as e:  # noqa — erro num fundo não derruba os outros
            erros[n] = f"erro ao criar a aba/escrever a fórmula: {e!r}"
            falhas_seguidas += 1

    # o add-in conecta no terminal de forma assíncrona: nos primeiros segundos a célula
    # mostra #BCONN — reescreve a fórmula até a conexão sair (sem mexer em quem já está buscando)
    pendentes = set(nums) - set(erros)
    limite_bconn = max(DESISTE_BCONN_S, 0.4 * timeout)
    t0 = ultimo_reenvio = time.time()
    t_consulta = None      # 1ª volta de consulta que o Excel respondeu (o prazo do #BCONN conta daí)
    reenvios = 0
    ultimo_print = 0.0
    while pendentes and time.time() - t0 < timeout:
        pythoncom.PumpWaitingMessages()
        time.sleep(1.0)
        # Excel ocupado (o add-in calculando) recusa chamadas por uns segundos: nesse caso
        # larga a volta e tenta de novo na seguinte, até o --timeout — nada de desistir cedo
        ocupado = False
        for n in list(pendentes):
            if time.time() - t0 >= timeout:
                break
            try:
                textos[n] = topo(n, tentativas=4)
            except Exception as e:  # noqa
                falhas[n] = f"erro ao consultar a célula: {e!r}"
                ocupado = True
                break
            if textos[n] not in PENDENTES:
                pendentes.discard(n)
                falhas.pop(n, None)
        if ocupado:
            continue
        if t_consulta is None and textos:
            t_consulta = time.time()
        if (pendentes and t_consulta and reenvios >= 1 and time.time() - t_consulta > limite_bconn
                and all(textos.get(n) == "#BCONN" for n in pendentes) and not any(t == "DRF" for t in textos.values())):
            break  # nada conectou mesmo reescrevendo: terminal aberto mas deslogado
        if pendentes and time.time() - ultimo_reenvio > 6:
            for n in [n for n in pendentes if textos.get(n, "") in REENVIAR]:
                try:
                    f = _formula_bch(n, inicio, fim)
                    _com(lambda: setattr(abas[n].Range("A1"), atributo, f), tentativas=4)
                except Exception as e:  # noqa — ocupado: reescreve na próxima rodada
                    falhas[n] = f"erro ao reescrever a fórmula: {e!r}"
                    break
            else:
                reenvios += 1  # só conta a rodada de reenvio que foi até o fim
            ultimo_reenvio = time.time()
        if pendentes and time.time() - ultimo_print > 10:
            print(f"[INFO] aguardando Broadcast... {len(nums) - len(pendentes)}/{len(nums)} fundo(s) prontos")
            ultimo_print = time.time()

    res, falhas_seguidas = {}, 0
    for n in nums:
        if n in erros:
            res[n] = erros[n]
            continue
        if falhas_seguidas >= 3:  # Excel travado: não gasta mais 1 min por fundo
            res[n] = falhas.get(n) or "não lido: Excel parou de responder"
            continue
        try:
            t = topo(n)
            if t == "DRF":
                res[n] = _ler_resultado(abas[n])
            elif n in pendentes or (n in falhas and t in PENDENTES):
                res[n] = f"sem resposta em {int(time.time() - t0)}s ({t or 'vazio'})" + (
                    f": {ERROS_BC[t]}" if t in ERROS_BC else "")
            else:
                res[n] = f"{t}: {ERROS_BC.get(t, 'erro desconhecido do add-in')}"
            falhas_seguidas = 0
        except Exception as e:  # noqa — erro ao ler um fundo não derruba os outros
            res[n] = falhas.get(n) or f"erro ao ler o resultado: {e!r}"
            falhas_seguidas += 1
    return res


# ── CVM ──────────────────────────────────────────────────────────────────────
def _baixar(url, destino, max_idade_h):
    """Baixa url -> destino, reaproveitando o cache se tiver menos de max_idade_h horas.
    Só troca o cache por um zip completo e íntegro; se o download falhar e houver cache
    (mesmo velho), usa o cache com aviso."""
    if os.path.exists(destino) and time.time() - os.path.getmtime(destino) < max_idade_h * 3600:
        return destino
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    tmp = destino + ".part"
    try:
        req = urllib.request.Request(url, headers=UA)
        with urllib.request.urlopen(req, timeout=300) as resp, open(tmp, "wb") as fh:
            esperado = resp.headers.get("Content-Length")
            shutil.copyfileobj(resp, fh)
        if esperado and os.path.getsize(tmp) != int(esperado):
            raise IOError(f"download incompleto ({os.path.getsize(tmp)} de {esperado} bytes)")
        try:
            with zipfile.ZipFile(tmp) as z:
                ruim = z.testzip()
        except zipfile.BadZipFile:
            raise IOError("resposta não é um zip (página de erro/manutenção?)")
        if ruim:
            raise IOError(f"zip corrompido ({ruim})")
        os.replace(tmp, destino)
    except Exception as e:  # noqa
        if os.path.exists(tmp):
            os.remove(tmp)
        if os.path.exists(destino):  # inclusive 404 passageiro: o arquivo existia, o cache vale
            print(f"[AVISO] falha ao atualizar {url} ({e!r}); usando cache de "
                  f"{datetime.fromtimestamp(os.path.getmtime(destino)):%d/%m/%Y %H:%M}.")
            return destino
        raise
    return destino


def _abrir_zip(url, max_idade_h):
    """ZipFile do arquivo da CVM (cache); cache corrompido é apagado e baixado de novo."""
    destino = os.path.join(CACHE_CVM, os.path.basename(url))
    caminho = _baixar(url, destino, max_idade_h)
    try:
        return zipfile.ZipFile(caminho)
    except zipfile.BadZipFile:
        os.remove(caminho)
        return zipfile.ZipFile(_baixar(url, destino, max_idade_h))


def _ler_informe(z, mes, alvo):
    """Linhas do informe diário (zip aberto) dos CNPJs em `alvo`, no mês `mes` (AAAAMM).
    Devolve (partes, última DT_COMPTC publicada no mês por qualquer fundo)."""
    partes, ultima = [], ""
    for nome in z.namelist():
        if not nome.lower().endswith(".csv"):
            continue
        # zip anual: 1 CSV por mês (2005-2020) ou 1 CSV do ano todo (2000-2004)
        m = re.search(r"(\d{6})\.csv$", nome)
        if m and m.group(1) != mes:
            continue
        with z.open(nome) as fh:
            df = pd.read_csv(fh, sep=";", dtype=str, encoding="cp1252", encoding_errors="replace")
        # antes da RCVM 175 as colunas eram CNPJ_FUNDO / TP_FUNDO e não havia subclasse
        df = df.rename(columns={"CNPJ_FUNDO": "CNPJ_FUNDO_CLASSE", "TP_FUNDO": "TP_FUNDO_CLASSE"})
        for col in ("ID_SUBCLASSE", "TP_FUNDO_CLASSE"):
            if col not in df.columns:
                df[col] = ""
        do_mes = df["DT_COMPTC"].str.startswith(f"{mes[:4]}-{mes[4:]}", na=False)
        ultima = max(ultima, df.loc[do_mes, "DT_COMPTC"].max() if do_mes.any() else "")
        df = df[do_mes]
        df = df.assign(cnpj=df["CNPJ_FUNDO_CLASSE"].map(normalizar_cnpj))
        partes.append(df[df["cnpj"].isin(alvo)])
    return partes, ultima


def _zip_do_mes(mes_dia):
    if mes_dia.year >= 2021:
        # arquivo do mês corrente/anterior é atualizado todo dia pela CVM
        recente = (datetime.now().date() - mes_dia).days <= 62
        return _abrir_zip(CVM_INF_MENSAL.format(mes_dia), 6 if recente else 24 * 30)
    return _abrir_zip(CVM_INF_ANUAL.format(mes_dia), 24 * 30)


def _serie(df):
    """Linhas do informe (colunas da CVM) -> série no formato interno (data + CAMPOS)."""
    s = df.rename(columns=CVM_CAMPOS).assign(data=pd.to_datetime(df["DT_COMPTC"]))
    return s[["data"] + CAMPOS].drop_duplicates("data", keep="last").sort_values("data").reset_index(drop=True)


def _consolidar_cvm(g):
    """Linhas de um CNPJ no informe -> DataFrame (data + CAMPOS + obs + prov_sub + n_sub),
    1 linha por dia. Classe que só informa por subclasse: soma as subclasses, checa a
    consistência de cada uma (a soma não tem cota p/ checar) e marca a soma incompleta."""
    subclasses = {sid: _serie(s) for sid, s in g[g["ID_SUBCLASSE"] != ""].groupby("ID_SUBCLASSE")}
    datas = sorted(g["DT_COMPTC"].unique())
    linhas = []
    for data in datas:
        gd = g[g["DT_COMPTC"] == data]
        classe = gd[gd["ID_SUBCLASSE"] == ""].drop_duplicates(subset=list(CVM_CAMPOS))
        obs, prov, n_sub, previa_sub = [], set(), 0, False
        if not classe.empty:
            if len(classe) > 1:
                novas = classe[classe["TP_FUNDO_CLASSE"].isin(CVM_TIPOS_NOVOS)]
                usada = (novas if not novas.empty else classe).iloc[-1]
                obs.append(f"CVM: {len(classe)} linhas diferentes para o dia "
                           f"({', '.join(classe['TP_FUNDO_CLASSE'])}); usada a de tipo {usada['TP_FUNDO_CLASSE']!r}")
            else:
                usada = classe.iloc[0]
            rec = {campo: usada[col] for col, campo in CVM_CAMPOS.items()}
        else:
            # só há linhas por subclasse: soma fluxos, PL e cotistas; cota não é somável
            sub_dia = set(gd["ID_SUBCLASSE"])
            n_sub = len(sub_dia)
            rec = {campo: gd[col].sum(min_count=1) for col, campo in CVM_CAMPOS.items() if campo != "cota"}
            rec["cota"] = float("nan")
            obs.append(f"CVM: soma de {n_sub} subclasse(s), cota não informada")
            # subclasse vista antes na janela (e não encerrada: último PL > 0) que faltou no dia
            data_ts = pd.Timestamp(data)
            faltam = [sid for sid, s in subclasses.items() if sid not in sub_dia
                      and not s[s["data"] < data_ts].empty and _num(s[s["data"] < data_ts]["patrimonio"].iloc[-1]) != 0]
            if faltam:
                prov |= set(FLUXOS)
                obs.append(f"CVM: soma incompleta — {len(faltam)} subclasse(s) sem informe no dia")
            marcadas = []
            for sid in sorted(sub_dia):
                a = analisar_serie(subclasses[sid], data_ts.date())
                previa_sub = previa_sub or a["previa"]
                if a["provisorios"]:
                    prov |= a["provisorios"] & set(FLUXOS)
                    marcadas.append(f"{sid} ({a['status']})")
            if marcadas:
                obs.append("CVM: subclasse(s) provisória(s): " + ", ".join(marcadas))
        rec.update(data=pd.Timestamp(data), obs="; ".join(obs), prov_sub=prov, n_sub=n_sub,
                   n_sub_janela=len(subclasses), previa_sub=previa_sub)
        linhas.append(rec)
    return pd.DataFrame(linhas).sort_values("data").reset_index(drop=True)


def cvm_informe_diario(dia, cnpjs, inicio=None):
    """Série do Informe Diário da CVM de `inicio` (padrão: `dia`) até `dia`, p/ os CNPJs
    pedidos — lê os meses do período e, se `inicio` cai no começo do mês, também o anterior
    (histórico p/ a checagem de D-1). Devolve dict:
      series       {cnpj: DataFrame(data, CAMPOS, obs, prov_sub, n_sub)} até `dia` inclusive
      presentes    CNPJs que aparecem nos arquivos lidos
      ultima_fundo {cnpj: último dia com informe do fundo (mesmo depois de `dia`)}
      ultima_geral último dia publicado pela CVM no mês de `dia` ('' se o mês ainda não existe)
      sem_d1       True se o mês anterior (histórico p/ a checagem) não pôde ser lido"""
    alvo = set(cnpjs)
    desde = inicio or dia
    info = dict(series={}, presentes=set(), ultima_fundo={}, ultima_geral="", sem_d1=False)
    m = (desde.replace(day=1) - timedelta(days=1)).replace(day=1) if desde.day <= 10 else desde.replace(day=1)
    meses = []
    while (m.year, m.month) < (dia.year, dia.month):
        meses.append(m)
        m = (m + timedelta(days=32)).replace(day=1)
    meses.append(dia)  # mês de `dia` (a própria data: dela sai ultima_geral)
    so_historico = [m for m in meses if (m.year, m.month) < (desde.year, desde.month)]
    partes = []
    for mes_dia in meses:
        try:
            with _zip_do_mes(mes_dia) as z:
                p, ultima = _ler_informe(z, f"{mes_dia:%Y%m}", alvo)
            partes += p
            if mes_dia == dia:
                info["ultima_geral"] = ultima
        except Exception as e:  # noqa
            if mes_dia not in so_historico:
                if isinstance(e, urllib.error.HTTPError) and e.code == 404:
                    continue  # CVM ainda não criou o arquivo do mês: nada publicado
                raise
            # mês anterior ao período é só histórico p/ a checagem
            info["sem_d1"] = True
            print(f"[AVISO] informe da CVM de {mes_dia:%m/%Y} indisponível ({e!r}): sem D-1 p/ a checagem do PL")
    if not partes:
        return info
    df = pd.concat(partes, ignore_index=True)
    info["presentes"] = set(df["cnpj"])
    info["ultima_fundo"] = df.groupby("cnpj")["DT_COMPTC"].max().to_dict()
    df = df[df["DT_COMPTC"] <= dia.isoformat()]
    for col in CVM_CAMPOS:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["ID_SUBCLASSE"] = df["ID_SUBCLASSE"].fillna("").str.strip()
    df["TP_FUNDO_CLASSE"] = df["TP_FUNDO_CLASSE"].fillna("").str.strip()
    for cnpj, g in df.groupby("cnpj"):
        datas = sorted(g["DT_COMPTC"].unique())
        primeira = next((k for k, d in enumerate(datas) if d >= desde.isoformat()), len(datas))
        datas = datas[max(0, primeira - JANELA_BCH_DIAS):]  # o período e até 10 dias antes dele
        info["series"][cnpj] = _consolidar_cvm(g[g["DT_COMPTC"].isin(datas)])
    return info


def cadastro_vazio():
    return dict(nomes={}, tipos={}, situacoes={}, inicio={}, n_sub={})


def cvm_cadastro(cnpjs):
    """Cadastro da CVM p/ os CNPJs: dict com nomes (denominação social), tipos (tipo de
    classe), situacoes, inicio (Data_Inicio da classe, AAAA-MM-DD) e n_sub (subclasses não
    canceladas da classe). Classe primeiro, depois fundo; registro ativo tem preferência."""
    alvo = set(cnpjs)
    cad = cadastro_vazio()
    with _abrir_zip(CVM_CADASTRO, 24) as z:
        id_classe = {}
        for arquivo, col, col_tipo in (("registro_classe.csv", "CNPJ_Classe", "Tipo_Classe"),
                                       ("registro_fundo.csv", "CNPJ_Fundo", "Tipo_Fundo")):
            extras = ["ID_Registro_Classe", "Data_Inicio"] if col_tipo == "Tipo_Classe" else []
            with z.open(arquivo) as fh:
                df = pd.read_csv(fh, sep=";", dtype=str, encoding="cp1252", encoding_errors="replace",
                                 usecols=[col, "Denominacao_Social", "Situacao", col_tipo] + extras)
            df = df.assign(cnpj=df[col].map(normalizar_cnpj))
            df = df[df["cnpj"].isin(alvo - set(cad["nomes"]))]
            df = df.assign(cancelado=df["Situacao"].fillna("").str.contains("Cancel"))
            for cnpj, g in df.sort_values("cancelado", kind="stable").groupby("cnpj", sort=False):
                # o cadastro tem espaços não separáveis (U+00A0) em alguns nomes
                cad["nomes"][cnpj] = str(g["Denominacao_Social"].iloc[0] or "").replace("\xa0", " ").strip()
                cad["situacoes"][cnpj] = str(g["Situacao"].iloc[0] or "")
                if col_tipo == "Tipo_Classe":
                    cad["tipos"][cnpj] = str(g[col_tipo].iloc[0] or "")
                    cad["inicio"][cnpj] = str(g["Data_Inicio"].iloc[0] or "")
                    id_classe[str(g["ID_Registro_Classe"].iloc[0])] = cnpj
        if id_classe and "registro_subclasse.csv" in z.namelist():
            with z.open("registro_subclasse.csv") as fh:
                sub = pd.read_csv(fh, sep=";", dtype=str, encoding="cp1252", encoding_errors="replace",
                                  usecols=["ID_Registro_Classe", "Situacao"])
            sub = sub[sub["ID_Registro_Classe"].isin(id_classe) & ~sub["Situacao"].fillna("").str.contains("Cancel")]
            for idc, n in sub.groupby("ID_Registro_Classe").size().items():
                cad["n_sub"][id_classe[idc]] = int(n)
    return cad


# ── consistência e combinação das fontes ──────────────────────────────────────
def _num(x):
    return float(x) if x is not None and pd.notna(x) else None


# regra que marcou a linha -> (status, campos provisórios)
REGRAS = {
    "previa": ("prévia (cotistas = 0)", set(FLUXOS)),
    "repetida": ("linha repetida de D-1", set(CAMPOS)),
    "possivel_previa": ("possível prévia (fluxos zerados)", set(FLUXOS)),
    "pl_rolado": ("PL provisório", {"patrimonio"}),
    "pl_repetido": ("PL provisório", {"patrimonio"}),
}


def _linha_igual(v, va):
    """Duas linhas (dicts de CAMPOS) iguais: cota até 1e-9 relativo; demais até R$ 0,005 / exato."""
    if not all(_tem(v[c]) and _tem(va[c]) for c in CAMPOS) or not va["cota"]:
        return False
    return (abs(v["cota"] / va["cota"] - 1) <= 1e-9
            and all(abs(v[c] - va[c]) < 0.005 for c in ("captacao", "resgate", "patrimonio", "cotistas")))


def analisar_serie(serie, dia, _nivel=0):
    """Analisa a linha de `dia` de uma série (Broadcast ou CVM: colunas data + CAMPOS) contra
    os dias anteriores. Devolve dict: valores (ou None), status, regra (chave de REGRAS ou
    None), provisorios (campos), previa, pl_repetido, residuo_pl, fluxo_tipico, liquido,
    amortizacao (valor, se o resgate for amortização) e d1_provisorio (D-1 também marcado)."""
    r = dict(valores=None, regra=None, provisorios=set(), previa=False, pl_repetido=False,
             residuo_pl=None, fluxo_tipico=None, liquido=None, amortizacao=None, d1_provisorio=None)
    if serie is None:
        return dict(r, status="não consultado")
    if isinstance(serie, str):
        return dict(r, status=serie)
    alvo = pd.Timestamp(dia)
    linha = serie[serie["data"] == alvo]
    if linha.empty:
        ult = serie["data"].max()
        return dict(r, status=(f"sem dado em {dia:%d/%m/%Y} (último: {ult:%d/%m/%Y})" if pd.notna(ult)
                               else f"sem dados de {dia - timedelta(days=JANELA_BCH_DIAS):%d/%m/%Y} a {dia:%d/%m/%Y}"))
    v = {c: _num(linha.iloc[-1].get(c)) for c in CAMPOS}
    r["valores"] = v
    pl, cota = v["patrimonio"], v["cota"]
    capt, resg = v["captacao"] or 0.0, v["resgate"] or 0.0
    r["liquido"] = liquido = capt - resg
    r["previa"] = v["cotistas"] == 0 and bool(pl) and pl > 0
    regra = "previa" if r["previa"] else None

    anteriores = serie[serie["data"] < alvo].tail(JANELA_BCH_DIAS)
    if not anteriores.empty and pl:
        a = anteriores.iloc[-1]
        va = {c: _num(a.get(c)) for c in CAMPOS}
        if _nivel == 0:  # a checagem de D usa D-1 como base: D-1 marcado deixa a checagem sem valor
            r["d1_provisorio"] = analisar_serie(serie[serie["data"] < alvo], a["data"].date(), _nivel=1)["regra"]
        if va["patrimonio"] is not None:
            r["pl_repetido"] = abs(va["patrimonio"] - pl) < 0.005
        cotas = [x for x in (_num(c) for c in anteriores["cota"]) if x]
        linhas_ant = [{c: _num(x.get(c)) for c in CAMPOS} for _, x in anteriores.iterrows()]
        # a cota mexe todo dia? (>= 80% dos dias da janela) — fundo com cota atualizada só de
        # vez em quando, ou parado, tem linhas iguais de verdade. Dia que já era cópia do
        # anterior não conta como "cota parada" (senão duas cópias seguidas desarmam a regra)
        mudancas = [abs(y["cota"] / x["cota"] - 1) > 1e-9 for x, y in zip(linhas_ant, linhas_ant[1:])
                    if x["cota"] and y["cota"] and not _linha_igual(y, x)]
        cota_diaria = len(mudancas) >= 3 and sum(mudancas) >= 0.8 * len(mudancas)
        if regra is None and _linha_igual(v, va) and dia_util(dia) and (capt or resg or cota_diaria):
            regra = "repetida"
        if va["patrimonio"] is not None and va["cota"] and cota:
            pl_ant = va["patrimonio"]
            razao = cota / va["cota"]
            rolado = pl_ant * razao
            r["residuo_pl"] = pl - rolado - liquido
            tol = max(1.0, 1e-7 * pl)            # TOL_PL_ROLADO
            relevante = max(1000.0, 1e-5 * pl)   # FLUXO_RELEVANTE
            pl_rolado = abs(pl - rolado) <= tol
            cota_mexeu = abs(razao - 1) > 1e-9
            # amortização lançada como resgate: a cota cai no valor amortizado e o PL só "rola"
            # (não desconta o resgate de novo — num resgate comum PL = rolado - resgate).
            # valor esperado = PL(D-1) x (1 + rendimento típico da cota - razão); tolerância 10%
            rend = [b / a_ - 1 for a_, b in zip(cotas, cotas[1:])]
            rend_tipico = statistics.median(rend) if rend else 0.0
            amort_esperada = pl_ant * (1 + rend_tipico - razao)
            if (liquido < 0 and razao < 1 and amort_esperada > 0
                    and abs(pl - rolado) <= 0.10 * -liquido                  # PL não absorveu o fluxo à parte
                    and abs(-liquido - amort_esperada) <= 0.10 * -liquido):  # queda da cota = valor amortizado
                r["amortizacao"] = resg or -liquido
            # fluxo típico: o fundo movimenta quase todo dia?
            brutos = [abs(_num(x.get("captacao")) or 0.0) + abs(_num(x.get("resgate")) or 0.0)
                      for _, x in anteriores.iterrows()]
            com_fluxo = [b for b in brutos if b > 0]
            diario = len(brutos) >= 3 and len(com_fluxo) >= 0.8 * len(brutos)
            r["fluxo_tipico"] = statistics.median(com_fluxo) if com_fluxo else 0.0
            if regra is None:
                if (capt == 0 and resg == 0 and va["cotistas"] == v["cotistas"]
                        and (pl_rolado or r["pl_repetido"]) and diario
                        and r["fluxo_tipico"] > max(1000.0, 1e-3 * pl)):   # FLUXO_TIPICO_RELEVANTE
                    regra = "possivel_previa"
                elif pl_rolado and abs(liquido) > relevante and not r["amortizacao"]:
                    regra = "pl_rolado"
                elif r["pl_repetido"] and cota_mexeu and abs(pl - rolado) > relevante:
                    regra = "pl_repetido"
    r["regra"] = regra
    if regra:
        r["status"], r["provisorios"] = REGRAS[regra][0], set(REGRAS[regra][1])
    else:
        r["status"] = "ok"
    return r


def _motivo(fonte, a):
    """Texto de Observações p/ o que a análise marcou numa fonte (pela regra que disparou)."""
    regra = a.get("regra")
    textos = {
        "previa": f"{fonte} em prévia (cotistas = 0; dia ainda não fechado)",
        "repetida": f"{fonte}: linha idêntica à do dia anterior (cota, PL, fluxos e cotistas repetidos) — dado de D-1 reenviado",
        "possivel_previa": (f"{fonte}: possível prévia — captação e resgate zerados, cotistas iguais aos da véspera e PL "
                            + ("idêntico ao da véspera" if a.get("pl_repetido") else "só corrigido pela cota")
                            + f", num fundo que movimenta ~R$ {_brl(a.get('fluxo_tipico') or 0)} por dia"),
        "pl_rolado": (f"{fonte}: PL provisório — igual ao do dia anterior corrigido pela cota, sem o fluxo líquido "
                      f"do dia (R$ {_brl(a.get('liquido') or 0)})"),
        "pl_repetido": f"{fonte}: PL provisório — idêntico ao do dia anterior, embora a cota tenha mudado",
    }
    partes = [textos[regra]] if regra else []
    if a.get("amortizacao"):
        partes.append(f"{fonte}: resgate de R$ {_brl(a['amortizacao'])} é amortização (a cota caiu nesse valor)")
    if a.get("d1_provisorio") and a.get("valores"):
        partes.append(f"{fonte}: dia anterior também estava provisório ({REGRAS[a['d1_provisorio']][0]}) — "
                      f"checagem do PL deste dia não é conclusiva")
    return "; ".join(partes)


def _iguais(campo, a, b):
    if campo == "cotistas":
        return round(a) == round(b)
    if campo == "cota":
        return math.isclose(a, b, rel_tol=1e-8, abs_tol=1e-6)
    return math.isclose(a, b, rel_tol=1e-6, abs_tol=0.01)


def _status_cvm_sem_linha(f, dia, info, cadastro):
    """Motivo de a CVM não ter a linha do dia p/ o fundo."""
    tipo = cadastro["tipos"].get(f["cnpj"], "")
    situacao = cadastro["situacoes"].get(f["cnpj"], "")
    inicio = cadastro["inicio"].get(f["cnpj"], "")
    d = dia.isoformat()
    if inicio > d:
        return f"sem dado em {dia:%d/%m/%Y} (a classe só começou em {inicio[8:]}/{inicio[5:7]}/{inicio[:4]})"
    if f["cnpj"] in info["presentes"]:
        extra = f"; situação no cadastro: {situacao}" if situacao and situacao != "Em Funcionamento Normal" else ""
        if info["ultima_fundo"].get(f["cnpj"], "") > d:
            return f"sem informe do fundo em {dia:%d/%m/%Y} (feriado ou dia não reportado; há dias posteriores{extra})"
        if extra:
            ult = info["ultima_fundo"].get(f["cnpj"], "")
            return f"sem dado em {dia:%d/%m/%Y} (último informe do fundo: {ult[8:]}/{ult[5:7]}/{ult[:4]}{extra})"
        if info["ultima_geral"] >= d:
            return f"sem dado em {dia:%d/%m/%Y} (o fundo ainda não entregou; outros já entregaram)"
        return f"sem dado em {dia:%d/%m/%Y} (CVM ainda não publicou o dia)"
    if tipo and "FIF" not in tipo.upper():
        return f"classe do tipo {tipo.replace('Classes de Cotas de Fundos ', '')} não entrega informe diário à CVM"
    msg = f"CNPJ não consta no informe diário da CVM de {dia:%m/%Y}"
    if situacao and situacao != "Em Funcionamento Normal":
        msg += f" (situação no cadastro: {situacao})"
    elif not info["ultima_geral"]:
        msg += " (CVM ainda não publicou o mês)"
    return msg


def _ultimo_dia_com_dado(f, dia, bc, bc_motivo, cvm_info, cvm_motivo):
    """Último dia antes de `dia` (dentro da janela) com linha do fundo em alguma das fontes."""
    datas = []
    serie_bc = bc.get(f["num"]) if f["num"] and f["prefixo"] == "C" and not bc_motivo else None
    serie_cvm = cvm_info["series"].get(f["cnpj"]) if f["cnpj"] and not cvm_motivo else None
    for s in (serie_bc, serie_cvm):
        if isinstance(s, pd.DataFrame) and not s.empty:
            datas += [d for d in s["data"] if d < pd.Timestamp(dia)]
    return max(datas).date() if datas else None


def montar_linhas(fundos, dia, bc, bc_motivo, cvm_info, cvm_motivo, cadastro, ultimo_disponivel=True):
    """Combina Broadcast e CVM campo a campo. Devolve lista de dicts (1 por fundo).
    Fundo sem dado no dia em NENHUMA fonte (ex.: ETF, que a ANBIMA publica com atraso e a CVM
    não cobre) sai com o último dia disponível da janela — a 'Data da Análise' da linha mostra
    esse dia e a aba Fontes marca a defasagem (ultimo_disponivel=False deixa em branco)."""
    linhas = []
    for f in fundos:
        ln = _combinar(f, dia, bc, bc_motivo, cvm_info, cvm_motivo, cadastro)
        if ultimo_disponivel and not any(_tem(v) for v in ln["valor"].values()):
            anterior = _ultimo_dia_com_dado(f, dia, bc, bc_motivo, cvm_info, cvm_motivo)
            ln_ant = _combinar(f, anterior, bc, bc_motivo, cvm_info, cvm_motivo, cadastro) if anterior else None
            if ln_ant and any(_tem(v) for v in ln_ant["valor"].values()):
                ln_ant["fonte"] = {k: f"{fo} — dado de {anterior:%d/%m}" if fo else "" for k, fo in ln_ant["fonte"].items()}
                ln_ant["obs"] = "; ".join(x for x in (
                    f"sem dado de {dia:%d/%m/%Y} em nenhuma fonte (Broadcast: {ln['status_bc']}; CVM: {ln['status_cvm']}) "
                    f"— usado o último dia disponível, {anterior:%d/%m/%Y}", ln_ant["obs"]) if x)
                ln = ln_ant
        linhas.append(ln)
    return linhas


def _combinar(f, dia, bc, bc_motivo, cvm_info, cvm_motivo, cadastro):
    """Linha de um fundo num dia: Broadcast e CVM combinados campo a campo."""
    nomes = cadastro["nomes"]
    series = cvm_info["series"]
    obs = list(f["avisos"])
    if not f["num"]:
        b = analisar_serie("sem código ANBIMA", dia)
    elif f["prefixo"] == "S":
        # o ticker "FI" + número vale p/ a classe; p/ subclasse não foi confirmado
        b = analisar_serie("código de subclasse ANBIMA (S…): ticker do Broadcast não confirmado — não consultado", dia)
    elif bc_motivo:
        b = analisar_serie(bc_motivo, dia)
    else:
        b = analisar_serie(bc.get(f["num"]), dia)

    linha_cvm = None
    if not f["cnpj"]:
        c = analisar_serie("sem CNPJ", dia)
    elif cvm_motivo:
        c = analisar_serie(cvm_motivo, dia)
    else:
        s = series.get(f["cnpj"])
        if s is not None and (s["data"] == pd.Timestamp(dia)).any():
            linha_cvm = s[s["data"] == pd.Timestamp(dia)].iloc[-1]
        if linha_cvm is None:
            c = analisar_serie(_status_cvm_sem_linha(f, dia, cvm_info, cadastro), dia)
        elif f["prefixo"] == "S" and max(linha_cvm["n_sub_janela"], cadastro["n_sub"].get(f["cnpj"], 0)) > 1:
            # a CVM não liga a subclasse ANBIMA à subclasse dela: a soma da classe não serve
            n = max(linha_cvm["n_sub_janela"], cadastro["n_sub"].get(f["cnpj"], 0))
            c = analisar_serie(f"CVM só tem a classe inteira ({n} subclasses); "
                               f"sem correspondência com a subclasse ANBIMA "
                               f"{codigo_anbima(f['prefixo'], f['num'])}", dia)
            linha_cvm = None
        else:
            c = analisar_serie(s, dia)
            if linha_cvm["previa_sub"]:
                c["previa"] = True  # alguma subclasse em prévia: a soma dos fluxos/cotistas não vale
            if linha_cvm["obs"]:
                obs.append(linha_cvm["obs"])
            if linha_cvm["prov_sub"]:
                c["provisorios"] |= linha_cvm["prov_sub"]
                if c["status"] == "ok":
                    c["status"] = "provisório (subclasses)"
            if cvm_info["sem_d1"] and c["residuo_pl"] is None:
                obs.append("CVM: sem o dia anterior (mês anterior indisponível) — checagem do PL não feita")
    vb, vc = b["valores"], c["valores"]

    valor, fonte = {}, {}
    for campo in CAMPOS:
        vb_c = vb.get(campo) if vb else None
        vc_c = vc.get(campo) if vc else None
        # zeros de prévia (cotistas = 0) nunca vão para a planilha, de nenhuma fonte
        if b["previa"] and campo in ("captacao", "resgate", "cotistas"):
            vb_c = None
        if c["previa"] and campo in ("captacao", "resgate", "cotistas"):
            vc_c = None
        b_prov, c_prov = campo in b["provisorios"], campo in c["provisorios"]
        if _tem(vb_c) and not (b_prov and _tem(vc_c) and not c_prov):
            valor[campo], fonte[campo] = vb_c, "Broadcast" + (" (provisório)" if b_prov else "")
        elif _tem(vc_c):
            valor[campo], fonte[campo] = vc_c, "CVM" + (" (provisório)" if c_prov else "")
        else:
            valor[campo], fonte[campo] = None, ""

    for nome_fonte, a in (("Broadcast", b), ("CVM", c)):
        m = _motivo(nome_fonte, a)
        if m:
            obs.append(m)
    usados_cvm = [NOMES_CAMPOS[k] for k in CAMPOS if fonte[k].startswith("CVM") and k in b["provisorios"]]
    if usados_cvm:
        obs.append("usados da CVM: " + ", ".join(usados_cvm))
    brancos = [NOMES_CAMPOS[k] for k in ("captacao", "resgate", "cotistas") if not fonte[k]
               and (b["previa"] or c["previa"])]
    if brancos:
        obs.append(f"{'/'.join(brancos)} em branco até a ANBIMA ou a CVM fecharem o dia")
    if not b["provisorios"] and not c["provisorios"] and b["pl_repetido"]:
        obs.append("PL do Broadcast igual ao do dia anterior")
    if vb and vc:
        div = [NOMES_CAMPOS[k] for k in CAMPOS
               if k not in b["provisorios"] and k not in c["provisorios"]
               and _tem(vb.get(k)) and _tem(vc.get(k)) and not _iguais(k, vb[k], vc[k])]
        if div:
            obs.append("Broadcast × CVM divergem em: " + ", ".join(div))

    return dict(
        data=dia, codigo_anbima=codigo_anbima(f["prefixo"], f["num"]), codigo_broadcast=codigo_broadcast(f["num"]),
        cnpj=formatar_cnpj(f["cnpj"]) if f["cnpj"] else "",
        razao_social=nomes.get(f["cnpj"]) or f.get("nome", ""),
        valor=valor, fonte=fonte, bc=vb or {}, cvm=vc or {},
        residuo_bc=b["residuo_pl"], residuo_cvm=c["residuo_pl"],
        status_bc=b["status"], status_cvm=c["status"], obs="; ".join(obs))


# ── planilha ─────────────────────────────────────────────────────────────────
def _titulo(linhas):
    fontes = {fo.split(" (")[0] for ln in linhas for fo in ln["fonte"].values() if fo}
    if fontes == {"Broadcast", "CVM"}:
        return "Broadcast+ e CVM"
    return "Broadcast+" if fontes == {"Broadcast"} else "CVM" if fontes == {"CVM"} else "sem dados"


def salvar_xlsx(linhas, dia, caminho, periodo=False):
    """Grava as abas Captacao e Fontes. Um dia: a data vai nos nomes das colunas (como no
    comdinheiro). Período (periodo=True): 1 linha por fundo e dia, colunas sem data."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    d = f"{dia:%d/%m/%Y}"
    wb = Workbook()
    ws = wb.active
    ws.title = "Captacao"
    ws.append([_titulo(linhas)])  # linha 1, no lugar do "comdinheiro" do layout original
    ws.append([titulo.replace(" {d}", "") if periodo else titulo.format(d=d) for _, titulo in LAYOUT])
    for ln in linhas:
        v = ln["valor"]
        cot = v["cotistas"]
        ws.append([ln["data"], ln["codigo_anbima"], ln["cnpj"], ln["razao_social"], v["cota"], v["captacao"],
                   v["resgate"], int(cot) if _tem(cot) and float(cot).is_integer() else cot, v["patrimonio"]])
        ws.cell(row=ws.max_row, column=1).number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A3"
    for col, larg in zip("ABCDEFGHI", (14, 16, 20, 70, 16, 18, 18, 14, 20)):
        ws.column_dimensions[col].width = larg

    fs = wb.create_sheet("Fontes")
    cab = ["Data da Análise", "Código Anbima", "Código Broadcast", "CNPJ", "Razão Social",
           "Status Broadcast", "Status CVM"]
    cab += [f"Fonte {NOMES_CAMPOS[c]}" for c in CAMPOS]
    for c in CAMPOS:
        cab += [f"{NOMES_CAMPOS[c]} Broadcast", f"{NOMES_CAMPOS[c]} CVM"]
    cab += ["Checagem PL Broadcast: PL - PL(D-1)×cota/cota(D-1) - (capt - resg)",
            "Checagem PL CVM: PL - PL(D-1)×cota/cota(D-1) - (capt - resg)", "Observações"]
    fs.append(cab)
    for cel in fs[1]:
        cel.font = Font(bold=True)
    for ln in linhas:
        lin = [ln["data"], ln["codigo_anbima"], ln["codigo_broadcast"], ln["cnpj"], ln["razao_social"],
               ln["status_bc"], ln["status_cvm"]]
        lin += [ln["fonte"][c] for c in CAMPOS]
        for c in CAMPOS:
            b, cv = ln["bc"].get(c), ln["cvm"].get(c)
            lin += [b if _tem(b) else None, cv if _tem(cv) else None]
        lin += [ln["residuo_bc"] if _tem(ln["residuo_bc"]) else None,
                ln["residuo_cvm"] if _tem(ln["residuo_cvm"]) else None, ln["obs"]]
        fs.append(lin)
        fs.cell(row=fs.max_row, column=1).number_format = "DD/MM/YYYY"
    fs.freeze_panes = "F2"
    for j, larg in enumerate([14, 14, 14, 20, 50, 34, 34] + [22] * 5 + [16] * 10 + [22, 22, 110], start=1):
        fs.column_dimensions[fs.cell(row=1, column=j).column_letter].width = larg

    pasta = os.path.dirname(os.path.abspath(caminho))
    os.makedirs(pasta, exist_ok=True)
    try:
        wb.save(caminho)
    except PermissionError:
        if not os.path.isfile(caminho):
            raise
        # arquivo já existe e está aberto (Excel trava o arquivo): salva com outro nome
        base, _ = os.path.splitext(caminho)
        caminho = f"{base}_{datetime.now():%H%M%S}.xlsx"
        print(f"[AVISO] {os.path.basename(base)}.xlsx está aberto (no Excel?); salvando como {os.path.basename(caminho)}")
        wb.save(caminho)
    return caminho


def _n_fundos_broadcast(caminho):
    """Quantos fundos de uma planilha já gravada têm resposta do Broadcast (algum valor bruto
    nas colunas '<campo> Broadcast' da aba Fontes, usado ou não na planilha)."""
    from openpyxl import load_workbook
    try:
        fs = load_workbook(caminho, read_only=True)["Fontes"]
        linhas = list(fs.iter_rows(values_only=True))
    except Exception:  # noqa
        return 0
    cab = list(linhas[0]) if linhas else []
    cols = [j for j, t in enumerate(cab) if str(t or "").endswith(" Broadcast") and not str(t).startswith("Status")]
    return sum(any(lin[j] is not None for j in cols) for lin in linhas[1:])


def _destino(out, dia, selecao=False, inicio=None):
    """Caminho do .xlsx a partir de --out (arquivo, pasta existente, pasta terminada em
    barra, ou vazio = padrão). Nome sem extensão conhecida ganha '.xlsx'. Rodada só com
    alguns fundos (linha de comando) ganha '_selecao' p/ não tomar o lugar da lista inteira;
    período vira captacao_<inicio>_a_<fim>."""
    datas = f"{inicio:%Y%m%d}_a_{dia:%Y%m%d}" if inicio and inicio < dia else f"{dia:%Y%m%d}"
    nome_padrao = f"captacao_{datas}{'_selecao' if selecao else ''}.xlsx"
    if not out:
        return os.path.join(SAIDA_DIR, nome_padrao)
    if os.path.isdir(out) or out.endswith(("/", "\\")):
        return os.path.join(out, nome_padrao)
    base, ext = os.path.splitext(out)
    if ext.lower() == ".xlsx":
        return out
    if ext.lower() in (".xls", ".xlsm", ".csv", ".txt"):
        print(f"[AVISO] a saída é sempre .xlsx; gravando em {os.path.basename(base)}.xlsx")
        return base + ".xlsx"
    return out + ".xlsx"  # sem extensão ou com ponto no nome (ex.: captacao_18.09)


def _pasta_gravavel(caminho):
    """Cria a pasta de saída e confere que dá p/ gravar nela (antes de gastar minutos consultando)."""
    pasta = os.path.dirname(os.path.abspath(caminho))
    try:
        os.makedirs(pasta, exist_ok=True)
        teste = os.path.join(pasta, f".captacao_teste_{os.getpid()}")
        with open(teste, "w"):
            pass
        os.remove(teste)
        return True
    except OSError as e:
        print(f"[ERRO] não dá para gravar a saída em {pasta}: {e}")
        return False


# ── CLI ───────────────────────────────────────────────────────────────────────
def _data(s):
    return datetime.strptime(s, "%d/%m/%Y").date()


def _pascoa(ano):
    """Domingo de Páscoa (algoritmo de Meeus/Jones/Butcher)."""
    a, b, c = ano % 19, ano // 100, ano % 100
    d, e = b // 4, b % 4
    g = (b - (b + 8) // 25 + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    return date(ano, (h + l - 7 * m + 114) // 31, (h + l - 7 * m + 114) % 31 + 1)


def feriados_nacionais(ano):
    """Feriados nacionais (calendário da ANBIMA): fixos + Carnaval, Sexta-feira Santa e
    Corpus Christi; Consciência Negra (20/11) a partir de 2024."""
    p = _pascoa(ano)
    fixos = [(1, 1), (4, 21), (5, 1), (9, 7), (10, 12), (11, 2), (11, 15), (12, 25)] + ([(11, 20)] if ano >= 2024 else [])
    return {date(ano, m, d) for m, d in fixos} | {p - timedelta(days=48), p - timedelta(days=47),
                                                   p - timedelta(days=2), p + timedelta(days=60)}


def dia_util(d):
    return d.weekday() < 5 and d not in feriados_nacionais(d.year)


def dia_util_anterior(d):
    """Dia útil anterior a d (seg-sex, fora os feriados nacionais)."""
    d -= timedelta(days=1)
    while not dia_util(d):
        d -= timedelta(days=1)
    return d


def main(argv=None):
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa
            pass

    hoje = datetime.now().date()
    ap = argparse.ArgumentParser(description="Cota/captação/resgate/cotistas/PL de fundos no layout do "
                                             "comdinheiro — Broadcast+ (add-in Excel) completado pela CVM.")
    ap.add_argument("fundos", nargs="*",
                    help="CNPJs e/ou códigos ANBIMA (C0000142697, 142697 ou FI142697). "
                         "Sem nenhum, usa todos os fundos da --lista.")
    ap.add_argument("--lista", "--mapa", dest="lista",
                    help="CSV/XLSX com CNPJ, código ANBIMA e nome dos fundos (padrão: ../input/fundos_captacao.csv); "
                         "aceita o próprio export do comdinheiro")
    ap.add_argument("--data", type=_data, default=dia_util_anterior(hoje),
                    help="data da análise dd/mm/aaaa (padrão: dia útil anterior a hoje); num período, é o último dia")
    periodo = ap.add_mutually_exclusive_group()
    periodo.add_argument("--inicio", type=_data,
                         help="período: 1º dia dd/mm/aaaa (vai até --data); sai 1 linha por fundo e dia útil")
    periodo.add_argument("--dias", type=int,
                         help="período: os últimos N dias úteis até --data (ex.: --dias 20)")
    ap.add_argument("--sem-broadcast", action="store_true", help="não consulta o Broadcast (só CVM)")
    ap.add_argument("--sem-cvm", action="store_true", help="não consulta a CVM (só Broadcast; sem razão social)")
    ap.add_argument("--so-data-exata", action="store_true",
                    help="fundo sem dado no dia em nenhuma fonte fica em branco (padrão: usa o último dia "
                         "disponível, com a data real na coluna 'Data da Análise' da linha)")
    ap.add_argument("--timeout", type=float, default=120.0, help="espera máxima (s) pela resposta do Broadcast (padrão 120)")
    ap.add_argument("--out", help="arquivo .xlsx ou pasta de saída (padrão: ../output/captacao/captacao_<yyyymmdd>.xlsx)")
    args = ap.parse_args(argv)

    dia = args.data
    if dia > hoje:
        ap.error("--data no futuro")
    inicio = args.inicio
    if args.dias is not None:
        if args.dias < 1:
            ap.error("--dias precisa ser >= 1")
        inicio = dia
        for _ in range(args.dias - 1):
            inicio = dia_util_anterior(inicio)
    if inicio and inicio > dia:
        ap.error("--inicio depois de --data")
    # dias da análise: só o --data, ou os dias úteis do período
    dias = [dia]
    if inicio and inicio < dia:
        dias = [inicio + timedelta(days=k) for k in range((dia - inicio).days + 1)]
        dias = [d for d in dias if dia_util(d)] or [dia]
    periodo = len(dias) > 1 or bool(inicio and inicio < dia)

    lista_path = args.lista or LISTA_PADRAO
    if args.lista and not os.path.isfile(args.lista):
        print(f"[ERRO] lista não encontrada (ou não é um arquivo): {os.path.abspath(args.lista)}")
        return 1
    try:
        lista = carregar_lista(lista_path) if os.path.isfile(lista_path) else []
    except (ValueError, OSError, zipfile.BadZipFile) as e:
        print(f"[ERRO] não consegui ler a lista {os.path.abspath(lista_path)}: {e}")
        return 1
    nome_automatico = not args.out or os.path.isdir(args.out) or args.out.endswith(("/", "\\"))
    destino = _destino(args.out, dia, selecao=bool(args.fundos), inicio=dias[0] if periodo else None)
    if not _pasta_gravavel(destino):
        return 1
    fundos, invalidas = resolver_fundos(args.fundos, lista)
    if not fundos:
        if args.fundos:
            print(f"[ERRO] nenhuma das entradas da linha de comando é CNPJ ou código ANBIMA: {', '.join(invalidas)}"
                  + (" (a data da análise vai em --data dd/mm/aaaa)" if any(re.fullmatch(r"\d{2}/\d{2}/\d{4}", x)
                                                                           for x in invalidas) else ""))
        else:
            print(f"[ERRO] nenhum fundo para consultar: a lista\n"
                  f"       {os.path.abspath(lista_path)}\n"
                  f"       está vazia ou não existe. Passe CNPJs/códigos ANBIMA na linha de comando ou preencha a lista\n"
                  f"       (colunas cnpj;codigo_anbima;nome — ex.: 00.807.777/0001-62;C0000019119;DAYCOVAL ...).")
        return 1
    if periodo:
        print(f"[INFO] {len(fundos)} fundo(s), período {dias[0]:%d/%m/%Y} a {dia:%d/%m/%Y} ({len(dias)} dia(s) útil(eis))")
    else:
        print(f"[INFO] {len(fundos)} fundo(s), data da análise {dia:%d/%m/%Y}")

    # Broadcast
    bc, bc_motivo = {}, ""
    nums = list(dict.fromkeys(f["num"] for f in fundos if f["num"] and f["prefixo"] == "C"))
    if args.sem_broadcast:
        bc_motivo = "não consultado (--sem-broadcast)"
    elif not nums:
        bc_motivo = "sem código ANBIMA"
    elif not terminal_no_ar():
        bc_motivo = "terminal Broadcast fechado"
        print("[AVISO] terminal Broadcast fechado — seguindo só com a CVM. "
              "Abra o terminal e rode de novo para ter os dados do Broadcast.")
    else:
        print(f"[INFO] consultando {len(nums)} fundo(s) no Broadcast...")
        try:
            bc = consultar_broadcast(nums, dias[0] - timedelta(days=JANELA_BCH_DIAS), dia, args.timeout)
        except Exception as e:  # noqa
            bc_motivo = f"erro ao consultar o Broadcast: {e!r}"
            print(f"[AVISO] {bc_motivo} — seguindo só com a CVM.")
        if bc and all("#BCONN" in str(r) for r in bc.values()):
            bc, bc_motivo = {}, "terminal Broadcast sem conexão (#BCONN) — aberto mas deslogado?"
            print(f"[AVISO] {bc_motivo} Faça login no terminal e rode de novo. Seguindo só com a CVM.")

    # CVM
    cvm_info = dict(series={}, presentes=set(), ultima_fundo={}, ultima_geral="", sem_d1=False)
    cvm_motivo, cadastro = "", cadastro_vazio()
    cnpjs = list(dict.fromkeys(f["cnpj"] for f in fundos if f["cnpj"]))
    if args.sem_cvm:
        cvm_motivo = "não consultada (--sem-cvm)"
    elif cnpjs:
        print("[INFO] lendo dados abertos da CVM (informe diário e cadastro)...")
        try:
            cvm_info = cvm_informe_diario(dia, cnpjs, inicio=dias[0])
        except Exception as e:  # noqa
            cvm_motivo = f"erro no informe diário da CVM: {e!r}"
            print(f"[AVISO] {cvm_motivo}")
        try:
            cadastro = cvm_cadastro(cnpjs)
        except Exception as e:  # noqa
            print(f"[AVISO] erro no cadastro da CVM (razão social): {e!r}")

    if not periodo:
        linhas = montar_linhas(fundos, dia, bc, bc_motivo, cvm_info, cvm_motivo, cadastro,
                               ultimo_disponivel=not args.so_data_exata)
        for ln in linhas:
            rotulo = ln["codigo_anbima"] or ln["cnpj"]
            fontes = sorted({fo.split(" (")[0] for fo in ln["fonte"].values() if fo})
            preenchidos = sum(_tem(v) for v in ln["valor"].values())
            msg = f"{rotulo}: {preenchidos}/{len(CAMPOS)} campos ({' + '.join(fontes) or 'sem dados'})"
            if ln["status_bc"] != "ok" and not bc_motivo:
                msg += f" | Broadcast: {ln['status_bc']}"
            print(("[OK] " if preenchidos else "[AVISO] ") + msg + (f" | {ln['obs']}" if ln["obs"] else ""))
    else:
        # período: cada dia útil é uma linha por fundo; dia sem dado fica em branco (o dia
        # anterior já é linha própria — não repete o último disponível)
        linhas = []
        for d in dias:
            do_dia = montar_linhas(fundos, d, bc, bc_motivo, cvm_info, cvm_motivo, cadastro, ultimo_disponivel=False)
            linhas += do_dia
            com_dado = sum(any(_tem(v) for v in ln["valor"].values()) for ln in do_dia)
            prov = sum(any("provisório" in fo for fo in ln["fonte"].values()) for ln in do_dia)
            print(f"[INFO] {d:%d/%m/%Y}: {com_dado}/{len(do_dia)} fundo(s) com dados"
                  + (f", {prov} com valor provisório" if prov else ""))
        for f in fundos:
            rotulo = codigo_anbima(f["prefixo"], f["num"]) or formatar_cnpj(f["cnpj"])
            sem = [ln["data"] for ln in linhas if (ln["codigo_anbima"] or ln["cnpj"]) == rotulo
                   and not any(_tem(v) for v in ln["valor"].values())]
            if sem:
                print(f"[AVISO] {rotulo}: sem dados em {len(sem)} dia(s) ({', '.join(f'{x:%d/%m}' for x in sem[:8])}"
                      f"{'...' if len(sem) > 8 else ''})")

    # não troca a foto do dia por uma com menos Broadcast (terminal fechado, --sem-broadcast,
    # falha em parte dos fundos): a nova vai p/ outro arquivo e a antiga fica
    if nome_automatico and os.path.isfile(destino):
        n_novo = sum(any(_tem(x) for x in ln["bc"].values()) for ln in linhas)
        n_atual = _n_fundos_broadcast(destino)
        if n_novo < n_atual:
            sufixo = "_so_cvm" if n_novo == 0 else f"_parcial_{datetime.now():%H%M%S}"
            destino = os.path.splitext(destino)[0] + sufixo + ".xlsx"
            print(f"[AVISO] a planilha já tem {n_atual} linha(s) com dados do Broadcast e esta rodada tem "
                  f"{n_novo}; para não perder a anterior, esta vai para {os.path.basename(destino)}")
    try:
        destino = salvar_xlsx(linhas, dia, destino, periodo=periodo)
    except OSError as e:
        base_reserva = os.path.splitext(_destino(None, dia, bool(args.fundos), dias[0] if periodo else None))[0]
        reserva = f"{base_reserva}_{datetime.now():%H%M%S}.xlsx"
        print(f"[ERRO] não consegui gravar {os.path.abspath(destino)} ({e}); salvando em {reserva}")
        try:
            destino = salvar_xlsx(linhas, dia, reserva, periodo=periodo)
        except OSError as e2:
            print(f"[ERRO] também falhou: {e2}")
            return 1
    print(f"[OK] salvo em: {os.path.abspath(destino)}")
    return 0 if any(any(_tem(v) for v in ln["valor"].values()) for ln in linhas) else 1


if __name__ == "__main__":
    raise SystemExit(main())
