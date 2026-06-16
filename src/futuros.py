from __future__ import annotations

import os
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
import matplotlib.image as mpimg
from mpl_toolkits.axes_grid1.inset_locator import inset_axes
from openpyxl.styles import PatternFill, Font
import openpyxl
from datetime import date
import QuantLib as ql
import sql_connection as sql


# ================= CONFIG =================

BASE_OUTPUT   = "../output/rolagem"
INPUT_FUNDOS  = "../input/fundos.txt"
INPUT_ROLAGEM = "../input/fundos_rolagem.xlsx"
LOGO_PATH     = "../input/logo.png"

ITAU_DARK   = "#000D3C"
ITAU_ORANGE = "#EC7000"
ITAU_WHITE  = "#ffffff"
ITAU_GRAY   = "#e0e0e0"
ITAU_GRAY2  = "#b3b3b3"
ITAU_RED    = "#c0392b"
ITAU_GREEN  = "#1e7e34"
ROW_PAR     = "#eaf0fb"

MES_CODIGO = {
    1: "F", 2: "G", 3: "H", 4: "J",
    5: "K", 6: "M", 7: "N", 8: "Q",
    9: "U", 10: "V", 11: "X", 12: "Z",
}

GRUPO_DOL = {"DOL", "WDO"}

GRUPOS_EXCLUIR = {
    "DI1",
}


# ================= UTIL =================


def output_dir_for_date(dt):
    return os.path.join(BASE_OUTPUT, pd.to_datetime(dt).strftime("%Y%m%d"))


def formatar_vencimento(dt):
    dt = pd.to_datetime(dt)
    return f"{MES_CODIGO[dt.month]}{str(dt.year)[-2:]}"


def normalizar_ativo(a):
    return str(a).replace("FUT ", "").strip()


def eh_grupo_dol(a):
    return normalizar_ativo(a) in GRUPO_DOL


def ler_fundos():
    if not os.path.exists(INPUT_FUNDOS):
        return []
    with open(INPUT_FUNDOS) as f:
        return [l.strip() for l in f if l.strip()]


# ================= SQL =================


def consulta_futuro(dt, fundos):
    filtro = (
        f"AND clcli_cd IN ({','.join([repr(f) for f in fundos])})"
        if fundos else ""
    )

    q_fut = f"""
        SELECT clcli_cd,
               ISNULL(qt,0)*ISNULL(vl_varponto,0)*ISNULL(vl_precomerc,0) AS Notional,
               dt_vencto,
               fuvct_fuatv_cd
        FROM dbsx337.dbo.view_vcrapfut WITH(NOLOCK)
        WHERE dt='{dt}' {filtro}
    """

    q_pl = f"""
        SELECT clcli_cd, vl_patrliqtot1 AS PL
        FROM dbsx337.dbo.view_vcraplsi WITH(NOLOCK)
        WHERE dt='{dt}' {filtro}
    """

    db = sql.get_connection()
    return db.run_query(q_fut), db.run_query(q_pl)


# ================= CORE =================


def processar_grupo(df_fut, df_pl, grupo):
    if grupo == "DOL_WDO":
        df = df_fut[df_fut["fuvct_fuatv_cd"].apply(eh_grupo_dol)].copy()
        label = "FUT DOL + WDO"
    else:
        df = df_fut[df_fut["fuvct_fuatv_cd"].apply(normalizar_ativo) == grupo].copy()
        label = f"FUT {grupo}"

    if df.empty:
        return pd.DataFrame(), [], label

    vencimentos = sorted(df["vencimento"].dropna().unique())

    if not vencimentos:
        return pd.DataFrame(), [], label

    v1 = vencimentos[0]
    v2 = None

    for v in vencimentos:
        if v > v1:
            v2 = v
            break

    if v2 is None and len(vencimentos) >= 2:
        v2 = v1
        v1 = vencimentos[-2]

    if v2 is None:
        return pd.DataFrame(), [v1], label

    print(f"    [{grupo}] {v1} >> {v2}")

    df = df[df["vencimento"].isin([v1, v2])]

    df_grp = df.groupby(["clcli_cd", "vencimento"])["Notional"].sum().reset_index()

    total = df_grp.groupby("clcli_cd")["Notional"].sum().reset_index()
    total.rename(columns={"Notional": "total_fut"}, inplace=True)

    df_tab = df_pl.merge(total, on="clcli_cd", how="inner")

    n1 = df_grp[df_grp.vencimento == v1].groupby("clcli_cd")["Notional"].sum().reset_index()
    n2 = df_grp[df_grp.vencimento == v2].groupby("clcli_cd")["Notional"].sum().reset_index()

    n1.rename(columns={"Notional": "not_v1"}, inplace=True)
    n2.rename(columns={"Notional": "not_v2"}, inplace=True)

    df_tab = df_tab.merge(n1, on="clcli_cd", how="left")
    df_tab = df_tab.merge(n2, on="clcli_cd", how="left")

    df_tab = df_tab.fillna(0)
    df_tab = df_tab[(df_tab["not_v1"] != 0) | (df_tab["not_v2"] != 0)]

    if df_tab.empty:
        return pd.DataFrame(), [], label

    df_tab["FUTURO"] = label

    df_tab[f"Pos {v1} % PL"] = df_tab["not_v1"] / df_tab["PL"]
    df_tab[f"Pos {v2} % PL"] = df_tab["not_v2"] / df_tab["PL"]

    df_tab[f"{v1} % FUT"] = df_tab["not_v1"] / df_tab["total_fut"]
    df_tab[f"{v2} % FUT"] = df_tab["not_v2"] / df_tab["total_fut"]

    df_out = df_tab[[
        "clcli_cd", "PL", "FUTURO",
        f"Pos {v1} % PL",
        f"Pos {v2} % PL",
        f"{v1} % FUT",
        f"{v2} % FUT",
    ]].rename(columns={"clcli_cd": "FUNDO"})

    return df_out, [v1, v2], label


# ================= EXCEL (DADOS) =================


def gerar_excel_dados(resultados, dt, outdir):
    caminho = os.path.join(outdir, f"rolagem_{dt}.xlsx")

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        for grupo, df in resultados:
            nome_aba = grupo[:31]
            df_copy = df.copy().sort_values("FUNDO")

            df_copy.to_excel(writer, sheet_name=nome_aba, index=False)

            ws = writer.sheets[nome_aba]

            for cell in ws[1]:
                cell.fill = openpyxl.styles.PatternFill("solid", fgColor="000D3C")
                cell.font = openpyxl.styles.Font(color="FFFFFF", bold=True)

            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    col_name = ws.cell(row=1, column=cell.column).value

                    if cell.row % 2 == 0:
                        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="EAF0FB")

                    if col_name == "PL":
                        cell.number_format = '"R$"#,##0.00'

                    if "%" in col_name:
                        cell.number_format = "0.00%"
                        try:
                            if cell.value is not None:
                                if cell.value < 0:
                                    cell.font = openpyxl.styles.Font(color="C0392B", bold=True)
                                elif cell.value > 0:
                                    cell.font = openpyxl.styles.Font(color="1E7E34", bold=True)
                        except Exception:
                            pass

            for col in ws.columns:
                max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
                ws.column_dimensions[col[0].column_letter].width = max_len + 2

    print(f"[OK] Excel formatado: {caminho}")
    return caminho


# ================= IMAGEM =================


def gerar_imagem(df, grupo, label, vcts, dt, outdir, mostrar_logo=False, mostrar_titulo=False):
    os.makedirs(outdir, exist_ok=True)

    df_disp = df.copy()
    df_disp["PL"] = df_disp["PL"].map(lambda x: f"R$ {x:,.2f}")

    for c in df_disp.columns:
        if "%" in c:
            df_disp[c] = df_disp[c].map(lambda x: f"{x:.2%}")

    colunas = df_disp.columns.tolist()

    fig = plt.figure(figsize=(25, 3 + len(df) * 0.6))
    gs = gridspec.GridSpec(4, 1, height_ratios=[1, 1.5, len(df) * 0.6, 0.3])

    # HEADER
    ax_b = fig.add_subplot(gs[0])
    ax_b.set_facecolor(ITAU_DARK)
    ax_b.axis("off")

    if mostrar_logo and os.path.exists(LOGO_PATH):
        ax_logo = inset_axes(ax_b, width="10%", height="80%", loc="center right")
        ax_logo.imshow(mpimg.imread(LOGO_PATH))
        ax_logo.axis("off")

        ax_b.text(
            0.85, 0.5, "Itaú Asset",
            fontsize=18,
            fontweight="bold",
            color=ITAU_GRAY2,
            va="center",
            ha="left",
        )

    ax_b.axhline(y=0.05, color=ITAU_ORANGE)

    # TITLE
    ax = fig.add_subplot(gs[1])
    ax.axis("off")

    dt_fmt = pd.to_datetime(dt).strftime("%d/%m/%Y")
    titulo_venc = f"{vcts[0]} >> {vcts[1]}"

    if mostrar_titulo:
        ax.text(0.02, 0.8, "Relatorio de Rolagem", fontsize=24, weight="bold")
    ax.text(0.02, 0.3, label, color=ITAU_ORANGE, fontsize=18)
    ax.text(0.98, 0.8, f"Data: {dt_fmt}", ha="right")
    ax.text(0.98, 0.3, titulo_venc, ha="right")

    # TABLE
    ax = fig.add_subplot(gs[2])
    ax.axis("off")

    tbl = ax.table(cellText=df_disp.values, colLabels=colunas, loc="center", cellLoc="center")
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(11)
    tbl.scale(1, 2)

    for (r, c), cell in tbl.get_celld().items():
        if r == 0:
            cell.set_facecolor(ITAU_DARK)
            cell.set_text_props(color="white", weight="bold")
        else:
            bg = ROW_PAR if r % 2 == 0 else ITAU_WHITE
            cell.set_facecolor(bg)

            if "%" in colunas[c]:
                val = df.iloc[r - 1][colunas[c]]
                if val < 0:
                    cell.set_text_props(color=ITAU_RED, weight="bold")
                elif val > 0:
                    cell.set_text_props(color=ITAU_GREEN, weight="bold")

    nome = os.path.join(outdir, f"rolagem_{grupo}_{dt}.png")
    plt.savefig(nome, dpi=150, bbox_inches="tight")
    plt.close()


# ================= MAIN =================


def main():
    dt = ql.Brazil().advance(
        ql.Date.from_date(date.today()),
        ql.Period('-1d')
    ).to_date().strftime('%Y-%m-%d')

    print(dt)

    out_dir = output_dir_for_date(dt)
    fundos = ler_fundos()

    df_fut, df_pl = consulta_futuro(dt, fundos)

    df_fut["dt_vencto"] = pd.to_datetime(df_fut["dt_vencto"])
    df_fut["vencimento"] = df_fut["dt_vencto"].apply(formatar_vencimento)

    ano_ref = pd.to_datetime(dt).strftime("%y")
    df_fut = df_fut[df_fut["vencimento"].str.endswith(ano_ref)]

    ativos = df_fut["fuvct_fuatv_cd"].apply(normalizar_ativo).unique()
    grupos_set = set(["DOL_WDO"] + [a for a in ativos if a not in GRUPO_DOL]) - GRUPOS_EXCLUIR

    try:
        abas = pd.ExcelFile(INPUT_ROLAGEM).sheet_names
        grupos = [a for a in abas if a in grupos_set]
        grupos += sorted([g for g in grupos_set if g not in abas])
    except Exception:
        grupos = ["DOL_WDO"] + sorted([a for a in ativos if a not in GRUPO_DOL])

    # DOL_WDO sempre primeiro para garantir que o titulo "Relatorio de Rolagem" apareca no inicio
    if "DOL_WDO" in grupos:
        grupos.remove("DOL_WDO")
        grupos.insert(0, "DOL_WDO")

    resultados = []
    primeiro = True

    for g in grupos:
        df_res, vcts, label = processar_grupo(df_fut, df_pl, g)

        if df_res.empty:
            continue

        resultados.append((g, df_res))

        gerar_imagem(df_res, g, label, vcts, dt, out_dir,
                     mostrar_logo=primeiro, mostrar_titulo=primeiro)

        primeiro = False

    gerar_excel_dados(resultados, dt, out_dir)


if __name__ == "__main__":
    main()
