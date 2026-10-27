# ─────────────────────────────────────────────
# taxa_media/__main__.py
# Taxa média de aquisição dos CRIs de um fundo (padrão ITAIMOBCDIFII): taxa de cada lote em aberto na
# custódia, média por ativo (ponderada pela quantidade) e da carteira (ponderada pelo financeiro atual).
#
# Fontes (DBSX337 pelo sql_connection.py desta pasta, credenciais SQL_LOGIN e SQL_PASSWORD; todas as
# leituras WITH (NOLOCK)):
#   VCRAOPRF              custódia de renda fixa, uma linha por movimento de cada lote (CD). Só entra
#                         RFTP_CD = 'CRI-I': LTN-O, NTN-O e LFT O são compromissadas over; CRIOP e CRILA,
#                         compromissadas com lastro em CRI. Trades = SG_OPERACAO C/V (M, J e I são eventos).
#   FUNDOS_OPERACOES      boleta do front: exec_taxa = taxa negociada (CDI+: spread % a.a.; %CDI: % do CDI).
#   VW_FUNDOS_POSICAO     quantidade na data (conferência) e PU da última data com preço (total/quantidade).
#   FUNDOS_DADOS_INDICES  CDIAcumulado: CDI do dia, só para converter %CDI em CDI+ equivalente.
#
# Cálculo:
#   lote      qtd em aberto = compras - vendas do mesmo CD. Lote com sufixo A (ex.: L10947A) é relançamento
#             da VCRA (venda do original + compra um dia depois no mesmo PU) e já sai certo nessa conta.
#   boleta    cada compra casa com a boleta do mesmo ativo em até 5 dias, pela quantidade ou pelo PU (há
#             boletas com o dobro da quantidade da custódia); prefere a quantidade e a data mais próxima.
#   grupo     CDI+ se o ativo tem spread de emissão (VL_JUROS_EMISSAO > 0); senão %CDI, com a emissão no
#             IDDIR_PC. O indexador da VW_FUNDOS_POSICAO classifica errado alguns papéis.
#   taxa      %CDI: só a boleta — a VCRA grava %CDI como IDDIR_PC + spread, e isso não converte de forma
#             consistente para o % do CDI da boleta. CDI+: a boleta (--fonte boleta, padrão) ou a VCRA, que
#             sai do PU pago (--fonte vcra); sem boleta válida, a VCRA.
#             Boleta válida: %CDI >= 50 (% do CDI), CDI+ < 50 (spread). Lote sem taxa fica fora da média do
#             ativo; input/correcoes.csv (CD;taxa;motivo) corrige a taxa de lotes específicos.
#   ativo     média das taxas dos lotes, ponderada pela quantidade em aberto.
#   carteira  média ponderada pelo financeiro (qtd x último PU), por grupo e no total. O total junta os grupos
#             em CDI+ equivalente, com CDI flat: spread = (1 + %CDI x CDI diário)^252 / (1 + CDI) - 1.
#
# Saída: output/taxa_media_<fundo>_<aaaammdd>.xlsx — abas Carteira, Ativos, Lotes, Avisos e Parâmetros.
#
# Requisitos: ODBC Driver 18 for SQL Server, pyodbc, sqlalchemy, pandas, openpyxl.
#
# Uso (de dentro de taxa_media/, ou "python taxa_media" a partir de itau_asset/):
#   python __main__.py                          (ITAIMOBCDIFII na última data da posição)
#   python __main__.py --data 23/10/2026
#   python __main__.py --fonte vcra             (taxa dos CDI+ pela VCRA)
#   python __main__.py --fundo OUTROFUNDO --cdi 13.65
# ─────────────────────────────────────────────

import argparse
import os
import re
import sys
from datetime import datetime

import numpy as np
import pandas as pd

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA_DIR = os.path.join(AQUI, "output")
CORRECOES_PADRAO = os.path.join(AQUI, "input", "correcoes.csv")

FUNDO_PADRAO = "ITAIMOBCDIFII"
TIPO_CRI = "CRI-I"      # RFTP_CD dos CRIs na VCRAOPRF
JANELA_BOLETA = 5       # dias entre a compra na VCRA e a data da boleta
TOLERANCIA_PU = 1e-4    # PU da boleta x PU da VCRA, para casar quando a quantidade não bate
LIMITE_CONVENCAO = 50   # boleta de %CDI vem em % do CDI (>= 50); de CDI+, em spread % a.a. (< 50)
DIAS_PU = 30            # até quantos dias antes da data procurar o último PU com preço na posição
# boleta x VCRA na convenção do grupo acima disso vira aviso: CDI+ em p.p. de spread; %CDI em pontos de % do
# CDI (convertida com CDI flat, a VCRA costuma ficar até ~1,5 ponto abaixo da boleta)
DIVERGENCIA = {"CDI+": 0.25, "%CDI": 3.0}
# ordem dos avisos no console e na aba Avisos ("Correção manual" sai como [INFO])
ORDEM_AVISOS = ["Quantidade", "Sem lote", "Fora da média", "Média parcial", "Boleta inválida", "Sem boleta",
                "Correção manual", "Correção não aplicada", "Boleta × VCRA", "Venda sem compra", "Lote negativo"]


# ── SQL ──────────────────────────────────────────────────────────────────────
SQL_ULTIMA_DATA = """
SELECT MAX(date) AS data
FROM   VW_FUNDOS_POSICAO WITH (NOLOCK)
WHERE  codigo_IAM_fundo = '{fundo}'
"""

SQL_VCRA = """
SELECT  CD, CD_CETIP_SELIC AS codigo, ISIN AS isin, DT AS data, SG_OPERACAO AS op, QT AS qt,
        VL_PU_OPERACAO AS pu, IDDIR_PC AS perc_cdi, VL_JUROS_EMISSAO AS taxa_emissao,
        VL_JUROSOPER AS taxa_vcra
FROM    VCRAOPRF WITH (NOLOCK)
WHERE   CLCLI_CD = '{fundo}' AND DT <= '{data}'
  AND   RFTP_CD = '{tipo}' AND SG_OPERACAO IN ('C', 'V')
"""

SQL_BOLETAS = """
SELECT  codigo_IAM_ativo AS codigo, date_operacao AS data, exec_qt AS qt, exec_preco AS pu,
        exec_taxa AS taxa
FROM    FUNDOS_OPERACOES WITH (NOLOCK)
WHERE   codigo_IAM_fundo = '{fundo}'
  AND   date_operacao <= DATEADD(day, {janela}, '{data}')
  AND   codigo_IAM_ativo IN (SELECT CD_CETIP_SELIC FROM VCRAOPRF WITH (NOLOCK)
                             WHERE CLCLI_CD = '{fundo}' AND RFTP_CD = '{tipo}')
"""

SQL_POSICAO = """
SELECT  codigo_IAM_ativo AS codigo, isin, date AS data, quantidade, total
FROM    VW_FUNDOS_POSICAO WITH (NOLOCK)
WHERE   codigo_IAM_fundo = '{fundo}' AND tipo = 'Fixed Income'
  AND   date BETWEEN DATEADD(day, -{dias}, '{data}') AND '{data}'
"""

SQL_CDI = """
SELECT TOP 1 date AS data, pnl_perc
FROM   FUNDOS_DADOS_INDICES WITH (NOLOCK)
WHERE  codigo_IAM_indice = 'CDIAcumulado' AND date <= '{data}' AND pnl_perc > 0
ORDER BY date DESC
"""


def _consulta(query, sql, nome, obrigatoria=True):
    """SELECT pelo sql_connection.query, que devolve DataFrame vazio também quando a consulta dá erro
    (e imprime a mensagem)."""
    df = query(sql)
    if df.empty and obrigatoria:
        raise RuntimeError(f"{nome}: nenhuma linha (ou erro na consulta: veja a mensagem acima)")
    return df


def carregar_dados(fundo, data, cdi):
    """VCRA, boletas, posição e CDI do banco. data None = última data da posição do fundo; cdi (% a.a.)
    None = CDIAcumulado da FUNDOS_DADOS_INDICES. Devolve dict com os DataFrames, a data e o CDI (fração
    a.a.) com a origem."""
    sys.path.insert(0, AQUI)
    from sql_connection import disconnect, query

    try:
        if data is None:
            ultima = _consulta(query, SQL_ULTIMA_DATA.format(fundo=fundo), "VW_FUNDOS_POSICAO")["data"].iloc[0]
            if pd.isna(ultima):
                raise RuntimeError(f"fundo {fundo} sem posição na VW_FUNDOS_POSICAO")
            data = pd.Timestamp(ultima).date()
        p = {"fundo": fundo, "data": f"{data:%Y-%m-%d}", "tipo": TIPO_CRI, "janela": JANELA_BOLETA, "dias": DIAS_PU}
        vcra = _consulta(query, SQL_VCRA.format(**p), f"VCRAOPRF ({fundo}, {TIPO_CRI}, até {data:%d/%m/%Y})")
        boletas = _consulta(query, SQL_BOLETAS.format(**p), "FUNDOS_OPERACOES", obrigatoria=False)
        posicao = _consulta(query, SQL_POSICAO.format(**p), f"VW_FUNDOS_POSICAO ({fundo})")
        if cdi is None:
            linha = _consulta(query, SQL_CDI.format(**p), "FUNDOS_DADOS_INDICES (CDIAcumulado)").iloc[0]
            cdi = (1 + float(linha["pnl_perc"])) ** 252 - 1
            origem = f"CDIAcumulado da FUNDOS_DADOS_INDICES em {pd.Timestamp(linha['data']):%d/%m/%Y}"
        else:
            cdi, origem = cdi / 100, "informado em --cdi"
    finally:
        disconnect()

    if not 0.01 < cdi < 0.5:
        raise RuntimeError(f"CDI de {cdi:.2%} a.a. ({origem}) fora do esperado: informe com --cdi")
    if not (pd.to_datetime(posicao["data"]).dt.normalize() == pd.Timestamp(data)).any():
        raise RuntimeError(f"fundo {fundo} sem posição em {data:%d/%m/%Y} na VW_FUNDOS_POSICAO")
    if boletas.empty:
        print("[AVISO] nenhuma boleta na FUNDOS_OPERACOES: CDI+ pela VCRA e %CDI sem taxa")
    return {"vcra": vcra, "boletas": boletas, "posicao": posicao, "data": data, "cdi": cdi, "cdi_origem": origem}


# ── Correções manuais ────────────────────────────────────────────────────────
def carregar_correcoes(caminho):
    """CSV CD;taxa;motivo -> {CD: (taxa, motivo)}. Linha com # é comentário e o cabeçalho é ignorado; a taxa
    vai na convenção do grupo do ativo (CDI+: spread % a.a.; %CDI: % do CDI). Sem arquivo, sem correções."""
    if not os.path.exists(caminho):
        return {}
    try:
        with open(caminho, encoding="utf-8-sig") as f:
            texto = f.read()
    except UnicodeDecodeError:
        with open(caminho, encoding="cp1252") as f:
            texto = f.read()
    correcoes = {}
    for n, linha in enumerate(texto.splitlines(), start=1):
        campos = [c.strip() for c in linha.split(";")]
        if not linha.strip() or linha.lstrip().startswith("#") or campos[0].upper() == "CD":
            continue
        try:
            taxa = float(campos[1].replace(".", "").replace(",", ".") if "," in campos[1] else campos[1])
        except (IndexError, ValueError):
            raise ValueError(f"{caminho}, linha {n}: esperava CD;taxa;motivo, veio {linha.strip()!r}") from None
        correcoes[campos[0]] = (taxa, campos[2] if len(campos) > 2 else "")
    return correcoes


# ── Cálculo ──────────────────────────────────────────────────────────────────
def _br(x, casas=2):
    """Número no formato brasileiro (1.234,56)."""
    return f"{x:,.{casas}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _cdi_mais(x):
    return f"CDI{'+' if x >= 0 else '−'}{_br(abs(x))}"


def _aviso(avisos, tipo, ativo, detalhe, lote=""):
    avisos.append({"tipo": tipo, "ativo": ativo, "lote": lote, "detalhe": detalhe})


def _normalizar(df, textos=(), datas=(), numeros=()):
    df = df.copy()
    for c in textos:
        df[c] = df[c].str.strip()
    for c in datas:
        df[c] = pd.to_datetime(df[c]).dt.normalize()
    for c in numeros:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def _media_ponderada(df, chave, valor, peso):
    """Média de valor ponderada por peso em cada grupo de chave, só com as linhas que têm valor."""
    ok = df[valor].notna()
    s = df.assign(_v=(df[valor] * df[peso]).where(ok, 0.0), _p=df[peso].where(ok, 0.0)).groupby(chave)[["_v", "_p"]].sum()
    return s["_v"] / s["_p"].replace(0, np.nan)


def casar_boletas(compras, boletas):
    """Acrescenta a cada compra da VCRA a boleta do mesmo ativo em até JANELA_BOLETA dias com a mesma
    quantidade ou o mesmo PU; prefere a quantidade e, depois, a data mais próxima."""
    novas = ["data_boleta", "qt_boleta", "taxa_boleta", "casou_por"]
    if boletas.empty:
        return compras.assign(**{c: np.nan for c in novas})
    b = boletas.rename(columns={"data": "data_boleta", "qt": "qt_boleta", "pu": "pu_boleta", "taxa": "taxa_boleta"})
    m = compras[["codigo", "data", "qt", "pu"]].reset_index().merge(b, on="codigo")
    mesma_qt = np.isclose(m["qt_boleta"].abs(), m["qt"], rtol=0, atol=1e-6)
    mesmo_pu = ((m["pu_boleta"] - m["pu"]).abs() < TOLERANCIA_PU).to_numpy()
    m["dias"] = (m["data_boleta"] - m["data"]).dt.days.abs()
    m["prioridade"] = np.where(mesma_qt, 0, 1)
    m["casou_por"] = np.where(mesma_qt, "quantidade", "PU")
    m = m[(m["dias"].to_numpy() <= JANELA_BOLETA) & (mesma_qt | mesmo_pu)]
    m = m.sort_values(["index", "prioridade", "dias"]).drop_duplicates("index").set_index("index")
    return compras.join(m[novas])


def montar_lotes(vcra, boletas, fonte, correcoes, cdi_d, avisos):
    """Lotes em aberto da VCRA (compras - vendas do mesmo CD) com a boleta casada, o grupo e a taxa usada."""
    compras = casar_boletas(vcra[vcra["op"] == "C"].reset_index(drop=True), boletas)
    vendas = vcra[vcra["op"] == "V"].groupby("CD")["qt"].sum()

    chave = ["codigo", "CD"]
    lotes = compras.groupby(chave).agg(
        isin=("isin", "first"), dt_compra=("data", "min"), perc_cdi=("perc_cdi", "max"),
        taxa_emissao=("taxa_emissao", "max"), qt_comprada=("qt", "sum"), data_boleta=("data_boleta", "min"),
        casou_por=("casou_por", "first"))
    for coluna, nome in (("pu", "pu_compra"), ("taxa_vcra", "taxa_vcra"), ("taxa_boleta", "taxa_boleta")):
        lotes[nome] = _media_ponderada(compras, chave, coluna, "qt")
    lotes = lotes.reset_index()
    lotes["qt_vendida"] = lotes["CD"].map(vendas).fillna(0.0)
    lotes["qt_aberta"] = lotes["qt_comprada"] - lotes["qt_vendida"]

    for cd in sorted(set(vendas.index) - set(lotes["CD"])):
        v = vcra[(vcra["op"] == "V") & (vcra["CD"] == cd)]
        _aviso(avisos, "Venda sem compra", v["codigo"].iloc[0],
               f"venda de {_br(v['qt'].sum(), 0)} sem a compra do lote no histórico da VCRA", cd)
    for r in lotes[lotes["qt_aberta"] < 0].itertuples():
        _aviso(avisos, "Lote negativo", r.codigo,
               f"vendas ({_br(r.qt_vendida, 0)}) maiores que as compras ({_br(r.qt_comprada, 0)})", r.CD)

    # CDI+ tem spread de emissão; %CDI tem a emissão no IDDIR_PC (o mesmo grupo em todos os lotes do ativo)
    lotes["grupo"] = np.where(lotes.groupby("codigo")["taxa_emissao"].transform("max") > 0, "CDI+", "%CDI")
    cdi_mais = lotes["grupo"] == "CDI+"
    lotes["boleta_ok"] = np.where(cdi_mais, lotes["taxa_boleta"] < LIMITE_CONVENCAO,
                                  lotes["taxa_boleta"] >= LIMITE_CONVENCAO)
    usa_boleta = lotes["boleta_ok"] & (~cdi_mais | (fonte == "boleta"))
    lotes["taxa"] = np.where(usa_boleta, lotes["taxa_boleta"], np.where(cdi_mais, lotes["taxa_vcra"], np.nan))
    lotes["fonte_taxa"] = np.where(usa_boleta, "boleta", np.where(cdi_mais, "VCRA", ""))

    # VCRA na convenção do grupo, para comparar com a boleta: no %CDI, o % do CDI que rende o mesmo que
    # IDDIR_PC% do CDI + o spread da VCRA, com CDI flat
    p, s = lotes["perc_cdi"] / 100, lotes["taxa_vcra"] / 100
    lotes["taxa_vcra_conv"] = np.where(cdi_mais, lotes["taxa_vcra"],
                                       ((1 + p * cdi_d) * (1 + s) ** (1 / 252) - 1) / cdi_d * 100)
    lotes["dif_boleta_vcra"] = (lotes["taxa_boleta"] - lotes["taxa_vcra_conv"]).where(lotes["boleta_ok"])

    for cd, (taxa, motivo) in correcoes.items():
        sel = (lotes["CD"] == cd) & (lotes["qt_aberta"] > 0)
        if not sel.any():
            _aviso(avisos, "Correção não aplicada", "", "o lote não está entre os lotes em aberto", cd)
            continue
        lotes.loc[sel, "taxa"] = taxa
        lotes.loc[sel, "fonte_taxa"] = "manual"

    abertos = lotes[lotes["qt_aberta"] > 0].reset_index(drop=True)
    consequencia = {"boleta": "", "VCRA": "usada a taxa da VCRA", "manual": "usada a correção manual",
                    "": "lote sem taxa, fora da média do ativo"}
    for r in abertos.itertuples():
        if pd.isna(r.taxa_boleta):
            _aviso(avisos, "Sem boleta", r.codigo, f"nenhuma boleta casada: {consequencia[r.fonte_taxa]}", r.CD)
        elif not r.boleta_ok:
            _aviso(avisos, "Boleta inválida", r.codigo, f"boleta com {_br(r.taxa_boleta, 4)} fora da convenção "
                   f"do {r.grupo}: {consequencia[r.fonte_taxa]}", r.CD)
        elif abs(r.dif_boleta_vcra) > DIVERGENCIA[r.grupo]:
            _aviso(avisos, "Boleta × VCRA", r.codigo, f"boleta {_br(r.taxa_boleta)} × VCRA {_br(r.taxa_vcra_conv)} "
                   f"(PU {_br(r.pu_compra)}); vale a {'correção manual' if r.fonte_taxa == 'manual' else r.fonte_taxa}",
                   r.CD)
        if r.fonte_taxa == "manual":
            _aviso(avisos, "Correção manual", r.codigo, f"taxa {_br(r.taxa, 4)} ({correcoes[r.CD][1]})", r.CD)
    return abertos


def montar_ativos(lotes, posicao, data, cdi, cdi_d, avisos):
    """Taxa média por ativo (lotes ponderados pela quantidade), conferência com a quantidade da posição e
    financeiro pelo último PU com preço."""
    l = lotes.assign(sem_boleta=~lotes["boleta_ok"], manual=lotes["fonte_taxa"] == "manual",
                     sem_taxa=lotes["taxa"].isna())
    a = l.groupby("codigo").agg(
        isin=("isin", "first"), grupo=("grupo", "first"), perc_cdi=("perc_cdi", "max"),
        taxa_emissao=("taxa_emissao", "max"), n_lotes=("CD", "count"), n_sem_boleta=("sem_boleta", "sum"),
        n_manual=("manual", "sum"), n_sem_taxa=("sem_taxa", "sum"), qt_lotes=("qt_aberta", "sum"))
    a["taxa_vcra"] = _media_ponderada(l, "codigo", "taxa_vcra_conv", "qt_aberta")
    a["taxa_media"] = _media_ponderada(l, "codigo", "taxa", "qt_aberta")

    ref = pd.Timestamp(data)
    na_data = posicao[(posicao["data"] == ref) & posicao["isin"].notna()]
    a = a.join(na_data.groupby("codigo")["quantidade"].sum().rename("qt_posicao"), how="outer")
    a["isin"] = a["isin"].fillna(na_data.groupby("codigo")["isin"].first())
    precos = (posicao[posicao["total"].notna()].groupby(["codigo", "data"])[["quantidade", "total"]].sum()
              .reset_index().sort_values("data").groupby("codigo").tail(1).set_index("codigo"))
    a["pu"] = (precos["total"] / precos["quantidade"].replace(0, np.nan)).reindex(a.index)
    a["dt_pu"] = precos["data"].reindex(a.index)
    a["dif_qt"] = a["qt_lotes"].fillna(0) - a["qt_posicao"].fillna(0)
    a["financeiro"] = a["qt_lotes"] * a["pu"]
    q = a["taxa_media"] / 100
    a["taxa_cdi_mais"] = np.where(a["grupo"] == "CDI+", a["taxa_media"],
                                  ((1 + q * cdi_d) ** 252 / (1 + cdi) - 1) * 100)
    a = a.rename_axis("codigo").reset_index().sort_values(["grupo", "codigo"], na_position="last")

    for r in a.itertuples():
        if pd.isna(r.qt_lotes):
            _aviso(avisos, "Sem lote", r.codigo,
                   f"{_br(r.qt_posicao, 0)} na posição sem lote {TIPO_CRI} em aberto na VCRA")
            continue
        if abs(r.dif_qt) > 1e-6:
            pos = _br(r.qt_posicao, 0) if pd.notna(r.qt_posicao) else "nada"
            _aviso(avisos, "Quantidade", r.codigo, f"lotes {_br(r.qt_lotes, 0)} × posição {pos} em {ref:%d/%m/%Y}")
        if pd.isna(r.taxa_media):
            _aviso(avisos, "Fora da média", r.codigo, "nenhum lote com taxa")
        elif r.n_sem_taxa:
            _aviso(avisos, "Média parcial", r.codigo,
                   f"{int(r.n_sem_taxa)} de {int(r.n_lotes)} lotes sem taxa: média só dos outros")
        if pd.isna(r.pu):
            _aviso(avisos, "Fora da média", r.codigo, f"sem PU na posição nos {DIAS_PU} dias até {ref:%d/%m/%Y}")
    return a.reset_index(drop=True)


def _resumo(grupo, g):
    fin = g["financeiro"].sum()
    return {"grupo": grupo, "n_ativos": len(g), "financeiro": fin,
            "taxa_media": (g["financeiro"] * g["taxa_media"]).sum() / fin,
            "taxa_cdi_mais": (g["financeiro"] * g["taxa_cdi_mais"]).sum() / fin}


def consolidar(ativos):
    """Média por grupo (na unidade do grupo) e da carteira (em CDI+ equivalente), ponderadas pelo financeiro."""
    ok = ativos[ativos["financeiro"].notna() & ativos["taxa_media"].notna()]
    linhas = [_resumo(grupo, g) for grupo, g in ok.groupby("grupo")]
    if len(ok):
        linhas.append({**_resumo("CARTEIRA", ok), "taxa_media": np.nan})
    return pd.DataFrame(linhas, columns=["grupo", "n_ativos", "financeiro", "taxa_media", "taxa_cdi_mais"])


def calcular(vcra, boletas, posicao, data, cdi, fonte, correcoes):
    """-> (lotes em aberto, ativos, carteira, avisos). cdi = CDI a.a. em fração; fonte = 'boleta' ou 'vcra'."""
    vcra = _normalizar(vcra, textos=("CD", "codigo", "op", "isin"), datas=("data",),
                       numeros=("qt", "pu", "perc_cdi", "taxa_emissao", "taxa_vcra"))
    vcra["qt"] = vcra["qt"].abs()
    if not boletas.empty:
        boletas = _normalizar(boletas, textos=("codigo",), datas=("data",), numeros=("qt", "pu", "taxa"))
    posicao = _normalizar(posicao, textos=("codigo", "isin"), datas=("data",), numeros=("quantidade", "total"))

    avisos = []
    cdi_d = (1 + cdi) ** (1 / 252) - 1
    lotes = montar_lotes(vcra, boletas, fonte, correcoes, cdi_d, avisos)
    ativos = montar_ativos(lotes, posicao, data, cdi, cdi_d, avisos)
    avisos.sort(key=lambda av: ORDEM_AVISOS.index(av["tipo"]))
    return lotes, ativos, consolidar(ativos), avisos


# ── Saída ────────────────────────────────────────────────────────────────────
TAXA, PU, QTD, VALOR, DATA = "0.0000", "#,##0.000000", "#,##0", "#,##0.00", "dd/mm/yyyy"
COLUNAS = {
    "Carteira": [("grupo", "Grupo", None), ("n_ativos", "Ativos", "0"), ("financeiro", "Financeiro (R$)", VALOR),
                 ("taxa_media", "Taxa média (CDI+: spread; %CDI: % do CDI)", TAXA),
                 ("taxa_cdi_mais", "Taxa média em CDI+ equivalente", TAXA)],
    "Ativos": [("codigo", "Ativo", None), ("isin", "ISIN", None), ("grupo", "Grupo", None),
               ("perc_cdi", "% do CDI (emissão)", TAXA), ("taxa_emissao", "Spread de emissão", TAXA),
               ("n_lotes", "Lotes", "0"), ("n_sem_boleta", "Lotes sem boleta válida", "0"),
               ("n_manual", "Lotes corrigidos", "0"), ("qt_lotes", "Qtd nos lotes", QTD),
               ("qt_posicao", "Qtd na posição", QTD), ("dif_qt", "Diferença", QTD),
               ("taxa_vcra", "Taxa média VCRA (convenção do grupo)", TAXA), ("taxa_media", "Taxa média", TAXA),
               ("taxa_cdi_mais", "Taxa média em CDI+ equivalente", TAXA), ("pu", "PU", PU),
               ("dt_pu", "Data do PU", DATA), ("financeiro", "Financeiro (R$)", VALOR)],
    "Lotes": [("codigo", "Ativo", None), ("isin", "ISIN", None), ("grupo", "Grupo", None), ("CD", "Lote", None),
              ("dt_compra", "Compra", DATA), ("qt_comprada", "Qtd comprada", QTD),
              ("qt_vendida", "Qtd vendida", QTD), ("qt_aberta", "Qtd em aberto", QTD),
              ("pu_compra", "PU de compra", PU), ("perc_cdi", "% do CDI (emissão)", TAXA),
              ("taxa_emissao", "Spread de emissão", TAXA), ("taxa_vcra", "Taxa VCRA (bruta)", TAXA),
              ("taxa_vcra_conv", "Taxa VCRA (convenção do grupo)", TAXA), ("data_boleta", "Data da boleta", DATA),
              ("casou_por", "Boleta casou por", None), ("taxa_boleta", "Taxa da boleta", TAXA),
              ("boleta_ok", "Boleta válida", None), ("dif_boleta_vcra", "Boleta − VCRA", TAXA),
              ("taxa", "Taxa usada", TAXA), ("fonte_taxa", "Fonte da taxa", None)],
    "Avisos": [("tipo", "Aviso", None), ("ativo", "Ativo", None), ("lote", "Lote", None), ("detalhe", "Detalhe", None)],
    "Parâmetros": [("parametro", "Parâmetro", None), ("valor", "Valor", None)],
}


def salvar_xlsx(caminho, lotes, ativos, carteira, avisos, parametros, notas):
    """Abas Carteira (com as notas embaixo), Ativos, Lotes, Avisos e Parâmetros."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    tabelas = {"Carteira": carteira, "Ativos": ativos,
               "Lotes": lotes.sort_values(["grupo", "codigo", "dt_compra"]).assign(
                   boleta_ok=lambda x: x["boleta_ok"].map({True: "sim", False: "não"})),
               "Avisos": pd.DataFrame(avisos, columns=["tipo", "ativo", "lote", "detalhe"]),
               "Parâmetros": pd.DataFrame(parametros, columns=["parametro", "valor"])}
    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        for aba, df in tabelas.items():
            colunas = COLUNAS[aba]
            df[[c for c, _, _ in colunas]].rename(columns={c: t for c, t, _ in colunas}).to_excel(
                xw, sheet_name=aba, index=False)
            ws = xw.book[aba]
            ws.freeze_panes = "B2"
            for j, (c, titulo, fmt) in enumerate(colunas, start=1):
                ws.cell(row=1, column=j).font = Font(bold=True)
                largura = 14 if fmt else max([len(str(v)) for v in df[c].dropna()] + [0])
                ws.column_dimensions[get_column_letter(j)].width = min(max(len(titulo), largura) + 2, 90)
                if fmt:
                    for (cel,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        cel.number_format = fmt
        ws = xw.book["Carteira"]
        ws.append([])
        for nota in notas:
            ws.append([nota])


def _data(s):
    return datetime.strptime(s, "%d/%m/%Y").date()


def main(argv=None):
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa
            pass

    ap = argparse.ArgumentParser(description="Taxa média de aquisição dos CRIs de um fundo: lotes da VCRAOPRF, "
                                             "taxa da boleta (FUNDOS_OPERACOES) e financeiro pela VW_FUNDOS_POSICAO.")
    ap.add_argument("--fundo", default=FUNDO_PADRAO, help=f"codigo_IAM_fundo (padrão: {FUNDO_PADRAO})")
    ap.add_argument("--data", type=_data, help="data de referência dd/mm/aaaa (padrão: última data da posição do fundo)")
    ap.add_argument("--fonte", choices=("boleta", "vcra"), default="boleta",
                    help="taxa dos CDI+: boleta (negociada, padrão) ou vcra (pelo PU pago); %%CDI usa sempre a boleta")
    ap.add_argument("--cdi", type=float, help="CDI em %% a.a. para converter %%CDI em CDI+ equivalente "
                                               "(padrão: CDIAcumulado da FUNDOS_DADOS_INDICES na data)")
    ap.add_argument("--correcoes", default=CORRECOES_PADRAO,
                    help="CSV CD;taxa;motivo com a taxa corrigida de lotes específicos (padrão: input/correcoes.csv)")
    args = ap.parse_args(argv)
    if not re.fullmatch(r"\w+", args.fundo):
        ap.error(f"--fundo inválido: {args.fundo!r}")
    if args.cdi is not None and args.cdi <= 0:
        ap.error("--cdi em % a.a. (ex.: 13.65)")
    try:
        correcoes = carregar_correcoes(args.correcoes)
    except (OSError, ValueError) as e:
        ap.error(str(e))

    try:
        d = carregar_dados(args.fundo, args.data, args.cdi)
    except Exception as e:  # noqa — banco (credenciais, driver, permissão) ou dado faltando
        sys.exit(f"[ERRO] {e}")
    data, cdi = d["data"], d["cdi"]
    print(f"[INFO] {args.fundo} em {data:%d/%m/%Y} | taxa dos CDI+: {args.fonte} | CDI {_br(cdi * 100)}% a.a. "
          f"({d['cdi_origem']}) | {len(correcoes)} correção(ões) manual(is)")

    lotes, ativos, carteira, avisos = calcular(d["vcra"], d["boletas"], d["posicao"], data, cdi, args.fonte,
                                               correcoes)
    if lotes.empty:
        sys.exit(f"[ERRO] nenhum lote {TIPO_CRI} em aberto na VCRAOPRF até {data:%d/%m/%Y}")
    print(f"[INFO] {len(lotes)} lotes {TIPO_CRI} em aberto em {lotes['codigo'].nunique()} ativos | boleta casada "
          f"em {lotes['taxa_boleta'].notna().sum()} ({(lotes['casou_por'] == 'PU').sum()} pelo PU)")
    if (ativos["dif_qt"].abs() <= 1e-6).all():
        print(f"[OK] quantidade dos lotes = posição de {data:%d/%m/%Y} nos {len(ativos)} ativos")
    for av in avisos:
        lote = f" ({av['lote']})" if av["lote"] else ""
        print(f"[{'INFO' if av['tipo'] == 'Correção manual' else 'AVISO'}] {av['tipo']} — {av['ativo']}{lote}: "
              f"{av['detalhe']}")
    for r in carteira.itertuples():
        if r.grupo == "%CDI":
            taxa = f"{_br(r.taxa_media)}% do CDI ({_cdi_mais(r.taxa_cdi_mais)} equivalente)"
        elif r.grupo == "CDI+":
            taxa = _cdi_mais(r.taxa_media)
        else:
            taxa = f"{_cdi_mais(r.taxa_cdi_mais)} (%CDI convertido com CDI de {_br(cdi * 100)}% a.a.)"
        nome = "Carteira" if r.grupo == "CARTEIRA" else r.grupo
        print(f"[OK] {nome}: {r.n_ativos} ativos, R$ {_br(r.financeiro, 0)} — {taxa}")

    parametros = [
        ("Fundo", args.fundo),
        ("Data de referência", f"{data:%d/%m/%Y}"),
        ("Taxa dos CDI+", "boleta (negociada)" if args.fonte == "boleta" else "VCRA (pelo PU pago)"),
        ("Taxa dos %CDI", "boleta (% do CDI)"),
        ("CDI para %CDI → CDI+", f"{_br(cdi * 100, 4)}% a.a. — {d['cdi_origem']}"),
        ("Casamento com a boleta", f"mesmo ativo, até {JANELA_BOLETA} dias, mesma quantidade ou PU "
                                   f"(tolerância {TOLERANCIA_PU})"),
        ("Boleta válida", f"%CDI >= {LIMITE_CONVENCAO} (% do CDI); CDI+ < {LIMITE_CONVENCAO} (spread)"),
        ("Correções manuais", f"{len(correcoes)} lote(s) — {args.correcoes}"),
        ("Gerado em", f"{datetime.now():%d/%m/%Y %H:%M}"),
    ]
    notas = ["CDI+: spread sobre 100% do CDI (% a.a.). %CDI: % do CDI.",
             f"CDI+ equivalente: %CDI convertido com CDI flat de {_br(cdi * 100)}% a.a. ({d['cdi_origem']}).",
             "Médias ponderadas pelo financeiro (qtd em aberto x último PU da posição); ativo sem taxa ou sem PU "
             "fica de fora (aba Avisos)."]
    os.makedirs(SAIDA_DIR, exist_ok=True)
    xlsx = os.path.join(SAIDA_DIR, f"taxa_media_{args.fundo}_{data:%Y%m%d}.xlsx")
    try:
        salvar_xlsx(xlsx, lotes, ativos, carteira, avisos, parametros, notas)
    except PermissionError:
        sys.exit(f"[ERRO] não deu para gravar {xlsx}: o arquivo está aberto no Excel?")
    print(f"[OK] {xlsx}")


if __name__ == "__main__":
    main()
