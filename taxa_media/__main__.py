# ─────────────────────────────────────────────
# taxa_media/__main__.py
# Taxa média de aquisição da renda fixa de uma lista de fundos: crédito (CRI, debênture, CRA...) e títulos públicos
# em CDI+, %CDI, SELIC+, pré e IPCA+, e as cotas de FIDC. Taxa de cada lote em aberto na custódia, média por ativo
# (ponderada pela quantidade) e da carteira (ponderada pelo financeiro atual), por grupo e no total em CDI+
# equivalente. Com --desde, também a série da carteira desde o início do fundo.
#
# Entrada: input/fundos.xlsx (colunas S/N aceitam S, Sim, X ou 1)
#   Fundos      fundo (codigo_IAM_fundo) · rodar (S/N, vazio = S) · nome (livre)
#   Tipos       RFTP_CD da VCRAOPRF · incluir (S/N) · descricao. Tipo que aparece na VCRA e não está na aba vira
#               aviso e fica fora. As compromissadas ficam com N: LTN-O, NTN-O e LFT O são over do caixa; CRIOP e
#               CRILA usam o código do CRI de lastro e somariam a compromissada ao CRI comprado em definitivo.
#               FIDC não é RFTP_CD: liga e desliga as cotas de FIDC da posição (sem a linha, entram).
#   Correções   fundo · CD (lote) · taxa · aplicar (S/N, vazio = S) · motivo. Taxa na convenção do grupo do ativo;
#               na cota de FIDC, o CD é o código da cota.
#
# Fontes (DBSX337 pelo sql_connection.py desta pasta, credenciais SQL_LOGIN e SQL_PASSWORD; todas as
# leituras WITH (NOLOCK)):
#   VCRAOPRF              custódia de renda fixa, uma linha por movimento de cada lote (CD). Trades = SG_OPERACAO
#                         C/V (M, J e I são eventos).
#   FUNDOS_OPERACOES      boleta do Lake: exec_taxa = taxa negociada (%CDI: % do CDI; os outros grupos: taxa ou
#                         spread % a.a.).
#   VW_FUNDOS_POSICAO     quantidade na data (conferência) e PU da última data com preço (total/quantidade); nas
#                         cotas de FIDC (código 5 + letra), a quantidade e o valor da cota.
#   FUNDOS_DADOS_INDICES  CDIAcumulado: CDI do dia, para levar %CDI e pré ao CDI+ equivalente.
#   FUNDOS_CADASTRO_ATIVOS  cadastro dos ativos (emissor, rating, vencimento, descrição...), uma linha por campo e
#                         data de alteração, com o valor vigente na data: a ficha do ativo no HTML e o indexador e a
#                         taxa_emissao das cotas de FIDC.
#
# Cálculo (por fundo):
#   escopo    tipos com incluir = S e indexador de um dos grupos (CDI, SELIC, pré ou IPCA); IGP-M, dólar e
#             outros ficam fora, com aviso. Nos títulos públicos a VCRA usa o código SELIC (ex.: 210100) e a
#             posição e o Lake, o ISIN: o código da VCRA vira o da posição pelo ISIN.
#   lote      qtd em aberto = compras - vendas do mesmo CD. Lote com sufixo A (ex.: L10947A) é relançamento
#             da VCRA (venda do original + compra um dia depois no mesmo PU) e já sai certo nessa conta.
#   boleta    cada compra casa com a boleta do mesmo ativo em até 5 dias, pela quantidade ou pelo PU (há
#             boletas com o dobro da quantidade da custódia); prefere a quantidade e a data mais próxima.
#   grupo     pelo indexador da VCRA (IDDIR_CD): CDI vira CDI+ se o ativo tem spread de emissão
#             (VL_JUROS_EMISSAO > 0), senão %CDI, com a emissão no IDDIR_PC; SELIC vira SELIC+, pré, Pré, e IPCA,
#             IPCA+. O indexador da VW_FUNDOS_POSICAO classifica errado alguns papéis.
#   taxa      %CDI: só a boleta — a VCRA grava %CDI como IDDIR_PC + spread, e isso não converte de forma
#             consistente para o % do CDI da boleta. Os outros grupos: a boleta (--fonte lake, padrão) ou a VCRA,
#             que sai do PU pago (--fonte vcra); sem boleta válida, a VCRA. A boleta vem do Lake
#             (FUNDOS_OPERACOES), o sistema que não é a VCRA: a fonte da taxa sai como Lake, VCRA, cadastro ou
#             manual. Boleta válida: %CDI >= 50 (% do CDI), os outros < 50 (taxa ou spread). Lote sem taxa fica
#             fora da média do ativo; a aba Correções corrige a taxa de lotes específicos.
#   FIDC      as cotas de FIDC da posição (código 5 + letra, como o SAC_Cota.py cadastra) não estão na VCRA: cada
#             uma é um lote com a quantidade da posição e o próprio código no CD. Grupo pelo indexador do cadastro;
#             taxa = média das compras do Lake até a data (exec_qt > 0), ponderada pela quantidade, ou, sem boleta
#             válida, a taxa_emissao do cadastro. Sem indexador no cadastro, a cota fica fora (aviso).
#   ativo     média das taxas dos lotes, ponderada pela quantidade em aberto.
#   carteira  média ponderada pelo financeiro (qtd x último PU), por grupo e no total. O total junta os grupos
#             em CDI+ equivalente, com CDI flat: %CDI vira (1 + %CDI x CDI diário)^252 / (1 + CDI) - 1; SELIC+
#             entra como está (SELIC = CDI); pré vira (1 + taxa) / (1 + CDI) - 1; IPCA+ vira (1 + IPCA) x
#             (1 + spread) / (1 + CDI) - 1, só com --ipca (sem ele, o IPCA+ sai à parte e fora do total).
#   série     com --desde, o mesmo cálculo em cada data (a última de cada mês ou, com --frequencia diaria,
#             todas as datas com posição), só com o que existia nela: VCRA e boletas até a data, PU dos
#             DIAS_PU dias anteriores e CDI do dia.
#   trades    cada compra e venda da VCRA no escopo, com a taxa pelas mesmas regras dos lotes, e a média por
#             dia e operação, ponderada pelo financeiro do trade (qtd x PU da operação). Relançamento da VCRA
#             (venda do lote X + compra do lote XA no mesmo PU) não é trade; as cotas de FIDC ficam fora.
#
# Saída: output/taxa_media_<aaaammdd>.xlsx (com um fundo só, taxa_media_<fundo>_<aaaammdd>.xlsx; com --desde,
#   taxa_media_historico_...) — abas Histórico (só com --desde: uma linha por fundo e data), Carteira, Ativos e
#   Lotes (da última data), Trades por dia e Trades (todo o histórico até a data; com --desde, desde o início
#   da série), Avisos (da última data), Cadastro (as versões de cada campo dos ativos, até a data), Fundos e
#   Parâmetros, com a coluna Fundo (menos o Cadastro, que é do ativo). Erro num fundo não para os outros: o fundo
#   sai na aba Fundos com o erro.
#   Ao lado, o mesmo nome em .html: relatório com as mesmas tabelas, gráficos e a ficha de cada ativo
#   (html_report.py, no layout do attribution). --html regera o relatório de um .xlsx já gravado, sem o banco.
#   Com --diario, também um relatório por data em output/diario/aaaammdd.html (carteira, ativos, lotes e avisos da
#   data e os trades desde a data anterior) e o output/index.html: calendário com todos os dias da pasta, o
#   relatório do dia escolhido e o histórico e os trades consolidados. Rodar todo dia com --diario acumula os dias.
#
# Requisitos: ODBC Driver 18 for SQL Server, pyodbc, sqlalchemy, pandas, openpyxl.
#
# Uso (de dentro de taxa_media/, ou "python taxa_media" a partir de itau_asset/):
#   python __main__.py                          (fundos da aba Fundos, cada um na última data da posição)
#   python __main__.py --data 23/10/2026
#   python __main__.py --fundo ITAIMOBCDIFII    (só esse fundo; tipos e correções continuam vindo do Excel)
#   python __main__.py --fonte vcra --cdi 13.65
#   python __main__.py --ipca 4,5                (IPCA projetado: o IPCA+ entra na carteira em CDI+ equivalente)
#   python __main__.py --fundo ITAIMOBCDIFII --desde            (série mensal desde o início do fundo)
#   python __main__.py --desde 01/06/2026 --frequencia diaria   (série diária a partir de 01/06/2026)

#   python __main__.py --fundo ITAIMOBCDIFII --desde                       (mensal, desde o início do fundo)
#   python __main__.py --fundo ITAIMOBCDIFII --desde --frequencia diaria   (todos os dias)
#   python __main__.py --fundo ITAIMOBCDIFII --desde 01/06/2026            (a partir de uma data)
#   python __main__.py --html output/taxa_media_historico_ITAIMOBCDIFII_20261023.xlsx   (só o relatório HTML)
#   python __main__.py --fundo ITAIMOBCDIFII --desde --diario    (um relatório por dia desde o início + index.html)
#   python __main__.py --diario                                   (a última data, somada aos dias já na pasta)

# ─────────────────────────────────────────────

import argparse
import os
import re
import sys
import unicodedata
from collections import Counter
from datetime import datetime

import numpy as np
import pandas as pd

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA_DIR = os.path.join(AQUI, "output")
ENTRADA_PADRAO = os.path.join(AQUI, "input", "fundos.xlsx")

# tipos da VCRAOPRF quando o Excel não tem a aba Tipos (True = entra no cálculo). FIDC não é RFTP_CD: são as cotas de
# FIDC da posição, que não estão na VCRA. LTN, NTN-B e NTN-F são os códigos prováveis dos títulos públicos em definitivo
TIPOS_PADRAO = {"CRI-I": True, "LFT": True, "LTN": True, "NTN-B": True, "NTN-F": True, "FIDC": True,
                "LTN-O": False, "NTN-O": False, "LFT O": False, "CRIOP": False, "CRILA": False}
# grupos de taxa, na ordem das saídas, e o sufixo das colunas de cada um (taxa_<sufixo>, fin_<sufixo>). O grupo sai
# do indexador: CDI vira %CDI ou CDI+ (com spread de emissão); SELIC (LFT), SELIC+; pré, Pré; IPCA, IPCA+.
GRUPOS = [("%CDI", "pct"), ("CDI+", "mais"), ("SELIC+", "selic"), ("Pré", "pre"), ("IPCA+", "ipca")]
PADRAO_FIDC = r"^5[A-Z]"  # cotas de FIDC na posição: o mesmo filtro com que o SAC_Cota.py cadastra os FIDC
JANELA_BOLETA = 5       # dias entre a compra na VCRA e a data da boleta
TOLERANCIA_PU = 1e-4    # PU da boleta x PU da VCRA, para casar quando a quantidade não bate
LIMITE_CONVENCAO = 50   # boleta de %CDI vem em % do CDI (>= 50); dos outros grupos, em taxa ou spread % a.a. (< 50)
DIAS_PU = 30            # até quantos dias antes da data procurar o último PU com preço na posição
# boleta x VCRA na convenção do grupo acima disso vira aviso: em p.p. de taxa ou spread; no %CDI, em pontos de % do
# CDI (convertida com CDI flat, a VCRA costuma ficar até ~1,5 ponto abaixo da boleta)
DIVERGENCIA = {"CDI+": 0.25, "%CDI": 3.0, "SELIC+": 0.25, "Pré": 0.25, "IPCA+": 0.25}
# ordem dos avisos no console e na aba Avisos ("Correção manual" e "FIDC pelo cadastro" saem como [INFO])
ORDEM_AVISOS = ["Quantidade", "Sem lote", "Fora do escopo", "Tipo novo", "Fora da média", "Média parcial",
                "FIDC sem taxa", "FIDC sem indexador", "Boleta inválida", "Sem boleta", "Correção manual",
                "FIDC pelo cadastro", "Correção não aplicada", "Boleta × VCRA", "Venda sem compra", "Lote negativo"]
INFORMATIVOS = {"Correção manual", "FIDC pelo cadastro"}


# ── Entrada (input/fundos.xlsx) ──────────────────────────────────────────────
def _chave(texto):
    """Nome de aba ou de coluna sem acento, sem espaços nas pontas e em minúsculas."""
    return unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().strip().lower()


def _sim(valor, padrao=False):
    if pd.isna(valor) or not str(valor).strip():
        return padrao
    return str(valor).strip().upper() in ("S", "SIM", "X", "1")


def _texto(valor):
    return "" if pd.isna(valor) else str(valor).strip()


def _strip(s):
    """Coluna de texto sem espaços nas pontas; vazio e não texto ficam como estão (vale para coluna toda nula)."""
    return s.map(lambda v: v.strip() if isinstance(v, str) else v)


def _e_fidc(codigos):
    """Máscara das cotas de FIDC (código 5 + letra) numa coluna de códigos."""
    return codigos.map(lambda c: isinstance(c, str) and re.match(PADRAO_FIDC, c) is not None).astype(bool)


def _numero(texto):
    texto = str(texto).strip()
    return float(texto.replace(".", "").replace(",", ".") if "," in texto else texto)


def _aba(abas, nome, colunas, caminho):
    """Aba do Excel com as colunas pedidas (nomes sem acento e minúsculos); None se a aba não existir."""
    df = abas.get(_chave(nome))
    if df is None:
        return None
    df = df.rename(columns=_chave)
    faltam = [c for c in colunas if c not in df.columns]
    if faltam:
        raise ValueError(f"{caminho}, aba {nome}: faltam as colunas {faltam}")
    return df.dropna(how="all")


def carregar_entrada(caminho):
    """input/fundos.xlsx -> (fundos [(fundo, nome)], tipos {RFTP_CD: incluir}, correções {fundo: {CD: (taxa,
    motivo)}}). Sem a aba Tipos vale TIPOS_PADRAO; sem a aba Correções, nenhuma correção."""
    abas = {_chave(nome): df for nome, df in pd.read_excel(caminho, sheet_name=None, dtype=str).items()}

    fundos = []
    df = _aba(abas, "Fundos", ["fundo"], caminho)
    for i, r in ([] if df is None else df.iterrows()):
        fundo = _texto(r["fundo"])
        if not fundo or not _sim(r.get("rodar"), padrao=True) or fundo in dict(fundos):
            continue
        if not re.fullmatch(r"\w+", fundo):
            raise ValueError(f"{caminho}, aba Fundos, linha {i + 2}: código de fundo inválido {fundo!r}")
        fundos.append((fundo, _texto(r.get("nome"))))

    df = _aba(abas, "Tipos", ["rftp_cd", "incluir"], caminho)
    tipos = dict(TIPOS_PADRAO) if df is None else {
        _texto(r["rftp_cd"]).upper(): _sim(r["incluir"]) for _, r in df.iterrows() if _texto(r["rftp_cd"])}

    correcoes = {}
    df = _aba(abas, "Correções", ["fundo", "cd", "taxa"], caminho)
    for i, r in ([] if df is None else df.iterrows()):
        if not _sim(r.get("aplicar"), padrao=True):
            continue
        linha = f"{caminho}, aba Correções, linha {i + 2}"
        if not (_texto(r["fundo"]) and _texto(r["cd"]) and _texto(r["taxa"])):
            raise ValueError(f"{linha}: precisa de fundo, CD e taxa")
        try:
            taxa = _numero(r["taxa"])
        except ValueError:
            raise ValueError(f"{linha}: taxa inválida {r['taxa']!r}") from None
        correcoes.setdefault(_texto(r["fundo"]), {})[_texto(r["cd"])] = (taxa, _texto(r.get("motivo")))
    return fundos, tipos, correcoes


# ── SQL ──────────────────────────────────────────────────────────────────────
SQL_DATAS = """
SELECT MIN(date) AS primeira, MAX(date) AS ultima
FROM   VW_FUNDOS_POSICAO WITH (NOLOCK)
WHERE  codigo_IAM_fundo = '{fundo}'
"""

SQL_VCRA = """
SELECT  CD, CD_CETIP_SELIC AS codigo, ISIN AS isin, RFTP_CD AS tipo, IDDIR_CD AS indexador, DT AS data,
        SG_OPERACAO AS op, QT AS qt, VL_PU_OPERACAO AS pu, IDDIR_PC AS perc_cdi,
        VL_JUROS_EMISSAO AS taxa_emissao, VL_JUROSOPER AS taxa_vcra
FROM    VCRAOPRF WITH (NOLOCK)
WHERE   CLCLI_CD = '{fundo}' AND DT <= '{data}' AND SG_OPERACAO IN ('C', 'V')
"""

SQL_BOLETAS = """
SELECT  codigo_IAM_ativo AS codigo, date_operacao AS data, exec_qt AS qt, exec_preco AS pu,
        exec_taxa AS taxa
FROM    FUNDOS_OPERACOES WITH (NOLOCK)
WHERE   codigo_IAM_fundo = '{fundo}'
  AND   date_operacao <= DATEADD(day, {janela}, '{data}')
  AND   codigo_IAM_ativo IN ({codigos})
"""

# as cotas de FIDC são cadastradas como Fixed Income (SAC_Cota.py); o 'Fund' com código de FIDC é por garantia
SQL_POSICAO = """
SELECT  codigo_IAM_ativo AS codigo, isin, date AS data, quantidade, total
FROM    VW_FUNDOS_POSICAO WITH (NOLOCK)
WHERE   codigo_IAM_fundo = '{fundo}'
  AND   (tipo = 'Fixed Income' OR (tipo = 'Fund' AND codigo_IAM_ativo LIKE '5[A-Z]%'))
  AND   date BETWEEN DATEADD(day, -{dias}, '{desde}') AND '{data}'
"""

SQL_CDI = """
SELECT date AS data, pnl_perc
FROM   FUNDOS_DADOS_INDICES WITH (NOLOCK)
WHERE  codigo_IAM_indice = 'CDIAcumulado' AND pnl_perc > 0
"""

SQL_CADASTRO = """
SELECT  codigo_IAM_ativo AS codigo, campo, date_change AS desde, valor_str, valor_num, valor_date, fonte
FROM    FUNDOS_CADASTRO_ATIVOS WITH (NOLOCK)
WHERE   codigo_IAM_ativo IN ({codigos}) AND date_change <= '{data}'
"""


def _consulta(query, sql, nome, obrigatoria=True):
    """SELECT pelo sql_connection.query, que devolve DataFrame vazio também quando a consulta dá erro
    (e imprime a mensagem)."""
    df = query(sql)
    if df.empty and obrigatoria:
        raise RuntimeError(f"{nome}: nenhuma linha (ou erro na consulta: veja a mensagem acima)")
    return df


def _lista_sql(valores):
    return ", ".join("'" + str(v).replace("'", "''") + "'" for v in valores)


def carregar_fundo(query, fundo, data, tipos, desde=None):
    """VCRA (todos os tipos), boletas dos ativos dos tipos incluídos e posição de um fundo até a data. data None =
    última data da posição do fundo. desde (início da série): None = só a data; "inicio" = primeira data da
    posição do fundo; ou uma data."""
    if data is None or desde == "inicio":
        datas = _consulta(query, SQL_DATAS.format(fundo=fundo), "VW_FUNDOS_POSICAO").iloc[0]
        if pd.isna(datas["ultima"]):
            raise RuntimeError("sem posição na VW_FUNDOS_POSICAO")
        data = data or pd.Timestamp(datas["ultima"]).date()
        if desde == "inicio":
            desde = pd.Timestamp(datas["primeira"]).date()
    desde = desde or data
    if desde > data:
        raise RuntimeError(f"--desde {desde:%d/%m/%Y} depois da última data, {data:%d/%m/%Y}")
    p = {"fundo": fundo, "data": f"{data:%Y-%m-%d}", "desde": f"{desde:%Y-%m-%d}", "janela": JANELA_BOLETA,
         "dias": DIAS_PU}
    vcra = _consulta(query, SQL_VCRA.format(**p), f"VCRAOPRF (até {data:%d/%m/%Y})")
    posicao = _consulta(query, SQL_POSICAO.format(**p), "VW_FUNDOS_POSICAO")
    vcra = _codigo_da_posicao(vcra, posicao)
    incluidos = _strip(vcra["tipo"]).fillna("").str.upper().map(lambda t: tipos.get(t, False)).astype(bool)
    codigos = set(_strip(vcra.loc[incluidos, "codigo"]).dropna())
    if tipos.get("FIDC", True):  # as boletas do Lake das cotas de FIDC também
        cods = _strip(posicao["codigo"])
        codigos |= set(cods[_e_fidc(cods)])
    codigos = sorted(codigos)
    boletas = pd.DataFrame()
    if codigos:
        boletas = _consulta(query, SQL_BOLETAS.format(codigos=_lista_sql(codigos), **p), "FUNDOS_OPERACOES",
                            obrigatoria=False)
    # data de cada linha, para fatiar o que existia em cada data da série
    dt = {nome: None if df.empty else pd.to_datetime(df["data"]).dt.normalize()
          for nome, df in (("vcra", vcra), ("boletas", boletas), ("posicao", posicao))}
    if not (dt["posicao"] == pd.Timestamp(data)).any():
        raise RuntimeError(f"sem posição em {data:%d/%m/%Y} na VW_FUNDOS_POSICAO")
    # cadastro dos ativos da VCRA no escopo e da posição, para a ficha do ativo no HTML
    todos = sorted(set(codigos) | set(_strip(posicao["codigo"]).dropna()))
    cadastro = pd.DataFrame()
    if todos:
        cadastro = _consulta(query, SQL_CADASTRO.format(codigos=_lista_sql(todos), **p), "FUNDOS_CADASTRO_ATIVOS",
                             obrigatoria=False)
    return {"vcra": vcra, "boletas": boletas, "posicao": posicao, "data": data, "desde": desde, "dt": dt,
            "cadastro": versoes_cadastro(cadastro)}


def _codigo_da_posicao(vcra, posicao):
    """Troca o código da VCRA (CD_CETIP_SELIC) pelo da posição quando o ativo não está na posição com o mesmo código
    mas o ISIN bate: nos títulos públicos a VCRA usa o código SELIC e a posição e o Lake, o ISIN."""
    pos = posicao.dropna(subset=["codigo", "isin"])
    por_isin = dict(zip(_strip(pos["isin"]), _strip(pos["codigo"])))
    codigo = _strip(vcra["codigo"])
    novo = _strip(vcra["isin"]).map(por_isin)
    troca = novo.notna() & ~codigo.isin(set(_strip(pos["codigo"]))) & (novo != codigo)
    return vcra.assign(codigo=codigo.where(~troca, novo)) if troca.any() else vcra


def datas_da_serie(d, frequencia):
    """Datas em que a carteira é calculada: as datas com posição de renda fixa de d["desde"] a d["data"], todas
    (diaria) ou a última de cada mês (mensal). Sem --desde, desde = data e sai só a data."""
    datas = d["dt"]["posicao"].drop_duplicates().sort_values()
    datas = datas[datas.between(pd.Timestamp(d["desde"]), pd.Timestamp(d["data"]))]
    if frequencia == "mensal":
        datas = datas.groupby(datas.dt.to_period("M")).max()
    return [x.date() for x in datas]


def fatiar(d, data):
    """VCRA, boletas e posição do fundo como as consultas viriam rodando só nessa data: VCRA até a data, boletas
    até a data + JANELA_BOLETA e posição dos DIAS_PU dias até a data."""
    ref, dt, boletas = pd.Timestamp(data), d["dt"], d["boletas"]
    if dt["boletas"] is not None:
        boletas = boletas[dt["boletas"] <= ref + pd.Timedelta(days=JANELA_BOLETA)]
    return (d["vcra"][dt["vcra"] <= ref], boletas,
            d["posicao"][dt["posicao"].between(ref - pd.Timedelta(days=DIAS_PU), ref)])


def carregar_cdi(query):
    """Série do CDIAcumulado (pnl_perc = CDI do dia, em fração), lida uma vez por execução."""
    cdis = _consulta(query, SQL_CDI, "FUNDOS_DADOS_INDICES (CDIAcumulado)")
    cdis = cdis.assign(data=pd.to_datetime(cdis["data"]).dt.normalize(), pnl_perc=pd.to_numeric(cdis["pnl_perc"]))
    return cdis.sort_values("data", kind="stable")


def cdi_na_data(cdis, data):
    """CDI a.a. (fração) do CDIAcumulado no último dia até a data, e de onde ele veio."""
    ate = cdis[cdis["data"] <= pd.Timestamp(data)]
    if ate.empty:
        raise RuntimeError(f"sem CDIAcumulado até {data:%d/%m/%Y}: informe com --cdi")
    linha = ate.iloc[-1]
    cdi = (1 + float(linha["pnl_perc"])) ** 252 - 1
    if not 0.01 < cdi < 0.5:
        raise RuntimeError(f"CDI de {cdi:.2%} a.a. no CDIAcumulado fora do esperado: informe com --cdi")
    return cdi, f"CDIAcumulado da FUNDOS_DADOS_INDICES em {linha['data']:%d/%m/%Y}"


def versoes_cadastro(cad):
    """FUNDOS_CADASTRO_ATIVOS (uma linha por ativo, campo e data de alteração) -> as versões de cada campo: só as
    linhas em que o valor muda (a tabela regrava o mesmo valor a cada carga, ex.: vencimento), com o valor no tipo do
    campo (texto, número ou data). Colunas: codigo, campo, desde, valor, fonte."""
    colunas = ["codigo", "campo", "desde", "valor", "fonte"]
    if cad.empty:
        return pd.DataFrame(columns=colunas)
    cad = cad.copy()
    for c in ("codigo", "campo", "valor_str", "fonte"):
        cad[c] = cad[c].map(lambda v: (v.strip() or None) if isinstance(v, str) else None)
    cad["desde"] = pd.to_datetime(cad["desde"]).dt.normalize()
    datas = pd.to_datetime(cad["valor_date"]).map(lambda v: None if pd.isna(v) else v.date())
    numeros = pd.to_numeric(cad["valor_num"], errors="coerce").astype(object)
    cad["valor"] = cad["valor_str"].where(cad["valor_str"].notna(), numeros.where(numeros.notna(), datas))
    cad = cad[cad["valor"].notna()].sort_values(["codigo", "campo", "desde"], kind="stable")
    texto = cad["valor"].astype(str)
    muda = texto != texto.groupby([cad["codigo"], cad["campo"]]).shift()
    return cad.loc[muda, colunas].reset_index(drop=True)


def _cdi(cdi_fixo, cdis, data):
    """(CDI a.a. em fração, origem) na data: o --cdi (% a.a.), se veio, ou o CDIAcumulado."""
    return (cdi_fixo / 100, "informado em --cdi") if cdi_fixo is not None else cdi_na_data(cdis, data)


# ── Cálculo ──────────────────────────────────────────────────────────────────
def _br(x, casas=2):
    """Número no formato brasileiro (1.234,56)."""
    return f"{x:,.{casas}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def _cdi_mais(x):
    return f"CDI{'+' if x >= 0 else '−'}{_br(abs(x))}"


def _aviso(avisos, aviso, ativo, detalhe, lote=""):
    avisos.append({"aviso": aviso, "ativo": ativo, "lote": lote, "detalhe": detalhe})


def _normalizar(df, textos=(), datas=(), numeros=()):
    df = df.copy()
    for c in textos:
        df[c] = _strip(df[c])
    for c in datas:
        df[c] = pd.to_datetime(df[c]).dt.normalize()
    for c in numeros:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def _media_ponderada(df, chave, valor, peso):
    """Média de valor ponderada por peso em cada grupo de chave, só com as linhas que têm valor."""
    ok = df[valor].notna()
    s = df.assign(_v=(df[valor] * df[peso]).where(ok, 0.0), _p=df[peso].where(ok, 0.0)).groupby(chave)[["_v", "_p"]].sum()
    return s["_v"] / s["_p"].replace(0, np.nan)


def _familia(indexador):
    """Família do indexador (IDDIR_CD da VCRA ou indexador do cadastro): CDI, SELIC, PRE, IPCA ou None (fora)."""
    i = re.sub(r"[^A-Z]", "", _chave(indexador).upper())  # sem acento, espaço, %, + e sufixos como D-3 ou M-1
    for prefixo, familia in (("CDI", "CDI"), ("DI", "CDI"), ("SELIC", "SELIC"), ("IPCA", "IPCA"), ("PRE", "PRE")):
        if i.startswith(prefixo):
            return familia
    return None


def _grupo(familia, taxa_emissao):
    """Grupo de taxa: CDI+ (CDI com spread de emissão) ou %CDI, SELIC+, Pré ou IPCA+; None fora dos grupos."""
    if familia == "CDI":
        return "CDI+" if taxa_emissao > 0 else "%CDI"
    return {"SELIC": "SELIC+", "PRE": "Pré", "IPCA": "IPCA+"}.get(familia)


def _grupo_cadastro(indexador, taxa):
    """Grupo de uma cota de FIDC pelo indexador do cadastro. No CDI, o % ou o + do indexador decide; sem eles, a
    convenção da taxa (do cadastro ou da boleta): >= LIMITE_CONVENCAO é % do CDI."""
    familia = _familia(indexador)
    if familia != "CDI":
        return _grupo(familia, taxa)
    if "%" in str(indexador) or ("+" not in str(indexador) and pd.notna(taxa) and taxa >= LIMITE_CONVENCAO):
        return "%CDI"
    return "CDI+"


def _vigente(cadastro, data, campos):
    """{codigo: valor} do primeiro dos campos que o ativo tem no cadastro (versoes_cadastro), na última versão até
    a data."""
    if cadastro is None or cadastro.empty:
        return {}
    c = cadastro[cadastro["desde"] <= pd.Timestamp(data)].assign(campo=lambda x: x["campo"].str.lower())
    valores = {}
    for campo in reversed(campos):  # o primeiro campo da lista sobrescreve os outros
        valores.update(c[c["campo"] == campo].sort_values("desde").groupby("codigo")["valor"].last().to_dict())
    return valores


def _grupo_por_ativo(vcra):
    """{codigo: grupo} pelo indexador e pelo maior spread de emissão do ativo na VCRA (o mesmo grupo em todos os
    lotes e trades do ativo)."""
    a = vcra.groupby("codigo").agg(indexador=("indexador", "first"), emissao=("taxa_emissao", "max"))
    return {c: _grupo(_familia(r.indexador), r.emissao) for c, r in a.iterrows()}


def filtrar_escopo(vcra, tipos, avisos):
    """Movimentos que entram no cálculo: tipo com incluir = S e indexador de um dos GRUPOS (CDI, SELIC, pré ou
    IPCA). Devolve (VCRA no escopo, {codigo: motivo} dos ativos que ficaram fora, para explicar quem está na posição
    sem lote)."""
    tipo = vcra["tipo"].fillna("").str.upper()
    novos = vcra[~tipo.isin(list(tipos))]
    for t, g in novos.groupby(tipo.loc[novos.index]):
        indexadores = ", ".join(sorted(g["indexador"].dropna().unique())) or "sem indexador"
        _aviso(avisos, "Tipo novo", "", f"RFTP_CD {t or '(vazio)'} ({indexadores}; {len(g)} movimento"
               f"{'s' if len(g) > 1 else ''}) não está na aba Tipos: ficou fora")
    incluido = tipo.map(lambda t: tipos.get(t, False)).astype(bool)
    conhecido = vcra["indexador"].map(_familia).notna()
    fora = {}
    for r in vcra[~incluido].itertuples():
        fora.setdefault(r.codigo, f"na VCRA como {r.tipo or '(vazio)'}, tipo fora do cálculo")
    for codigo, g in vcra[incluido & ~conhecido].groupby("codigo"):
        indexadores = ", ".join(sorted(g["indexador"].dropna().unique())) or "vazio"
        fora[codigo] = f"indexador {indexadores}: só entram CDI, SELIC, pré e IPCA"
    return vcra[incluido & conhecido], fora


def casar_boletas(compras, boletas):
    """Acrescenta a cada compra da VCRA a boleta do mesmo ativo em até JANELA_BOLETA dias com a mesma
    quantidade ou o mesmo PU; prefere a quantidade e, depois, a data mais próxima."""
    novas = ["data_boleta", "qt_boleta", "taxa_boleta", "casou_por"]
    if boletas.empty:
        return compras.assign(**{c: np.nan for c in novas})
    b = boletas.rename(columns={"data": "data_boleta", "qt": "qt_boleta", "pu": "pu_boleta", "taxa": "taxa_boleta"})
    m = compras[["codigo", "data", "qt", "pu"]].reset_index().merge(b, on="codigo")
    mesma_qt = np.isclose(m["qt_boleta"].abs(), m["qt"], rtol=0, atol=1e-6)
    mesmo_pu = ((m["pu_boleta"] - m["pu"]).abs() < TOLERANCIA_PU).to_numpy()
    m["dias"] = (m["data_boleta"] - m["data"]).dt.days.abs()
    m["prioridade"] = np.where(mesma_qt, 0, 1)
    m["casou_por"] = np.where(mesma_qt, "quantidade", "PU")
    m = m[(m["dias"].to_numpy() <= JANELA_BOLETA) & (mesma_qt | mesmo_pu)]
    m = m.sort_values(["index", "prioridade", "dias"]).drop_duplicates("index").set_index("index")
    return compras.join(m[novas])


def _taxa_usada(df, fonte):
    """(boleta válida, taxa usada, fonte da taxa) de cada lote ou trade, a partir de grupo, taxa_boleta e
    taxa_vcra: o %CDI só usa a boleta do Lake (a VCRA grava %CDI como IDDIR_PC + spread); os outros grupos, em taxa
    ou spread a.a., usam a boleta (fonte lake) ou a VCRA (fonte vcra), e a VCRA sem boleta válida. A fonte sai como
    Lake, VCRA ou vazia (sem taxa)."""
    pct = df["grupo"] == "%CDI"
    boleta_ok = np.where(pct, df["taxa_boleta"] >= LIMITE_CONVENCAO, df["taxa_boleta"] < LIMITE_CONVENCAO)
    usa_boleta = boleta_ok & (pct | (fonte != "vcra"))
    return (boleta_ok, np.where(usa_boleta, df["taxa_boleta"], np.where(pct, np.nan, df["taxa_vcra"])),
            np.where(usa_boleta, "Lake", np.where(pct, "", "VCRA")))


def _vcra_na_convencao(df, cdi_d):
    """Taxa da VCRA na convenção do grupo, para comparar com a boleta: no %CDI, o % do CDI que rende o mesmo que
    IDDIR_PC% do CDI + o spread da VCRA, com CDI flat; nos outros grupos, o próprio VL_JUROSOPER. cdi_d = CDI diário
    (um número, ou um por linha)."""
    p, s = df["perc_cdi"] / 100, df["taxa_vcra"] / 100
    return np.where(df["grupo"] == "%CDI", ((1 + p * cdi_d) * (1 + s) ** (1 / 252) - 1) / cdi_d * 100, df["taxa_vcra"])


def _em_cdi_mais(grupo, taxa, cdi, ipca=None):
    """Taxa em CDI+ equivalente, com CDI flat (cdi = CDI a.a. em fração, um número ou um por linha): CDI+ e SELIC+
    ficam como estão (a SELIC conta como igual ao CDI); o %CDI vira o spread sobre 100% do CDI que rende o mesmo em
    um ano; o pré, (1 + taxa) / (1 + CDI) - 1; o IPCA+, (1 + IPCA) x (1 + spread) / (1 + CDI) - 1 com o IPCA
    projetado (ipca, em fração) e vazio sem ele."""
    grupo = np.asarray(grupo, dtype=object)
    taxa, cdi = np.asarray(taxa, dtype=float), np.asarray(cdi, dtype=float)
    cdi_d = (1 + cdi) ** (1 / 252) - 1
    pct = ((1 + taxa / 100 * cdi_d) ** 252 / (1 + cdi) - 1) * 100
    pre = ((1 + taxa / 100) / (1 + cdi) - 1) * 100
    ipca_mais = ((1 + ipca) * (1 + taxa / 100) / (1 + cdi) - 1) * 100 if ipca is not None else np.nan * taxa
    return np.select([grupo == "%CDI", grupo == "Pré", grupo == "IPCA+"], [pct, pre, ipca_mais], default=taxa)


def montar_lotes(vcra, boletas, fonte, correcoes, cdi_d, avisos):
    """Lotes em aberto da VCRA (compras - vendas do mesmo CD) com a boleta casada, o grupo e a taxa usada."""
    compras = casar_boletas(vcra[vcra["op"] == "C"].reset_index(drop=True), boletas)
    vendas = vcra[vcra["op"] == "V"].groupby("CD")["qt"].sum()

    chave = ["codigo", "CD"]
    lotes = compras.groupby(chave).agg(
        isin=("isin", "first"), tipo=("tipo", "first"), indexador=("indexador", "first"), dt_compra=("data", "min"),
        perc_cdi=("perc_cdi", "max"), taxa_emissao=("taxa_emissao", "max"), qt_comprada=("qt", "sum"),
        data_boleta=("data_boleta", "min"), casou_por=("casou_por", "first"))
    for coluna, nome in (("pu", "pu_compra"), ("taxa_vcra", "taxa_vcra"), ("taxa_boleta", "taxa_boleta")):
        lotes[nome] = _media_ponderada(compras, chave, coluna, "qt")
    lotes = lotes.reset_index()
    lotes["qt_vendida"] = lotes["CD"].map(vendas).fillna(0.0)
    lotes["qt_aberta"] = lotes["qt_comprada"] - lotes["qt_vendida"]

    for cd in sorted(set(vendas.index) - set(lotes["CD"])):
        v = vcra[(vcra["op"] == "V") & (vcra["CD"] == cd)]
        _aviso(avisos, "Venda sem compra", v["codigo"].iloc[0],
               f"venda de {_br(v['qt'].sum(), 0)} sem a compra do lote no histórico da VCRA", cd)
    for r in lotes[lotes["qt_aberta"] < 0].itertuples():
        _aviso(avisos, "Lote negativo", r.codigo,
               f"vendas ({_br(r.qt_vendida, 0)}) maiores que as compras ({_br(r.qt_comprada, 0)})", r.CD)

    # grupo pelo indexador; no CDI, CDI+ tem spread de emissão e %CDI tem a emissão no IDDIR_PC
    lotes["grupo"] = lotes["codigo"].map(_grupo_por_ativo(vcra))
    lotes["boleta_ok"], lotes["taxa"], lotes["fonte_taxa"] = _taxa_usada(lotes, fonte)
    lotes["taxa_vcra_conv"] = _vcra_na_convencao(lotes, cdi_d)
    lotes["dif_boleta_vcra"] = (lotes["taxa_boleta"] - lotes["taxa_vcra_conv"]).where(lotes["boleta_ok"])

    for cd, (taxa, motivo) in correcoes.items():
        sel = (lotes["CD"] == cd) & (lotes["qt_aberta"] > 0)
        if not sel.any():
            _aviso(avisos, "Correção não aplicada", "", "o lote não está entre os lotes em aberto do fundo", cd)
            continue
        lotes.loc[sel, "taxa"] = taxa
        lotes.loc[sel, "fonte_taxa"] = "manual"

    abertos = lotes[lotes["qt_aberta"] > 0].reset_index(drop=True)
    consequencia = {"Lake": "", "VCRA": "usada a taxa da VCRA", "manual": "usada a correção manual",
                    "": "lote sem taxa, fora da média do ativo"}
    vale = {"Lake": "vale a taxa do Lake", "VCRA": "vale a taxa da VCRA", "manual": "vale a correção manual"}
    for r in abertos.itertuples():
        if pd.isna(r.taxa_boleta):
            _aviso(avisos, "Sem boleta", r.codigo, f"nenhuma boleta casada: {consequencia[r.fonte_taxa]}", r.CD)
        elif not r.boleta_ok:
            _aviso(avisos, "Boleta inválida", r.codigo, f"boleta com {_br(r.taxa_boleta, 4)} fora da convenção "
                   f"do {r.grupo}: {consequencia[r.fonte_taxa]}", r.CD)
        elif abs(r.dif_boleta_vcra) > DIVERGENCIA[r.grupo]:
            _aviso(avisos, "Boleta × VCRA", r.codigo, f"boleta {_br(r.taxa_boleta)} × VCRA {_br(r.taxa_vcra_conv)} "
                   f"(PU {_br(r.pu_compra)}); {vale.get(r.fonte_taxa, 'lote sem taxa')}", r.CD)
        if r.fonte_taxa == "manual":
            _aviso(avisos, "Correção manual", r.codigo, f"taxa {_br(r.taxa, 4)} ({correcoes[r.CD][1]})", r.CD)
    return abertos


def _valida(grupo, taxa):
    """Taxa na convenção do grupo: % do CDI (>= LIMITE_CONVENCAO) no %CDI; taxa ou spread % a.a. nos outros."""
    return bool(pd.notna(taxa) and (taxa >= LIMITE_CONVENCAO if grupo == "%CDI" else taxa < LIMITE_CONVENCAO))


def _num(valor):
    """Valor do cadastro como número (NaN se vazio ou não numérico)."""
    try:
        return float(valor)
    except (TypeError, ValueError):
        return np.nan


def _media_qt(b, coluna):
    return (b[coluna] * b["qt"]).sum() / b["qt"].sum() if len(b) else np.nan


def montar_fidc(posicao, boletas, cadastro, data, correcoes, avisos):
    """Cotas de FIDC da posição na data (código 5 + letra; não estão na VCRA), cada uma como um lote com a
    quantidade da posição e o próprio código no lugar do CD. Grupo pelo indexador do cadastro; taxa da boleta do Lake
    (compras até a data, exec_qt > 0, média ponderada pela quantidade) ou, sem boleta válida, a taxa_emissao do
    cadastro. correcoes = {codigo da cota: (taxa, motivo)}."""
    ref = pd.Timestamp(data)
    pos = posicao[(posicao["data"] == ref) & _e_fidc(posicao["codigo"]) & (posicao["quantidade"].abs() > 1e-9)]
    if pos.empty:
        return pd.DataFrame()
    indexadores = _vigente(cadastro, data, ["indexador"])
    taxas_cadastro = _vigente(cadastro, data, ["taxa_emissao", "taxa"])
    compras = (boletas[(boletas["data"] <= ref) & (boletas["qt"] > 0) & boletas["taxa"].notna()]
               if not boletas.empty else pd.DataFrame(columns=["codigo", "data", "qt", "pu", "taxa"]))
    linhas = []
    for codigo, g in pos.groupby("codigo"):
        b = compras[compras["codigo"] == codigo]
        indexador, taxa_cad, bruta = indexadores.get(codigo), _num(taxas_cadastro.get(codigo)), _media_qt(b, "taxa")
        grupo = _grupo_cadastro(indexador, taxa_cad if pd.notna(taxa_cad) else bruta) if _familia(indexador) else None
        ok = b[b["taxa"].map(lambda t: _valida(grupo, t)).astype(bool)] if grupo else b.iloc[0:0]
        l = {"codigo": codigo, "CD": codigo, "isin": g["isin"].dropna().iloc[0] if g["isin"].notna().any() else None,
             "tipo": "FIDC", "indexador": indexador, "grupo": grupo,
             "dt_compra": b["data"].min() if len(b) else pd.NaT,
             "perc_cdi": taxa_cad if grupo == "%CDI" else np.nan,
             "taxa_emissao": taxa_cad if grupo and grupo != "%CDI" else np.nan,
             "qt_comprada": np.nan, "qt_vendida": np.nan, "qt_aberta": g["quantidade"].sum(),
             "pu_compra": _media_qt(b, "pu"), "taxa_vcra": np.nan, "taxa_vcra_conv": np.nan, "dif_boleta_vcra": np.nan,
             "data_boleta": b["data"].max() if len(b) else pd.NaT,
             "casou_por": f"média de {len(ok)} boleta{'s' if len(ok) > 1 else ''}" if len(ok) else None,
             "taxa_boleta": _media_qt(ok, "taxa") if len(ok) else bruta, "boleta_ok": bool(len(ok)),
             "taxa": np.nan, "fonte_taxa": ""}
        boleta = "sem boleta do Lake" if b.empty else f"boleta do Lake com {_br(bruta, 4)}"
        if not grupo:
            _aviso(avisos, "FIDC sem indexador", codigo, f"sem indexador no cadastro (FUNDOS_CADASTRO_ATIVOS): fora "
                   f"da média ({boleta})")
            if codigo in correcoes:
                _aviso(avisos, "Correção não aplicada", codigo, "a cota não tem indexador no cadastro: o grupo da "
                       "taxa é desconhecido")
            linhas.append(l)
            continue
        if len(ok):
            l.update(taxa=l["taxa_boleta"], fonte_taxa="Lake")
        elif _valida(grupo, taxa_cad):
            l.update(taxa=taxa_cad, fonte_taxa="cadastro")
        if codigo in correcoes:
            l.update(taxa=correcoes[codigo][0], fonte_taxa="manual")
        sem_lake = boleta if b.empty else f"{boleta} fora da convenção do {grupo}"
        if l["fonte_taxa"] == "cadastro":
            _aviso(avisos, "FIDC pelo cadastro", codigo, f"{sem_lake}: taxa do cadastro, {_br(taxa_cad, 4)} "
                   f"({indexador})")
        elif not l["fonte_taxa"]:
            cad = (f"taxa {_br(taxa_cad, 4)} do cadastro fora da convenção" if pd.notna(taxa_cad)
                   else "sem taxa no cadastro")
            _aviso(avisos, "FIDC sem taxa", codigo, f"{sem_lake} e {cad}: fora da média")
        elif l["fonte_taxa"] == "manual":
            _aviso(avisos, "Correção manual", codigo, f"taxa {_br(l['taxa'], 4)} ({correcoes[codigo][1]})")
        linhas.append(l)
    return pd.DataFrame(linhas)


def _ordem_grupo(s):
    """Chave de ordenação de uma coluna na ordem dos GRUPOS (para sort_values(key=...)); as outras ficam como estão."""
    ordem = {g: i for i, (g, _) in enumerate(GRUPOS)}
    return s.map(ordem) if s.name == "grupo" else s


def montar_ativos(lotes, posicao, data, cdi, fora, avisos, ipca=None):
    """Taxa média por ativo (lotes ponderados pela quantidade), conferência com a quantidade da posição e
    financeiro pelo último PU com preço. fora: {codigo: motivo} dos ativos fora do escopo; ipca = IPCA projetado
    (fração) para o IPCA+ em CDI+ equivalente, ou None."""
    l = lotes.assign(sem_boleta=~lotes["boleta_ok"].astype(bool), manual=lotes["fonte_taxa"] == "manual",
                     sem_taxa=lotes["taxa"].isna())
    a = l.groupby("codigo").agg(
        isin=("isin", "first"), tipo=("tipo", "first"), grupo=("grupo", "first"), perc_cdi=("perc_cdi", "max"),
        taxa_emissao=("taxa_emissao", "max"), n_lotes=("CD", "count"), n_sem_boleta=("sem_boleta", "sum"),
        n_manual=("manual", "sum"), n_sem_taxa=("sem_taxa", "sum"), qt_lotes=("qt_aberta", "sum"))
    a["taxa_vcra"] = _media_ponderada(l, "codigo", "taxa_vcra_conv", "qt_aberta")
    a["taxa_media"] = _media_ponderada(l, "codigo", "taxa", "qt_aberta")

    ref = pd.Timestamp(data)
    fidc = _e_fidc(posicao["codigo"])  # cota de FIDC zerada não conta (as outras sem ISIN são caixa e over)
    na_data = posicao[(posicao["data"] == ref) & ((posicao["isin"].notna() & ~fidc)
                                                  | (fidc & (posicao["quantidade"].abs() > 1e-9)))]
    a = a.join(na_data.groupby("codigo")["quantidade"].sum().rename("qt_posicao"), how="outer")
    a["isin"] = a["isin"].fillna(na_data.groupby("codigo")["isin"].first())
    precos = (posicao[posicao["total"].notna()].groupby(["codigo", "data"])[["quantidade", "total"]].sum()
              .reset_index().sort_values("data").groupby("codigo").tail(1).set_index("codigo"))
    a["pu"] = (precos["total"] / precos["quantidade"].replace(0, np.nan)).reindex(a.index)
    a["dt_pu"] = precos["data"].reindex(a.index)
    a["dif_qt"] = a["qt_lotes"].fillna(0) - a["qt_posicao"].fillna(0)
    a["financeiro"] = a["qt_lotes"] * a["pu"]
    a["taxa_cdi_mais"] = _em_cdi_mais(a["grupo"], a["taxa_media"], cdi, ipca)
    a = a.rename_axis("codigo").reset_index().sort_values(["grupo", "codigo"], key=_ordem_grupo, na_position="last")

    for r in a.itertuples():
        if pd.isna(r.qt_lotes):
            if r.codigo in fora:
                _aviso(avisos, "Fora do escopo", r.codigo, f"{_br(r.qt_posicao, 0)} na posição; {fora[r.codigo]}")
            else:
                _aviso(avisos, "Sem lote", r.codigo, f"{_br(r.qt_posicao, 0)} na posição sem lote em aberto na VCRA")
            continue
        if abs(r.dif_qt) > 1e-6:
            pos = _br(r.qt_posicao, 0) if pd.notna(r.qt_posicao) else "nada"
            _aviso(avisos, "Quantidade", r.codigo, f"lotes {_br(r.qt_lotes, 0)} × posição {pos} em {ref:%d/%m/%Y}")
        if pd.isna(r.taxa_media):
            if r.tipo != "FIDC":  # a cota de FIDC sem taxa já tem o aviso do FIDC
                _aviso(avisos, "Fora da média", r.codigo, "nenhum lote com taxa")
        elif r.n_sem_taxa:
            _aviso(avisos, "Média parcial", r.codigo,
                   f"{int(r.n_sem_taxa)} de {int(r.n_lotes)} lotes sem taxa: média só dos outros")
        if pd.isna(r.pu):
            _aviso(avisos, "Fora da média", r.codigo, f"sem PU na posição nos {DIAS_PU} dias até {ref:%d/%m/%Y}")
    return a.reset_index(drop=True)


def _resumo(grupo, g):
    fin = g["financeiro"].sum()
    return {"grupo": grupo, "n_ativos": len(g), "financeiro": fin,
            "taxa_media": (g["financeiro"] * g["taxa_media"]).sum(min_count=1) / fin,
            "taxa_cdi_mais": (g["financeiro"] * g["taxa_cdi_mais"]).sum(min_count=1) / fin}


def consolidar(ativos):
    """Média por grupo (na unidade do grupo) e da carteira (em CDI+ equivalente), ponderadas pelo financeiro. A
    carteira junta os ativos com CDI+ equivalente: o IPCA+ só entra com o IPCA projetado (--ipca)."""
    ok = ativos[ativos["financeiro"].notna() & ativos["taxa_media"].notna()]
    linhas = [_resumo(grupo, ok[ok["grupo"] == grupo]) for grupo, _ in GRUPOS if (ok["grupo"] == grupo).any()]
    total = ok[ok["taxa_cdi_mais"].notna()]
    if len(total):
        linhas.append({**_resumo("CARTEIRA", total), "taxa_media": np.nan})
    return pd.DataFrame(linhas, columns=["grupo", "n_ativos", "financeiro", "taxa_media", "taxa_cdi_mais"])


def calcular(vcra, boletas, posicao, data, cdi, fonte, correcoes, tipos, cadastro=None, ipca=None):
    """Um fundo -> (lotes em aberto, ativos, carteira, avisos). cdi = CDI a.a. em fração; fonte = 'lake' ou
    'vcra'; correcoes = {CD ou código da cota de FIDC: (taxa, motivo)}; tipos = {RFTP_CD: incluir}; cadastro =
    versoes_cadastro (indexador e taxa das cotas de FIDC); ipca = IPCA projetado (fração) ou None."""
    vcra = _normalizar(vcra, textos=("CD", "codigo", "op", "isin", "tipo", "indexador"), datas=("data",),
                       numeros=("qt", "pu", "perc_cdi", "taxa_emissao", "taxa_vcra"))
    vcra["qt"] = vcra["qt"].abs()
    if not boletas.empty:
        boletas = _normalizar(boletas, textos=("codigo",), datas=("data",), numeros=("qt", "pu", "taxa"))
    posicao = _normalizar(posicao, textos=("codigo", "isin"), datas=("data",), numeros=("quantidade", "total"))

    avisos = []
    cdi_d = (1 + cdi) ** (1 / 252) - 1
    vcra, fora = filtrar_escopo(vcra, tipos, avisos)
    cotas = set(posicao.loc[_e_fidc(posicao["codigo"]) & (posicao["data"] == pd.Timestamp(data)), "codigo"])
    lotes = montar_lotes(vcra, boletas, fonte, {cd: c for cd, c in correcoes.items() if cd not in cotas}, cdi_d,
                         avisos)
    if tipos.get("FIDC", True):
        fidc = montar_fidc(posicao, boletas, cadastro, data, {cd: c for cd, c in correcoes.items() if cd in cotas},
                           avisos)
        if len(fidc):  # sem concat com tabela vazia (FutureWarning do pandas)
            lotes = fidc if lotes.empty else pd.concat([lotes, fidc], ignore_index=True)
    else:
        fora.update({c: "cota de FIDC, com FIDC = N na aba Tipos" for c in cotas})
    ativos = montar_ativos(lotes, posicao, data, cdi, fora, avisos, ipca)
    avisos.sort(key=lambda av: ORDEM_AVISOS.index(av["aviso"]))
    return lotes, ativos, consolidar(ativos), avisos


def resumo_da_data(data, cdi, lotes, ativos, carteira, avisos):
    """Linha das abas Fundos e Histórico: taxa e financeiro da carteira e de cada grupo, contagens e avisos por
    tipo (na ordem da lista, que já vem na ORDEM_AVISOS)."""
    g = carteira.set_index("grupo").reindex(["CARTEIRA"] + [grupo for grupo, _ in GRUPOS])
    if pd.notna(g.at["CARTEIRA", "financeiro"]):
        status = "OK"
    elif len(carteira):  # só grupos sem CDI+ equivalente: IPCA+ sem o --ipca
        status = "sem taxa da carteira: só IPCA+, que precisa do --ipca"
    else:
        status = "sem ativo com taxa e PU"
    linha = {"data": data, "status": status, "taxa_carteira": g.at["CARTEIRA", "taxa_cdi_mais"],
             "financeiro": g.at["CARTEIRA", "financeiro"]}
    for grupo, sufixo in GRUPOS:
        linha.update({f"taxa_{sufixo}": g.at[grupo, "taxa_media"], f"fin_{sufixo}": g.at[grupo, "financeiro"]})
    return {**linha, "n_lotes": len(lotes), "n_ativos": int(ativos["qt_lotes"].notna().sum()), "cdi": cdi * 100,
            "n_avisos": len(avisos),
            "avisos": " · ".join(f"{a} ({n})" for a, n in Counter(av["aviso"] for av in avisos).items())}


def calcular_trades(vcra, boletas, fonte, correcoes, tipos, cdi_do_dia, ipca=None):
    """Trades de um fundo: cada compra e venda da VCRA no escopo, com a boleta casada e a taxa pelas mesmas
    regras dos lotes. Relançamento da VCRA (venda do lote X e compra do lote X + "A", mesmo ativo, quantidade e
    PU, até JANELA_BOLETA dias depois) não é trade e fica fora; a correção do lote relançado vale para a compra
    original. As cotas de FIDC não estão na VCRA e ficam fora. cdi_do_dia(data) = CDI a.a. (fração) do dia do
    trade; ipca = IPCA projetado (fração) ou None. Devolve (trades, número de relançamentos)."""
    vcra = _normalizar(vcra, textos=("CD", "codigo", "op", "isin", "tipo", "indexador"), datas=("data",),
                       numeros=("qt", "pu", "perc_cdi", "taxa_emissao", "taxa_vcra"))
    vcra["qt"] = vcra["qt"].abs()
    if not boletas.empty:
        boletas = _normalizar(boletas, textos=("codigo",), datas=("data",), numeros=("qt", "pu", "taxa"))
    vcra = filtrar_escopo(vcra, tipos, [])[0].reset_index(drop=True)  # os avisos de escopo já saem no calcular

    pares = vcra[vcra["op"] == "V"].reset_index().merge(vcra[vcra["op"] == "C"].reset_index(), on=["codigo", "qt"],
                                                         suffixes=("_v", "_c"))
    pares = pares[(pares["CD_c"] == pares["CD_v"] + "A") & ((pares["pu_c"] - pares["pu_v"]).abs() < TOLERANCIA_PU)
                  & (pares["data_c"] - pares["data_v"]).dt.days.between(0, JANELA_BOLETA)]
    original = dict(zip(pares["CD_c"], pares["CD_v"]))

    trades = casar_boletas(vcra.drop(index=list(set(pares["index_v"]) | set(pares["index_c"]))).reset_index(drop=True),
                           boletas)
    trades["grupo"] = trades["codigo"].map(_grupo_por_ativo(vcra))
    trades["boleta_ok"], trades["taxa"], trades["fonte_taxa"] = _taxa_usada(trades, fonte)
    for cd, (taxa, _) in correcoes.items():
        while cd in original:  # lote relançado: a correção vale para a compra do lote original
            cd = original[cd]
        sel = (trades["op"] == "C") & (trades["CD"] == cd)
        trades.loc[sel, "taxa"] = taxa
        trades.loc[sel, "fonte_taxa"] = "manual"

    trades["cdi"] = trades["data"].map({dia: cdi_do_dia(dia.date()) for dia in trades["data"].drop_duplicates()})
    trades["taxa_vcra_conv"] = _vcra_na_convencao(trades, (1 + trades["cdi"]) ** (1 / 252) - 1)
    trades["taxa_cdi_mais"] = _em_cdi_mais(trades["grupo"], trades["taxa"], trades["cdi"], ipca)
    trades["financeiro"] = trades["qt"] * trades["pu"]
    trades["operacao"] = trades["op"].map({"C": "Compra", "V": "Venda"})
    return trades.sort_values(["data", "operacao", "codigo", "CD"]).reset_index(drop=True), len(pares)


def trades_por_dia(trades):
    """Uma linha por data e operação (compra ou venda): trades, ativos, financeiro e taxa média ponderada pelo
    financeiro de cada trade, por grupo (na unidade do grupo) e no total (em CDI+ equivalente, com o CDI do dia)."""
    chave = ["data", "operacao"]
    dia = trades.groupby(chave).agg(n_trades=("CD", "size"), n_ativos=("codigo", "nunique"),
                                    financeiro=("financeiro", "sum"), cdi=("cdi", "first"))
    dia["taxa_total"] = _media_ponderada(trades, chave, "taxa_cdi_mais", "financeiro")
    dia["fin_total"] = trades[trades["taxa_cdi_mais"].notna()].groupby(chave)["financeiro"].sum()
    for grupo, sufixo in GRUPOS:
        g = trades[(trades["grupo"] == grupo) & trades["taxa"].notna()]
        dia[f"taxa_{sufixo}"] = _media_ponderada(g, chave, "taxa", "financeiro")
        dia[f"fin_{sufixo}"] = g.groupby(chave)["financeiro"].sum()
    sem_taxa = trades[trades["taxa"].isna()].groupby(chave)["financeiro"]
    dia["n_sem_taxa"] = sem_taxa.size().reindex(dia.index, fill_value=0)
    dia["fin_sem_taxa"] = sem_taxa.sum()
    dia["cdi"] = dia["cdi"] * 100
    return dia.reset_index()


# ── Saída ────────────────────────────────────────────────────────────────────
TAXA, PU, QTD, VALOR, DATA = "0.0000", "#,##0.000000", "#,##0", "#,##0.00", "dd/mm/yyyy"
FUNDO = ("fundo", "Fundo", None)
# taxa média e financeiro de cada grupo, nas abas Histórico e Trades por dia
_UNIDADE = {"%CDI": "taxa média (% do CDI)", "CDI+": "spread médio", "SELIC+": "spread médio",
            "Pré": "taxa média (% a.a.)", "IPCA+": "spread médio"}
POR_GRUPO = [c for grupo, s in GRUPOS for c in ((f"taxa_{s}", f"{grupo}: {_UNIDADE[grupo]}", TAXA),
                                                (f"fin_{s}", f"{grupo}: financeiro (R$)", VALOR))]
COLUNAS = {
    "Histórico": [FUNDO, ("data", "Data", DATA), ("status", "Status", None),
                  ("taxa_carteira", "Carteira em CDI+ equivalente", TAXA), ("financeiro", "Financeiro (R$)", VALOR),
                  *POR_GRUPO,
                  ("n_ativos", "Ativos", "0"), ("n_lotes", "Lotes em aberto", "0"), ("cdi", "CDI (% a.a.)", TAXA),
                  ("n_avisos", "Avisos", "0"), ("avisos", "Avisos por tipo", None)],
    "Carteira": [FUNDO, ("data", "Data", DATA), ("grupo", "Grupo", None), ("n_ativos", "Ativos", "0"),
                 ("financeiro", "Financeiro (R$)", VALOR),
                 ("taxa_media", "Taxa média (na unidade do grupo)", TAXA),
                 ("taxa_cdi_mais", "Taxa média em CDI+ equivalente", TAXA)],
    "Ativos": [FUNDO, ("codigo", "Ativo", None), ("isin", "ISIN", None), ("tipo", "Tipo (RFTP_CD)", None),
               ("grupo", "Grupo", None), ("perc_cdi", "% do CDI (emissão)", TAXA),
               ("taxa_emissao", "Taxa de emissão", TAXA), ("n_lotes", "Lotes", "0"),
               ("n_sem_boleta", "Lotes sem boleta válida", "0"), ("n_manual", "Lotes corrigidos", "0"),
               ("qt_lotes", "Qtd nos lotes", QTD), ("qt_posicao", "Qtd na posição", QTD), ("dif_qt", "Diferença", QTD),
               ("taxa_vcra", "Taxa média VCRA (convenção do grupo)", TAXA), ("taxa_media", "Taxa média", TAXA),
               ("taxa_cdi_mais", "Taxa média em CDI+ equivalente", TAXA), ("pu", "PU", PU),
               ("dt_pu", "Data do PU", DATA), ("financeiro", "Financeiro (R$)", VALOR)],
    "Lotes": [FUNDO, ("codigo", "Ativo", None), ("isin", "ISIN", None), ("tipo", "Tipo (RFTP_CD)", None),
              ("indexador", "Indexador", None), ("grupo", "Grupo", None), ("CD", "Lote", None),
              ("dt_compra", "Compra", DATA), ("qt_comprada", "Qtd comprada", QTD), ("qt_vendida", "Qtd vendida", QTD),
              ("qt_aberta", "Qtd em aberto", QTD), ("pu_compra", "PU de compra", PU),
              ("perc_cdi", "% do CDI (emissão)", TAXA), ("taxa_emissao", "Taxa de emissão", TAXA),
              ("taxa_vcra", "Taxa VCRA (bruta)", TAXA), ("taxa_vcra_conv", "Taxa VCRA (convenção do grupo)", TAXA),
              ("data_boleta", "Data da boleta", DATA), ("casou_por", "Boleta casou por", None),
              ("taxa_boleta", "Taxa da boleta", TAXA), ("boleta_ok", "Boleta válida", None),
              ("dif_boleta_vcra", "Boleta − VCRA", TAXA), ("taxa", "Taxa usada", TAXA),
              ("fonte_taxa", "Fonte da taxa", None)],
    "Trades por dia": [FUNDO, ("data", "Data", DATA), ("operacao", "Operação", None), ("n_trades", "Trades", "0"),
                       ("n_ativos", "Ativos", "0"), ("financeiro", "Financeiro (R$)", VALOR),
                       ("taxa_total", "Taxa média em CDI+ equivalente", TAXA),
                       ("fin_total", "Financeiro na taxa média (R$)", VALOR), *POR_GRUPO,
                       ("n_sem_taxa", "Sem taxa: trades", "0"), ("fin_sem_taxa", "Sem taxa: financeiro (R$)", VALOR),
                       ("cdi", "CDI (% a.a.)", TAXA)],
    "Trades": [FUNDO, ("data", "Data", DATA), ("operacao", "Operação", None), ("codigo", "Ativo", None),
               ("isin", "ISIN", None), ("tipo", "Tipo (RFTP_CD)", None), ("indexador", "Indexador", None),
               ("grupo", "Grupo", None), ("CD", "Lote", None), ("qt", "Quantidade", QTD), ("pu", "PU", PU),
               ("financeiro", "Financeiro (R$)", VALOR), ("perc_cdi", "% do CDI (emissão)", TAXA),
               ("taxa_emissao", "Taxa de emissão", TAXA), ("taxa_vcra", "Taxa VCRA (bruta)", TAXA),
               ("taxa_vcra_conv", "Taxa VCRA (convenção do grupo)", TAXA), ("data_boleta", "Data da boleta", DATA),
               ("casou_por", "Boleta casou por", None), ("taxa_boleta", "Taxa da boleta", TAXA),
               ("boleta_ok", "Boleta válida", None), ("taxa", "Taxa usada", TAXA), ("fonte_taxa", "Fonte da taxa", None),
               ("taxa_cdi_mais", "Taxa em CDI+ equivalente", TAXA)],
    "Avisos": [FUNDO, ("aviso", "Aviso", None), ("ativo", "Ativo", None), ("lote", "Lote", None),
               ("detalhe", "Detalhe", None)],
    "Cadastro": [("codigo", "Ativo", None), ("campo", "Campo", None), ("valor", "Valor", None),
                 ("desde", "Desde", DATA), ("fonte", "Fonte", None)],
    "Fundos": [FUNDO, ("nome", "Nome", None), ("data", "Data", DATA), ("status", "Status", None),
               ("taxa_carteira", "Carteira em CDI+ equivalente", TAXA), ("financeiro", "Financeiro (R$)", VALOR),
               ("n_lotes", "Lotes em aberto", "0"), ("n_ativos", "Ativos", "0"), ("n_avisos", "Avisos", "0"),
               ("cdi", "CDI (% a.a.)", TAXA), ("cdi_origem", "Origem do CDI", None)],
    "Parâmetros": [("parametro", "Parâmetro", None), ("valor", "Valor", None)],
}
# títulos de colunas de versões anteriores, para o --html ler um .xlsx gravado antes
TITULOS_ANTIGOS = {"Taxa média (CDI+: spread; %CDI: % do CDI)": "taxa_media", "Spread de emissão": "taxa_emissao"}


def salvar_xlsx(caminho, tabelas, notas):
    """tabelas: {aba: DataFrame} com as abas de COLUNAS (a que não vier fica de fora, como o Histórico sem
    --desde); as notas vão embaixo da aba Carteira."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        for aba, colunas in COLUNAS.items():
            if aba not in tabelas:
                continue
            df = tabelas[aba].reindex(columns=[c for c, _, _ in colunas])
            df.rename(columns={c: t for c, t, _ in colunas}).to_excel(xw, sheet_name=aba, index=False)
            ws = xw.book[aba]
            ws.freeze_panes = "B2"
            for j, (c, titulo, fmt) in enumerate(colunas, start=1):
                ws.cell(row=1, column=j).font = Font(bold=True)
                largura = 14 if fmt else max([len(str(v)) for v in df[c].dropna()] + [0])
                ws.column_dimensions[get_column_letter(j)].width = min(max(len(titulo), largura) + 2, 90)
                if fmt:
                    for (cel,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        cel.number_format = fmt
        ws = xw.book["Carteira"]
        ws.append([])
        for nota in notas:
            ws.append([nota])


def _juntar(partes):
    """{aba: [DataFrame de cada fundo]} -> {aba: DataFrame}, sem as partes vazias."""
    return {aba: pd.concat([p for p in lista if len(p)], ignore_index=True) if any(len(p) for p in lista)
            else pd.DataFrame() for aba, lista in partes.items()}


def gravar_diario(pasta, diario, trades_fundo, cadastro, parametros, notas, html_report):
    """Um relatório HTML por data em pasta/aaaammdd.html (o mesmo layout, com os fundos que têm a data): carteira,
    ativos, lotes e avisos da data, os trades desde a data anterior da série do fundo e o cadastro vigente na data.
    diario e trades_fundo: como o main monta com --diario. O index.html é o html_report.gerar_indice, que lê a pasta
    inteira."""
    os.makedirs(pasta, exist_ok=True)
    for dia in sorted(diario):
        partes = {aba: [] for aba in ("Carteira", "Ativos", "Lotes", "Trades por dia", "Trades", "Avisos")}
        linhas, inicio = [], pd.Timestamp(dia)
        for fundo, (linha, resultado) in diario[dia].items():
            linhas.append(linha)
            if resultado is None:  # erro na data: o fundo sai só com a linha da aba Fundos
                continue
            lotes, ativos, carteira, avisos = resultado
            partes["Carteira"].append(carteira.assign(fundo=fundo, data=pd.Timestamp(dia)))
            partes["Ativos"].append(ativos.assign(fundo=fundo))
            partes["Lotes"].append(lotes.assign(fundo=fundo))
            partes["Avisos"].append(pd.DataFrame(avisos, columns=["aviso", "ativo", "lote", "detalhe"]).assign(fundo=fundo))
            if fundo in trades_fundo:
                trades, datas = trades_fundo[fundo]
                anterior = max((x for x in datas if x < dia), default=None)
                desde = pd.Timestamp(anterior) + pd.Timedelta(days=1) if anterior else pd.Timestamp(dia)
                janela = trades[trades["data"].between(desde, pd.Timestamp(dia))]
                inicio = min(inicio, desde)
                partes["Trades"].append(janela.assign(fundo=fundo))
                partes["Trades por dia"].append(trades_por_dia(janela).assign(fundo=fundo))
        tabelas = _juntar(partes)
        tabelas["Fundos"] = pd.DataFrame(linhas)
        tabelas["Cadastro"] = cadastro[cadastro["desde"] <= pd.Timestamp(dia)] if len(cadastro) else cadastro
        periodo = f"em {dia:%d/%m/%Y}" if inicio == pd.Timestamp(dia) else f"de {inicio:%d/%m/%Y} a {dia:%d/%m/%Y}"
        param = parametros.copy()
        param.loc[param["parametro"] == "Data de referência", "valor"] = f"{dia:%d/%m/%Y} (relatório do dia)"
        tabelas["Parâmetros"] = pd.concat([param, pd.DataFrame([("Trades deste relatório", periodo)],
                                                               columns=param.columns)], ignore_index=True)
        html_report.gerar_html(os.path.join(pasta, f"{dia:%Y%m%d}.html"), tabelas, notas, dia=dia, trades_desde=inicio)
    return len(diario)


def ler_saida(caminho):
    """Saída .xlsx já gravada -> (tabelas, notas) com os nomes internos das colunas, para regerar o HTML sem o banco.
    As notas ficam embaixo da aba Carteira, na coluna Fundo."""
    abas = pd.read_excel(caminho, sheet_name=None)
    tabelas = {aba: abas[aba].rename(columns={**TITULOS_ANTIGOS, **{t: c for c, t, _ in colunas}})
               for aba, colunas in COLUNAS.items() if aba in abas}
    carteira = tabelas.get("Carteira", pd.DataFrame(columns=["fundo", "grupo"]))
    notas = carteira.loc[carteira["grupo"].isna(), "fundo"].dropna().astype(str).tolist()
    tabelas["Carteira"] = carteira[carteira["grupo"].notna()].reset_index(drop=True)
    return tabelas, notas


def _taxa_grupo(grupo, x):
    """Taxa na unidade do grupo: 105,00% do CDI, CDI+1,50, SELIC+0,05, 13,20% a.a. ou IPCA+7,10."""
    if grupo == "%CDI":
        return f"{_br(x)}% do CDI"
    if grupo == "Pré":
        return f"{_br(x)}% a.a."
    return f"{grupo[:-1]}{'+' if x >= 0 else '−'}{_br(abs(x))}"


def imprimir_fundo(lotes, ativos, carteira, avisos, data, cdi, ipca=None):
    fidc = (lotes["tipo"] == "FIDC").sum() if len(lotes) else 0
    vcra = lotes[lotes["tipo"] != "FIDC"] if len(lotes) else lotes
    print(f"[INFO] {len(vcra)} lotes em aberto em {vcra['codigo'].nunique()} ativos da VCRA"
          f"{f' e {fidc} cota(s) de FIDC' if fidc else ''} | boleta casada em {vcra['taxa_boleta'].notna().sum()} "
          f"({(vcra['casou_por'] == 'PU').sum()} pelo PU)")
    if len(lotes) and (ativos["dif_qt"].abs() <= 1e-6).all():
        print(f"[OK] quantidade dos lotes = posição de {data:%d/%m/%Y} nos {len(ativos)} ativos")
    for av in avisos:
        alvo = " ".join(x for x in (av["ativo"], f"({av['lote']})" if av["lote"] else "") if x)
        print(f"[{'INFO' if av['aviso'] in INFORMATIVOS else 'AVISO'}] {av['aviso']}"
              f"{f' — {alvo}' if alvo else ''}: {av['detalhe']}")
    for r in carteira.itertuples():
        if r.grupo == "CARTEIRA":
            ipca_txt = f", IPCA {_br(ipca * 100)}%" if ipca is not None and (carteira["grupo"] == "IPCA+").any() else ""
            taxa = f"{_cdi_mais(r.taxa_cdi_mais)} em CDI+ equivalente (CDI {_br(cdi * 100)}% a.a.{ipca_txt})"
        elif r.grupo == "CDI+":
            taxa = _cdi_mais(r.taxa_media)
        elif pd.notna(r.taxa_cdi_mais):
            taxa = f"{_taxa_grupo(r.grupo, r.taxa_media)} ({_cdi_mais(r.taxa_cdi_mais)} equivalente)"
        else:
            taxa = f"{_taxa_grupo(r.grupo, r.taxa_media)} (fora da carteira: sem --ipca)"
        nome = "Carteira" if r.grupo == "CARTEIRA" else r.grupo
        print(f"[OK] {nome}: {r.n_ativos} ativo{'' if r.n_ativos == 1 else 's'}, R$ {_br(r.financeiro, 0)} — {taxa}")


def imprimir_trades(trades, relancamentos):
    """Linha do console com os trades do fundo no período: compras e vendas, financeiro e taxa média."""
    if trades.empty:
        print("[INFO] nenhum trade no período")
        return
    partes = []
    for operacao, g in trades.groupby("operacao"):
        ok = g[g["taxa_cdi_mais"].notna()]
        taxa = (ok["financeiro"] * ok["taxa_cdi_mais"]).sum() / ok["financeiro"].sum() if len(ok) else np.nan
        partes.append(f"{len(g)} {operacao.lower()}{'s' if len(g) > 1 else ''} (R$ {_br(g['financeiro'].sum(), 0)}"
                      f"{f', {_cdi_mais(taxa)} equivalente' if pd.notna(taxa) else ''})")
    fora = (f"; {relancamentos} relançamento{'s' if relancamentos > 1 else ''} da VCRA "
            f"{'ficaram' if relancamentos > 1 else 'ficou'} fora") if relancamentos else ""
    print(f"[INFO] trades de {trades['data'].min():%d/%m/%Y} a {trades['data'].max():%d/%m/%Y}, em "
          f"{trades['data'].nunique()} dias: {' e '.join(partes)}{fora}")


def _situacao(r):
    """(etiqueta, texto) de uma linha das abas Fundos ou Histórico, para o console."""
    if r["status"] == "OK":
        n = r["n_avisos"]
        return "OK", (f"carteira {_cdi_mais(r['taxa_carteira'])} ({r['n_ativos']} ativos, "
                      f"R$ {_br(r['financeiro'], 0)}; {n} aviso{'' if n == 1 else 's'})")
    if r["status"].startswith("erro: "):
        return "ERRO", r["status"][len("erro: "):]
    return "AVISO", r["status"]


def _data(s):
    return datetime.strptime(s, "%d/%m/%Y").date()


def _desde(s):
    if _chave(s) == "inicio":
        return "inicio"
    try:
        return _data(s)
    except ValueError:
        raise argparse.ArgumentTypeError(f"{s!r}: use dd/mm/aaaa ou inicio") from None


def _taxa_aa(s):
    """Taxa em % a.a. da linha de comando, com ponto ou vírgula (13.65 ou 4,5)."""
    try:
        return float(s.replace(",", "."))
    except ValueError:
        raise argparse.ArgumentTypeError(f"{s!r}: use % a.a., ex.: 4,5") from None


def main(argv=None):
    for _s in (sys.stdout, sys.stderr):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa
            pass

    ap = argparse.ArgumentParser(description="Taxa média de aquisição dos papéis de renda fixa (CDI+, %CDI, SELIC+, "
                                             "pré e IPCA+) e das cotas de FIDC de uma lista de fundos: lotes da "
                                             "VCRAOPRF, taxa da boleta do Lake (FUNDOS_OPERACOES) e financeiro pela "
                                             "VW_FUNDOS_POSICAO.")
    ap.add_argument("--entrada", default=ENTRADA_PADRAO,
                    help="Excel com as abas Fundos, Tipos e Correções (padrão: input/fundos.xlsx)")
    ap.add_argument("--fundo", nargs="+", metavar="FUNDO",
                    help="só estes fundos, no lugar da aba Fundos (tipos e correções continuam vindo do Excel)")
    ap.add_argument("--data", type=_data,
                    help="data de referência dd/mm/aaaa (padrão: última data da posição de cada fundo)")
    ap.add_argument("--desde", nargs="?", const="inicio", type=_desde, metavar="DATA",
                    help="série: recalcula a carteira em cada data desde dd/mm/aaaa ou, sem data, desde o início do "
                         "fundo, e grava a aba Histórico")
    ap.add_argument("--frequencia", type=_chave, choices=("mensal", "diaria"),
                    help="datas da série: mensal (última data com posição de cada mês, padrão) ou diaria")
    ap.add_argument("--fonte", type=lambda s: "lake" if _chave(s) == "boleta" else _chave(s), choices=("lake", "vcra"),
                    default="lake", help="taxa de CDI+, SELIC+, pré e IPCA+: lake (a boleta do Lake, negociada; padrão) "
                                         "ou vcra (pelo PU pago); %%CDI usa sempre o Lake. Aceita também boleta, o "
                                         "nome antigo do lake")
    ap.add_argument("--cdi", type=_taxa_aa, help="CDI em %% a.a. para levar %%CDI e pré ao CDI+ equivalente "
                                                  "(padrão: CDIAcumulado da FUNDOS_DADOS_INDICES na data de cada fundo)")
    ap.add_argument("--ipca", type=_taxa_aa, help="IPCA projetado em %% a.a. (ex.: 4,5) para levar o IPCA+ ao CDI+ "
                                                   "equivalente e incluí-lo na carteira; sem ele, o IPCA+ sai à parte")
    ap.add_argument("--html", metavar="XLSX", help="só regera o relatório HTML de uma saída .xlsx já gravada, sem "
                                                   "consultar o banco (grava o .html ao lado do .xlsx)")
    ap.add_argument("--diario", action="store_true",
                    help="grava também um relatório HTML por data em output/diario/aaaammdd.html e o output/index.html, "
                         "com o calendário de todos os dias da pasta; com --desde, a série passa a ser diária")
    args = ap.parse_args(argv)
    sys.path.insert(0, AQUI)
    import html_report
    if args.html:
        try:
            tabelas, notas = ler_saida(args.html)
        except (OSError, ValueError) as e:
            ap.error(f"--html: {e}")
        print(f"[OK] {html_report.gerar_html(os.path.splitext(args.html)[0] + '.html', tabelas, notas)}")
        return
    if args.cdi is not None and not 1 < args.cdi < 50:
        ap.error("--cdi em % a.a. (ex.: 13.65)")
    if args.ipca is not None and not -5 < args.ipca < 50:
        ap.error("--ipca em % a.a. (ex.: 4,5)")
    ipca = args.ipca / 100 if args.ipca is not None else None
    for fundo in args.fundo or []:
        if not re.fullmatch(r"\w+", fundo):
            ap.error(f"--fundo inválido: {fundo!r}")
    if args.frequencia and not args.desde:
        ap.error("--frequencia só vale com --desde")
    if args.desde not in (None, "inicio") and args.data and args.desde > args.data:
        ap.error("--desde depois de --data")
    args.frequencia = args.frequencia or ("diaria" if args.diario else "mensal")
    serie = "" if not args.desde else f"{args.frequencia}, desde " + (
        "o início do fundo" if args.desde == "inicio" else f"{args.desde:%d/%m/%Y}")

    if os.path.exists(args.entrada):
        try:
            fundos, tipos, correcoes = carregar_entrada(args.entrada)
        except (OSError, ValueError) as e:
            ap.error(str(e))
    elif args.fundo:
        fundos, tipos, correcoes = [], dict(TIPOS_PADRAO), {}
    else:
        ap.error(f"não achei {args.entrada}: crie a lista de fundos ou use --fundo")
    if args.fundo:
        nomes = dict(fundos)
        fundos = [(f, nomes.get(f, "")) for f in dict.fromkeys(args.fundo)]
    if not fundos:
        ap.error(f"nenhum fundo para rodar na aba Fundos de {args.entrada}")
    print(f"[INFO] {len(fundos)} fundo(s) | tipos incluídos: {', '.join(t for t, s in tipos.items() if s)} | "
          f"taxa (menos %CDI): {'Lake' if args.fonte == 'lake' else 'VCRA'} | IPCA+ "
          f"{f'com IPCA de {_br(args.ipca)}% a.a.' if ipca is not None else 'à parte (sem --ipca)'}"
          f"{f' | série {serie}' if serie else ''}")

    from sql_connection import disconnect, get_connection, query
    try:
        get_connection()
    except Exception as e:  # noqa — credenciais ou driver
        sys.exit(f"[ERRO] conexão com o banco: {e}")

    partes = {aba: [] for aba in ("Carteira", "Ativos", "Lotes", "Trades por dia", "Trades", "Avisos")}
    resumo, historico, cdis = [], [], None
    # com --diario: {data: {fundo: (linha da aba Fundos, (lotes, ativos, carteira, avisos) ou None se deu erro)}} e
    # {fundo: (trades, datas da série)}, para o relatório de cada data
    diario, trades_fundo = {}, {}
    cadastros = []  # versões do cadastro de cada fundo (FUNDOS_CADASTRO_ATIVOS), até a data do fundo

    def cdi_do_dia(dia):  # para os trades: sem CDI até o dia, o CDI+ equivalente dos %CDI fica vazio
        try:
            return _cdi(args.cdi, cdis, dia)[0]
        except RuntimeError:
            return np.nan

    try:
        for fundo, nome in fundos:
            print(f"\n── {fundo}{f' ({nome})' if nome else ''} {'─' * 40}")
            info = {"fundo": fundo, "nome": nome}
            resumo.append(info)
            corr = correcoes.get(fundo, {})
            try:
                d = carregar_fundo(query, fundo, args.data, tipos, args.desde)
                if args.cdi is None and cdis is None:
                    cdis = carregar_cdi(query)
                if d["cadastro"].empty:
                    print("[AVISO] nenhum ativo na FUNDOS_CADASTRO_ATIVOS: a ficha do ativo sai sem o cadastro")
                dias = datas_da_serie(d, args.frequencia)
                if args.desde:
                    print(f"[INFO] série {args.frequencia}: {len(dias)} datas, de {dias[0]:%d/%m/%Y} a "
                          f"{dias[-1]:%d/%m/%Y}")
                for dia in dias[:-1]:  # a última é a data de referência, calculada abaixo com o detalhe
                    resultado, origem_dia = None, ""
                    try:
                        vcra, boletas, posicao = fatiar(d, dia)
                        # lote que ainda não existia na data não conta como correção não aplicada (na cota de FIDC,
                        # o lote é o código da cota)
                        lancados = set(_strip(vcra["CD"])) | set(_strip(posicao["codigo"]).dropna())
                        cdi, origem_dia = _cdi(args.cdi, cdis, dia)
                        resultado = calcular(vcra, boletas, posicao, dia, cdi, args.fonte,
                                             {cd: c for cd, c in corr.items() if cd in lancados}, tipos,
                                             d["cadastro"], ipca)
                        r = resumo_da_data(dia, cdi, *resultado)
                    except Exception as e:  # noqa — erro numa data da série não para as outras
                        r = {"data": dia, "status": f"erro: {e}"}
                    historico.append({"fundo": fundo, **r})
                    if args.diario:
                        diario.setdefault(dia, {})[fundo] = ({"fundo": fundo, "nome": nome, **r,
                                                              "cdi_origem": origem_dia}, resultado)
                    etiqueta, texto = _situacao(r)
                    print(f"[{etiqueta}] {dia:%d/%m/%Y}: {texto}")
                cdi, origem = _cdi(args.cdi, cdis, d["data"])
                print(f"[INFO] {d['data']:%d/%m/%Y} | CDI {_br(cdi * 100)}% a.a. ({origem}) | "
                      f"{len(corr)} correção(ões) manual(is)")
                if d["boletas"].empty:
                    print("[AVISO] nenhuma boleta na FUNDOS_OPERACOES: taxa da VCRA, %CDI sem taxa e FIDC pelo "
                          "cadastro")
                lotes, ativos, carteira, avisos = calcular(*fatiar(d, d["data"]), d["data"], cdi, args.fonte, corr,
                                                           tipos, d["cadastro"], ipca)
                imprimir_fundo(lotes, ativos, carteira, avisos, d["data"], cdi, ipca)
                trades, relancamentos = calcular_trades(*fatiar(d, d["data"])[:2], args.fonte, corr, tipos, cdi_do_dia,
                                                        ipca)
                if args.desde:
                    trades = trades[trades["data"] >= pd.Timestamp(d["desde"])]
                por_dia = trades_por_dia(trades)
                imprimir_trades(trades, relancamentos)
            except Exception as e:  # noqa — erro num fundo (dado faltando, consulta) não para os outros
                print(f"[ERRO] {fundo}: {e}")
                info["status"] = f"erro: {e}"
                continue
            info.update(resumo_da_data(d["data"], cdi, lotes, ativos, carteira, avisos), cdi_origem=origem)
            if args.desde:
                historico.append(info)
            if args.diario:
                diario.setdefault(d["data"], {})[fundo] = (dict(info), (lotes, ativos, carteira, avisos))
                trades_fundo[fundo] = (trades, dias)
            partes["Carteira"].append(carteira.assign(fundo=fundo, data=pd.Timestamp(d["data"])))
            partes["Ativos"].append(ativos.assign(fundo=fundo))
            partes["Lotes"].append(lotes.assign(fundo=fundo))
            partes["Trades por dia"].append(por_dia.assign(fundo=fundo))
            partes["Trades"].append(trades.assign(fundo=fundo))
            partes["Avisos"].append(pd.DataFrame(avisos, columns=["aviso", "ativo", "lote", "detalhe"]).assign(fundo=fundo))
            cadastros.append(d["cadastro"])
    finally:
        disconnect()

    tabelas = _juntar(partes)
    if len(tabelas["Lotes"]):
        tabelas["Lotes"] = tabelas["Lotes"].sort_values(["fundo", "grupo", "codigo", "dt_compra"], key=_ordem_grupo,
                                                        na_position="last").assign(
            boleta_ok=lambda x: x["boleta_ok"].map({True: "sim", False: "não"}))
    if len(tabelas["Trades"]):
        tabelas["Trades"] = tabelas["Trades"].assign(boleta_ok=lambda x: x["boleta_ok"].map({True: "sim", False: "não"}))
    tabelas["Fundos"] = pd.DataFrame(resumo)
    tabelas["Cadastro"] = (pd.concat(cadastros, ignore_index=True).drop_duplicates(["codigo", "campo", "desde"])
                           .sort_values(["codigo", "campo", "desde"]).reset_index(drop=True)
                           if cadastros else versoes_cadastro(pd.DataFrame()))
    if args.desde:
        tabelas["Histórico"] = pd.DataFrame(historico)
    datas = [r["data"] for r in resumo if "data" in r]
    ref = args.data or (max(datas) if datas else datetime.now().date())
    tabelas["Parâmetros"] = pd.DataFrame([
        ("Entrada", args.entrada if os.path.exists(args.entrada) else "(sem arquivo: --fundo e tipos padrão)"),
        ("Fundos", ", ".join(f for f, _ in fundos)),
        ("Data de referência", f"{args.data:%d/%m/%Y}" if args.data else "última data da posição de cada fundo"),
        ("Série (aba Histórico)", serie or "não: só a data de referência"),
        ("Taxa (menos %CDI)", "Lake (boleta negociada)" if args.fonte == "lake" else "VCRA (pelo PU pago)"),
        ("Taxa dos %CDI", "Lake (boleta, % do CDI)"),
        ("CDI para %CDI e pré → CDI+", f"{_br(args.cdi, 4)}% a.a. (--cdi)" if args.cdi is not None
         else "CDIAcumulado na data de cada fundo (aba Fundos)"),
        ("IPCA para IPCA+ → CDI+", f"{_br(args.ipca, 4)}% a.a. (--ipca): IPCA+ na carteira" if ipca is not None
         else "não informado (--ipca): IPCA+ à parte, fora da carteira"),
        ("Tipos incluídos", ", ".join(t for t, s in tipos.items() if s)),
        ("Tipos excluídos", ", ".join(t for t, s in tipos.items() if not s)),
        ("Grupos", "pelo indexador: CDI → %CDI ou CDI+ (com spread de emissão), SELIC → SELIC+, pré → Pré, IPCA → "
                   "IPCA+; outros indexadores (IGP-M, dólar...) ficam fora"),
        ("Cotas de FIDC", "fora (FIDC = N na aba Tipos)" if not tipos.get("FIDC", True) else
         "códigos 5 + letra da posição; grupo pelo indexador do cadastro; taxa das compras no Lake (média pela "
         "quantidade) ou, sem elas, a taxa_emissao do cadastro"),
        ("SELIC+ e CDI+", "a SELIC conta como igual ao CDI: o spread sobre a SELIC entra como CDI+"),
        ("Casamento com a boleta", f"mesmo ativo, até {JANELA_BOLETA} dias, mesma quantidade ou PU "
                                   f"(tolerância {TOLERANCIA_PU})"),
        ("Boleta válida", f"%CDI >= {LIMITE_CONVENCAO} (% do CDI); demais grupos < {LIMITE_CONVENCAO} (taxa ou "
                          f"spread % a.a.)"),
        ("Trades", f"compras e vendas da VCRA no escopo, {'desde o início da série' if args.desde else 'todo o histórico'} "
                   f"até a data; relançamentos (lote X + XA) e cotas de FIDC ficam fora; taxa pelas regras dos lotes"),
        ("Correções manuais", f"{sum(map(len, correcoes.values()))} lote(s) com aplicar = S na aba Correções"),
        ("Gerado em", f"{datetime.now():%d/%m/%Y %H:%M}"),
    ], columns=["parametro", "valor"])
    notas = ["CDI+: spread sobre 100% do CDI (% a.a.). %CDI: % do CDI. SELIC+: spread sobre a SELIC. Pré: taxa "
             "% a.a. IPCA+: spread sobre o IPCA.",
             "CDI+ equivalente, com CDI flat: %CDI vira o spread que rende o mesmo em um ano; SELIC+ entra como CDI+; "
             "pré = (1 + taxa) / (1 + CDI) - 1; IPCA+ = (1 + IPCA) x (1 + spread) / (1 + CDI) - 1, só com o --ipca "
             "(sem ele, o IPCA+ sai à parte e fora da carteira). O CDI usado em cada fundo está na aba Fundos.",
             "Médias ponderadas pelo financeiro (qtd em aberto x último PU da posição; na cota de FIDC, qtd x valor "
             "da cota); ativo sem taxa ou sem PU fica de fora (aba Avisos).",
             "Trades por dia: média ponderada pelo financeiro de cada trade (qtd x PU da operação); %CDI e pré "
             "convertidos com o CDI do dia do trade; trade sem taxa fica de fora."]
    if args.desde:
        notas.append(f"Histórico: o mesmo cálculo em cada data, só com o que existia nela (VCRA e boletas até a "
                     f"data, PU dos {DIAS_PU} dias anteriores, CDI do dia). Carteira, Ativos, Lotes e Avisos são da "
                     f"última data.")

    os.makedirs(SAIDA_DIR, exist_ok=True)
    prefixo = "taxa_media_historico" if args.desde else "taxa_media"
    arquivo = f"{prefixo}_{fundos[0][0]}_{ref:%Y%m%d}.xlsx" if len(fundos) == 1 else f"{prefixo}_{ref:%Y%m%d}.xlsx"
    xlsx = os.path.join(SAIDA_DIR, arquivo)
    try:
        salvar_xlsx(xlsx, tabelas, notas)
    except PermissionError:
        sys.exit(f"[ERRO] não deu para gravar {xlsx}: o arquivo está aberto no Excel?")
    html = html_report.gerar_html(os.path.splitext(xlsx)[0] + ".html", tabelas, notas)
    if args.diario:
        pasta = os.path.join(SAIDA_DIR, "diario")
        print(f"\n[INFO] gravando {len(diario)} relatório(s) diário(s) em {pasta}")
        notas_dia = [n for n in notas if not n.startswith("Histórico:")] + [
            f"Relatório do dia: carteira, ativos, lotes e avisos só com o que existia na data (VCRA e boletas até a "
            f"data, PU dos {DIAS_PU} dias anteriores, CDI do dia); trades desde a data anterior da série."]
        gravar_diario(pasta, diario, trades_fundo, tabelas["Cadastro"], tabelas["Parâmetros"], notas_dia, html_report)
        indice = html_report.gerar_indice(SAIDA_DIR)

    print(f"\n── Resumo {'─' * 40}")
    for r in resumo:
        etiqueta, texto = _situacao(r)
        quando = f" em {r['data']:%d/%m/%Y}" if etiqueta == "OK" else ""
        print(f"[{etiqueta}] {r['fundo']}{quando}: {texto}")
    print(f"[OK] {xlsx}")
    print(f"[OK] {html}")
    if args.diario:
        print(f"[OK] {len(diario)} relatório(s) diário(s) em {pasta}")
        print(f"[OK] {indice}")
    if any(r["status"].startswith("erro") for r in resumo):
        sys.exit(1)


if __name__ == "__main__":
    main()
