from openpyxl import load_workbook
import pandas as pd
import win32com.client
from PIL import ImageGrab
import os

def inject_df_to_excel(
    df, 
    filepath, 
    position=(0,0), 
    sheet_name='Sheet1',
    include_index=True
):
    if include_index:
        aux = df.reset_index()
    else:
        aux = df
    book = load_workbook(filepath)
    with pd.ExcelWriter( 
        filepath, 
        engine='openpyxl',
        mode='a',
        datetime_format='yyyy-mm-dd',
        date_format='yyyy-mm-dd',
    ) as writer:
        writer.book = book
        writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
        worksheet = writer.sheets[sheet_name]
        for column_index, column in enumerate(aux.columns):
            for row_index, row in enumerate(aux.index):
                value = aux.loc[row, column]
                worksheet.cell(
                    row=position[0]+row_index+1, 
                    column=position[1]+column_index+1, 
                ).value = value

def excel_export_img(path_excel,path_img,sheet,init_cell,end_cell):
    path_excel_abs = os.path.abspath(path_excel)
    path_image_abs = os.path.abspath(path_img)
    o = win32com.client.gencache.EnsureDispatch("Excel.Application")
    wb = o.Workbooks.Open(path_excel_abs)
    if sheet:
        ws = wb.Worksheets(sheet)
    else:
        ws = wb.Worksheets("Sheet1")
        
    ws.Range(ws.Cells(init_cell[0],init_cell[1]),ws.Cells(end_cell[0],end_cell[1])).CopyPicture(Format= win32com.client.constants.xlBitmap)  
    img = ImageGrab.grabclipboard()
    img.save(path_image_abs)
    wb.Close(True)
